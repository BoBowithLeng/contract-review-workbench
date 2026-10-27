"""In-memory XLSX checklist template and bounded import (no persistence)."""
from __future__ import annotations

import io
import zipfile
from xml.etree.ElementTree import ParseError

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.utils.exceptions import InvalidFileException

from checklist_store import ChecklistError, legacy_content

MAX_UPLOAD_BYTES = 2 * 1024 * 1024
MAX_UNCOMPRESSED_BYTES = 20 * 1024 * 1024
MAX_NONEMPTY_ROWS = 200
MAX_FIELD_CHARS = 8000
COLUMNS = (
    ("检查事项", "title", "必填，简短描述要检查什么。", "试用期"),
    ("检查内容", "requirement", "必填，写清检查对象、判断条件和具体问题；可用Markdown小标题和列表分层。", "### 期限和次数\n- 合同约定试用期时，结合合同期限核对时长，并核对是否重复约定。"),
    ("适用情形", "when", "选填，只写在哪类交易情形下适用。普遍适用的项目可留空。", "合同约定试用期时。"),
    ("备注", "notes", "选填，其他业务说明。", ""),
)
HEADER_FIELDS = {header: field for header, field, _, _ in COLUMNS}
HEADER_FIELDS.update({"审核要求": "requirement", "核查要求": "requirement",
                      "什么时候检查": "when", "适用条件": "when",
                      "可能需要的信息": "facts", "判断前可能需要的信息": "facts",
                      "处理建议": "guidance", "判定要点": "guidance", "不同情况怎么处理": "guidance"})
IMPORT_FIELDS = ("title", "requirement", "when", "facts", "guidance", "notes")


def _record(message: str, field: str, code: str, row_number: int | None = None) -> dict:
    return {"row_number": row_number, "field": field, "code": code, "message": message}


def _reject(message: str, field: str = "sheet", code: str = "invalid", row_number: int | None = None) -> dict:
    return {"ok": False, "rows": [], "errors": [_record(message, field, code, row_number)], "warnings": []}


def _text(value: object) -> str:
    return "" if value is None else str(value)


def make_template() -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "清单"
    sheet.freeze_panes = "A2"
    sheet.sheet_view.showGridLines = False
    for i, (header, _, _, _) in enumerate(COLUMNS, 1):
        cell = sheet.cell(1, i, header)
        cell.font = Font(name="等线", bold=True, color="20344D")
        cell.fill = PatternFill("solid", fgColor="FFF0D6" if i <= 2 else "E8EEF7")
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        sheet.column_dimensions[get_column_letter(i)].width = 28 if i == 1 else 44
    sheet.row_dimensions[1].height = 30
    sheet.auto_filter.ref = "A1:D1"
    guide = workbook.create_sheet("说明与示例")
    guide.append(["导入只读取“清单”页；前两列必填。示例不导入，公式不执行；最多200项，每字段8000字。"])
    guide.append(["兼容旧模板：信息与处理建议两栏会带原栏标题并入检查内容，不自动去重；合并后超过8000字须人工整理。"])
    guide.append([])
    guide.append(["列名", "填写说明", "示例"])
    for header, _, description, example in COLUMNS:
        guide.append([header, description, example])
    for row in guide:
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    for key, width in (("A", 32), ("B", 60), ("C", 60)):
        guide.column_dimensions[key].width = width
    output = io.BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


make_template_bytes = make_template


def parse_workbook(data: bytes) -> dict:
    if not isinstance(data, bytes) or not data:
        return _reject("上传文件为空或内容无效。")
    if len(data) > MAX_UPLOAD_BYTES:
        return _reject("文件超过2 MiB上限。", "limit", "limit")
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            if sum(info.file_size for info in archive.infolist()) > MAX_UNCOMPRESSED_BYTES:
                return _reject("XLSX解压后超过20 MiB上限。", "limit", "limit")
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=False)
    except (InvalidFileException, zipfile.BadZipFile, OSError, ValueError, KeyError, ParseError, RuntimeError, EOFError):
        return _reject("文件不是可读取的XLSX工作簿。")
    try:
        if "清单" not in workbook.sheetnames:
            return _reject("工作簿缺少名称严格为“清单”的工作表。", "sheet", "sheet")
        sheet = workbook["清单"]
        # Ignore worksheet dimensions supplied by the file; iterate actual rows.
        sheet.reset_dimensions()
        iterator = sheet.iter_rows()
        headers = next(iterator, ())
        legacy_preview = {"通用目的", "类型细化与覆盖边界", "核查要求"} <= {_text(c.value).strip() for c in headers}
        names, seen, mapped = [], set(), {}
        warnings = []
        for index, cell in enumerate(headers):
            if cell.data_type == "f":
                return _reject("表头不能使用公式。", "header", "header", 1)
            name = _text(cell.value).strip()
            names.append(name)
            if name and name in seen:
                return _reject(f"表头存在重复列名：{name}。", "header", "header", 1)
            seen.add(name)
            field = HEADER_FIELDS.get(name)
            if name == "适用条件" and legacy_preview:
                field = "facts"
            if field:
                if field in mapped:
                    return _reject(f"表头存在映射到同一字段的重复列：{name}。", "header", "header", 1)
                mapped[field] = index
                if name not in {column[0] for column in COLUMNS}:
                    destination = {"facts": "检查内容／原可能需要的信息", "guidance": "检查内容／原处理建议"}.get(field)
                    destination = destination or next(c[0] for c in COLUMNS if c[1] == field)
                    message = f"旧列“{name}”已映射为“{destination}”，请确认语义。"
                    if field in {"facts", "guidance"}:
                        message += "旧文字按原栏标题保留，未做语义去重或重写。"
                    if name == "适用条件" and legacy_preview:
                        message += "识别为旧小样成套列，其内容是待了解信息，不自动作为触发条件或必答问题。"
                    warnings.append(_record(message, field, "legacy_column", 1))
            elif name:
                warnings.append(_record(f"“{name}”不自动导入检查内容；有内容的单元格将在警告中列出供确认。", name, "unmapped_column", 1))
        if not {"title", "requirement"} <= set(mapped):
            return _reject("缺少必填列：检查事项、检查内容（兼容审核要求、核查要求）。", "header", "header", 1)
        rows, errors = [], []
        for row_number, cells in enumerate(iterator, 2):
            if not any(_text(c.value).strip() for c in cells):
                continue
            if len(rows) >= MAX_NONEMPTY_ROWS:
                return _reject("非空检查项超过200行上限。", "limit", "limit", row_number)
            item = {"id": "", "row_number": row_number, "enabled": True}
            for field in IMPORT_FIELDS:
                index = mapped.get(field)
                cell = cells[index] if index is not None and index < len(cells) else None
                text = _text(cell.value) if cell is not None else ""
                if cell is not None and cell.data_type == "f":
                    text = ""
                    errors.append(_record("单元格包含公式，未执行也未导入；请改为文字或数值。", field, "formula", row_number))
                elif len(text) > MAX_FIELD_CHARS:
                    errors.append(_record("字段超过8000字上限。", field, "too_long", row_number))
                if field in {"title", "requirement"} and not text.strip():
                    errors.append(_record("必填内容为空。", field, "required", row_number))
                item[field] = text
            if item["requirement"].strip():
                try:
                    item["requirement"] = legacy_content(item)
                    item["facts"] = item["guidance"] = ""
                except ChecklistError as exc:
                    # Keep the original fields in this rejected preview row;
                    # never truncate useful old content to make it importable.
                    errors.append(_record(str(exc), "requirement", "legacy_merge", row_number))
            mapped_indices = set(mapped.values())
            for index, cell in enumerate(cells):
                if index in mapped_indices or not _text(cell.value).strip():
                    continue
                name = names[index] if index < len(names) and names[index] else f"未命名第{index + 1}列"
                if cell.data_type == "f":
                    errors.append(_record("未导入列含公式，未执行；请改为文字后确认内容。", name, "formula", row_number))
                else:
                    value = _text(cell.value)
                    if len(value) > MAX_FIELD_CHARS:
                        errors.append(_record("字段超过8000字上限。", name, "too_long", row_number))
                    warnings.append({**_record(f"“{name}”未自动导入，请确认残余内容。", name, "unmapped_content", row_number), "value": value[:MAX_FIELD_CHARS]})
            rows.append(item)
        return {"ok": True, "rows": rows, "errors": errors, "warnings": warnings}
    except (ValueError, KeyError, ParseError, zipfile.BadZipFile, OSError, RuntimeError, EOFError):
        return _reject("工作簿内容损坏，无法读取。")
    finally:
        workbook.close()
