"""Small synthetic delivery demonstration, not a canned AI/legal audit."""
from __future__ import annotations

import argparse
import datetime as dt
import io

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from docx_delivery import inspect_source
from delivery_bundle import write_bundle


def source_document() -> bytes:
    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = "宋体"
    normal.font.size = Pt(12)
    normal.element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "宋体")
    normal.paragraph_format.space_after = Pt(10)
    for section in doc.sections:
        section.top_margin = section.bottom_margin = Cm(2)
    doc.add_heading("劳动合同节选 · 合成交付示例", 0)  # P000
    doc.add_paragraph("本文件只演示文字修改与批注的交付，不是可直接签署的完整合同或法律审查结果。")  # P001
    doc.add_paragraph("甲方：星河示例有限公司　乙方：张示例")  # P002
    doc.add_heading("一、工资与交接", 1)  # P003
    p = doc.add_paragraph()  # P004: deliberately split one target across runs.
    p.add_run("离职工资将在")
    run = p.add_run("工作交接")
    run.bold = True
    run.font.color.rgb = RGBColor.from_string("25496A")
    run = p.add_run("完成后")
    run.underline = True
    p.add_run("结算。")
    doc.add_heading("二、通知与联系人", 1)  # P005
    p = doc.add_paragraph("双方确认通知方式为电子邮件。")  # P006
    p.add_run("\n示例备注：").italic = True
    p.add_run("保留此行及其中的原有格式。")
    doc.add_paragraph("重复说明：双方应核实实际通知渠道。")  # P007
    table = doc.add_table(rows=2, cols=2)
    table.style = "Table Grid"
    table.cell(0, 0).text = "项目"
    table.cell(0, 1).text = "约定"
    table.cell(1, 0).text = "业务联系人"
    table.cell(1, 1).text = "旧联系人"
    doc.add_paragraph("技术验收签记：\t只改变明确选中的文字。")  # P008
    stream = io.BytesIO()
    doc.save(stream)
    return stream.getvalue()


def demo_plan(data: bytes) -> dict:
    source = inspect_source(data)
    pieces = {p["id"]: p for p in source["pieces"]}

    def edit(edit_id, piece_id, old, new, *, insert_after=None):
        piece = pieces[piece_id]
        if insert_after is None:
            assert piece["text"].count(old) == 1
            start = piece["text"].index(old)
        else:
            start = piece["text"].index(insert_after) + len(insert_after)
        return {"id": edit_id, "piece_id": piece_id, "paragraph_sha256": piece["paragraph_sha256"],
                "start": start, "end": start + len(old), "old_text": old, "new_text": new, "confirmed": True}

    return {
        "schema_version": 1, "source_sha256": source["source_sha256"], "revision": 1,
        "author": "交付功能演示", "created_at": dt.datetime.now(dt.timezone.utc).isoformat(),
        "edits": [
            edit("E1", "P004", "工作交接完成后", "约定支付日"),
            edit("E2", "T0R1C1P0", "旧联系人", "陈示例"),
            edit("E3", "P006", "", "及短信", insert_after="电子邮件"),
            edit("E4", "P007", "重复说明：", ""),
        ],
        "comments": [
            {"id": "C1", "edit_id": "E1", "text": "请财务确认具体支付日；这里仅演示修订与批注的关联。", "confirmed": True},
            {"id": "C2", "edit_id": "E2", "text": "请业务部门核对实际联系人。", "confirmed": True},
            {"id": "C3", "edit_id": "E3", "text": "请确认是否实际使用短信作为通知渠道。", "confirmed": True},
        ],
    }


def main():
    parser = argparse.ArgumentParser(description="生成合成合同的洁净／修订／批注／报告，输出目录须不存在")
    parser.add_argument("output_dir")
    args = parser.parse_args()
    data = source_document()
    manifest = write_bundle(data, demo_plan(data), args.output_dir)
    print("已生成合成技术示例：" + args.output_dir)
    for item in manifest["files"]:
        print(item["name"])


if __name__ == "__main__":
    main()
