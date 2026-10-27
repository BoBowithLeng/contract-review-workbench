"""Derive presentation from source text and Word metadata, without changing it."""
import copy
import hashlib
import re
import io
import threading
from collections import OrderedDict
from html import escape

# Opening a report asks for it three times and every later request repeats the same source; parsing the
# Word file with python-docx took ~130 ms each time. The result depends only on the bytes.
_NAVIGATION_CACHE = OrderedDict()
_NAVIGATION_LOCK = threading.Lock()


def docx_headings(data):
    """Read explicit Word heading styles; IDs match body paragraphs of parser."""
    from docx import Document
    result = {}
    for i, paragraph in enumerate(Document(io.BytesIO(data)).paragraphs):
        level = _word_heading_level(paragraph)
        if level:
            result[f'P{i:03d}'] = level
    return result


def _word_heading_level(paragraph):
    from docx.oxml.ns import qn
    level = None
    ppr = paragraph._p.pPr
    if ppr is not None:
        value = ppr.find(qn('w:outlineLvl'))
        if value is not None:
            level = int(value.get(qn('w:val')))
    style = paragraph.style
    seen = set()
    while level is None and style is not None and style.style_id not in seen:
        seen.add(style.style_id)
        pr = style.element.find(qn('w:pPr'))
        node = pr.find(qn('w:outlineLvl')) if pr is not None else None
        if node is not None:
            level = int(node.get(qn('w:val')))
            break
        match = re.search(r'(?:Heading|标题)\s*(\d+)', style.name or '', re.I)
        if match:
            level = int(match.group(1))-1
            break
        style = style.base_style
    return level + 1 if level is not None and 0 <= level < 9 and paragraph.text.strip() else 0


def docx_navigation(data):
    """Presentation-only outline and rendered list labels, cached by source bytes; callers get a copy."""
    key = hashlib.sha256(data).hexdigest()
    with _NAVIGATION_LOCK:
        cached = _NAVIGATION_CACHE.get(key)
        if cached is not None:
            _NAVIGATION_CACHE.move_to_end(key)
    if cached is None:
        cached = _docx_navigation_uncached(data)
        with _NAVIGATION_LOCK:
            _NAVIGATION_CACHE[key] = cached
            while len(_NAVIGATION_CACHE) > 16:
                _NAVIGATION_CACHE.popitem(last=False)
    return copy.deepcopy(cached)


def _docx_navigation_uncached(data):
    """Presentation-only outline and rendered list labels keyed by parser IDs.

    Walk table-cell paragraphs too: list counters may continue across tables.
    Neither a number nor a short clause supplies evidence of a heading.
    """
    from docx import Document
    from docx.oxml.ns import qn
    from docx.text.paragraph import Paragraph
    from word_numbering import WordNumbering
    document = Document(io.BytesIO(data))
    numbering = WordNumbering(document)
    result = {'headings': {}, 'numbering': {}}

    def read(pid, node, body_heading=False):
        paragraph = Paragraph(node, document)
        level = _word_heading_level(paragraph) if body_heading else 0
        if level:
            result['headings'][pid] = level
        label = numbering.label(node)
        if label:
            result['numbering'][pid] = label
        else:
            raw_numbering = numbering.paragraph_numbering(node)
            if raw_numbering:
                num_id, depth = raw_numbering
                result['numbering'][pid] = {'label': '', 'level': depth, 'num_id': num_id}

    paragraph_index = table_index = 0
    for node in document.element.body:
        if node.tag == qn('w:p'):
            read(f'P{paragraph_index:03d}', node, True)
            paragraph_index += 1
        elif node.tag == qn('w:tbl'):
            for r, row in enumerate(node.findall(qn('w:tr'))):
                for c, cell in enumerate(row.findall(qn('w:tc'))):
                    for p, paragraph in enumerate(cell.findall(qn('w:p'))):
                        read(f'T{table_index}R{r}C{c}P{p}', paragraph)
            table_index += 1
    return result


def index_document(contract):
    pieces = contract.pieces if hasattr(contract, "pieces") else contract["pieces"]
    styles = {} if hasattr(contract, 'pieces') else contract.get('headings', {}) or {}
    numbering = {} if hasattr(contract, 'pieces') else contract.get('numbering', {}) or {}
    explicit = bool(styles)
    headings = {}  # Word outline levels, or conservative standalone chapter/section titles.
    heading_ids = {}
    list_path = {}
    list_id = None
    article = ""  # Position labels are not navigation entries.
    sub = ""
    result = []
    for i, item in enumerate(pieces):
        p = item.model_dump() if hasattr(item, "model_dump") else item
        text = p["text"]
        stripped = text.strip()
        number = numbering.get(p['id'], {})
        label = number.get('label', '') if isinstance(number, dict) else ''
        label = label if isinstance(label, str) else ''
        numbering_level = number.get('level') if isinstance(number, dict) else None
        if type(numbering_level) is not int or not 0 <= numbering_level <= 8:
            numbering_level = None
        display_text = (label + (' ' if not label[-1:].isspace() else '') + text) if label else text
        title = display_text.strip()
        level = 0
        styled = styles.get(p["id"])
        if type(styled) is int and 1 <= styled <= 9 and stripped:
            level = styled
        elif not explicit:
            # Only standalone chapter/section titles are inferred. Numbered
            # clauses and list items may be useful position labels, not headings.
            chapter_title = re.fullmatch(r"第[一二三四五六七八九十百零〇\d]+([章节])[ \t　:：、.．-]*(.{1,24})", stripped)
            if (chapter_title and len(stripped) <= 36
                    and not re.search(r"[。；;！？?]", stripped)):
                level = 1 if chapter_title.group(1) == "章" else 2
        if level:
            headings = {depth: title for depth, title in headings.items() if depth < level}
            heading_ids = {depth: pid for depth, pid in heading_ids.items() if depth < level}
            heading_parent = heading_ids[max(heading_ids)] if heading_ids else None
            heading_ids[level] = p['id']
            headings[level] = title
            article = sub = ""
            list_path = {}
            list_id = number.get('num_id') if label else None
        else:
            heading_parent = None
            if label and numbering_level is not None:
                current_list = number.get('num_id')
                if current_list != list_id:
                    list_path = {}
                list_id = current_list
                list_path = {depth: value for depth, value in list_path.items() if depth < numbering_level}
                list_path[numbering_level] = label
                article = sub = ''
            elif numbering_level is not None:
                # The source has numbering but its label is unsupported. Do not
                # assign this clause the preceding clause's rendered number.
                list_path = {}
            article_match = re.match(r"^(第[一二三四五六七八九十百零〇\d]+条)(?:[ \t　:：、.．]*)", stripped)
            ordinal = re.match(r"^([一二三四五六七八九十]+[、.．])\s*(.+)", stripped)
            if label and numbering_level is not None:
                pass  # Word numbering already supplied this clause's position.
            elif article_match:
                list_path = {}
                article, sub = article_match.group(1), ""
            elif ordinal and len(stripped) <= 25 and not re.search(r"[。；;！？?]|应当|应|不得|须|不存在", stripped):
                list_path = {}
                article, sub = stripped, ""
            else:
                decimal = re.match(r"^(\d+(?:[.．]\d+)+)(?!\d)", stripped)
                item_number = re.match(r"^([（(][一二三四五六七八九十\d]+[）)]|\d+[、.．])(?!\d)", stripped)
                if decimal or item_number:
                    sub = (decimal or item_number).group(1)
        path = [headings[depth] for depth in sorted(headings)]
        if article:
            path.append(article)
        if sub:
            path.append(sub)
        path.extend(list_path[depth] for depth in sorted(list_path))
        location = " / ".join(path) or f"正文第{i + 1}段"
        cell = re.fullmatch(r"T(\d+)R(\d+)C(\d+)P(\d+)", p["id"])
        if cell:
            t, r, c, _ = map(int, cell.groups())
            location = (" / ".join(path) + " / " if path else "") + f"表格{t+1}·第{r+1}行第{c+1}列"
        result.append({**p, "location": location, "level": level,
                       "display_text": display_text, "numbering_label": label,
                       "numbering_level": numbering_level, "heading_parent": heading_parent})
    return result


def sort_findings_by_document(findings, contract):
    """Return findings in first-cited source order without changing stable IDs.

    Exact quote offsets break ties inside one paragraph. Findings without a
    usable reference remain at the end in their original relative order.
    """
    pieces = contract.pieces if hasattr(contract, "pieces") else contract.get("pieces", [])
    order = {}
    text_by_id = {}
    for index, item in enumerate(pieces):
        piece = item.model_dump() if hasattr(item, "model_dump") else item
        if isinstance(piece, dict) and isinstance(piece.get("id"), str):
            order[piece["id"]] = index
            text_by_id[piece["id"]] = piece.get("text", "")

    def key(entry):
        original_index, finding = entry
        locations = []
        for ref in finding.get("refs", []) if isinstance(finding, dict) else []:
            if not isinstance(ref, dict) or ref.get("id") not in order:
                continue
            text = text_by_id[ref["id"]]
            quote = ref.get("quote")
            offset = text.find(quote) if isinstance(text, str) and isinstance(quote, str) else -1
            locations.append((order[ref["id"]], offset if offset >= 0 else len(text)))
        return (*min(locations), original_index) if locations else (float("inf"), float("inf"), original_index)

    return [finding for _, finding in sorted(enumerate(list(findings or [])), key=key)]


def reader_markup(contract, items, selected):
    """Escape all document/model content. Only exact quote ranges get marked."""
    doc = index_document(contract)
    by_piece = {}
    for item in items:
        for ref in item.get("refs", []):
            by_piece.setdefault(ref["id"], []).append((item["id"], ref["quote"]))
    sections, blocks = [], []
    first = None
    for p in doc:
        pid, text = p["id"], p["text"]
        if not text.strip():
            continue
        marks = by_piece.get(pid, [])
        selected_marks = [(fid, q) for fid, q in marks if fid == selected and q in text]
        ranges = []
        for fid, q in (selected_marks or marks):
            if not q or q not in text:
                continue
            start = text.index(q)
            ranges.append((start, start+len(q)))
        merged = []
        for start, end in sorted(ranges):
            if merged and start <= merged[-1][1]:
                merged[-1] = (merged[-1][0], max(end, merged[-1][1]))
            else:
                merged.append((start, end))
        chunks, cursor = [], 0
        for start, end in merged:
            chunks += [escape(text[cursor:start]), '<mark>' + escape(text[start:end]) + '</mark>']
            cursor = end
        chunks.append(escape(text[cursor:]))
        cls = "paragraph selected" if selected_marks else "paragraph"
        if selected_marks and first is None:
            first = pid
        buttons = "".join(f'<button type="button" data-issue="{escape(fid, quote=True)}">查看意见 {escape(fid)}</button>' for fid in dict.fromkeys(fid for fid, _ in marks))
        if p["level"]:
            sections.append({"id":pid, "text":p['display_text'].strip()})
        prefix = escape(p['numbering_label']) + ' ' if p['numbering_label'] else ''
        blocks.append(f'<section id="p-{escape(pid,quote=True)}" class="{cls}"><div class="location">{escape(p["location"])}</div><div class="body level-{p["level"]}">{prefix}{"".join(chunks)}</div>{buttons}</section>')
    return {"html":"".join(blocks), "target":"p-"+first if first else "", "sections":sections}
