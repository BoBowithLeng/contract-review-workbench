"""Render one confirmed document plan into a consistent, local delivery bundle.

This module never infers edits from review decisions and never calls a model.
The original bytes and explicit confirmation plan are the common inputs.
"""
from __future__ import annotations

import argparse
import hashlib
import io
import json
from pathlib import Path

from docx import Document
from docx.shared import Pt
from docx.oxml.ns import qn

from contract_core import parse_docx
from docx_delivery import inspect_source, render_docx, validate_plan
from document_view import index_document
from report_export import export_docx


FILENAMES = {
    "clean": "01-洁净版.docx",
    "tracked": "02-修订版.docx",
    "comments": "03-原稿批注版.docx",
    "tracked_comments": "04-修订与批注版.docx",
    "report": "05-合同审查报告.docx",
}
AUDIT_REPORT_FILENAME = "05-交付核对报告.docx"


def sha256(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def canonical_plan_bytes(plan: dict) -> bytes:
    return json.dumps(plan, ensure_ascii=False, sort_keys=True, separators=(",", ":")).encode("utf-8")


def delivery_report(data: bytes, plan: dict, review_bundle: dict | None = None) -> bytes:
    """Export a readable review; keep technical audits for the standalone CLI."""
    source = inspect_source(data)
    contract = parse_docx(data).model_dump()
    locations = {p["id"]: p["location"] for p in index_document(contract)}
    if review_bundle is not None:
        if review_bundle.get("contract", {}).get("fingerprint") != contract["fingerprint"]:
            raise ValueError("审查报告的合同文字与本次原稿不一致，不能拼入这批交付。")
        report = review_bundle.get("report", {})
        if report.get("fingerprint") != contract["fingerprint"]:
            raise ValueError("审查结果的合同指纹与本次原稿不一致。")
        return export_docx(report, review_bundle["sources"], contract,
                          reviews=review_bundle.get("reviews"),
                          selected_options=review_bundle.get("selected_options"), reader_mode=True,
                          reader_actions=review_bundle.get("actions"),
                          supplemental_cases=review_bundle.get("supplemental_cases"),
                          source_docx=data, confirmed_plan=plan)
    else:
        doc = Document()
        doc.add_heading("文档交付核对报告", 0)
        doc.add_paragraph("本次仅验证文件交付，没有运行法律审核；以下是明确确认的文字修改与协作批注，不是自动生成的法律结论。")
    normal = doc.styles["Normal"]
    normal.font.name = "宋体"
    normal.font.size = Pt(11)
    normal.element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "宋体")
    plan_hash = sha256(canonical_plan_bytes(plan))
    doc.add_paragraph("原稿 SHA-256：" + source["source_sha256"])
    doc.add_paragraph("确认版本：" + str(plan["revision"]) + "；确认人：" + plan["author"])
    doc.add_paragraph("修改计划 SHA-256：" + plan_hash)
    doc.add_paragraph("确认时间：" + plan["created_at"])
    doc.add_heading("已确认的文档修改", 1)
    if not plan["edits"]:
        doc.add_paragraph("本次不修改正文。")
    for edit in plan["edits"]:
        kind = ("新增独立段落" if edit.get("kind") == "insert_paragraph_after" else
                "插入" if edit["start"] == edit["end"] else "删除" if not edit["new_text"] else "替换")
        doc.add_heading(edit["id"] + " · " + kind, 2)
        doc.add_paragraph("原文位置：" + locations.get(edit["piece_id"], edit["piece_id"]))
        doc.add_paragraph("原文：" + (edit["old_text"] or "（在指定位置插入）"))
        doc.add_paragraph("确认改为：" + (edit["new_text"] or "（删除）"))
        if edit.get("finding_ids"):
            doc.add_paragraph("关联审阅意见：" + "、".join(edit["finding_ids"]))
        if edit.get("option_ids"):
            doc.add_paragraph("关联方案：" + "、".join(edit["option_ids"]))
    doc.add_heading("已确认的协作批注", 1)
    if not plan["comments"]:
        doc.add_paragraph("本次不附批注。")
    for comment in plan["comments"]:
        anchor = ("关联修改 " + comment["edit_id"]) if "edit_id" in comment else locations.get(comment["piece_id"], comment["piece_id"])
        doc.add_paragraph(comment["id"] + " · " + anchor)
        doc.add_paragraph(comment["text"])
    doc.add_heading("文件之间如何对应", 1)
    doc.add_paragraph("洁净版应用本批确认修改；修订版保留相同修改的原生接受／拒绝记录；原稿批注版不改正文；修订与批注版把两者组合。它们使用同一原稿与确认版本，未确认内容不混入正文。")
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def build_bundle(data: bytes, raw_plan: dict, *, variants=None, review_bundle=None) -> dict:
    """Validate everything before returning any files; no filesystem effects."""
    variants = list(FILENAMES) if variants is None else list(variants)
    if not variants or len(set(variants)) != len(variants) or any(v not in FILENAMES for v in variants):
        raise ValueError("请选择不重复的 clean/tracked/comments/tracked_comments/report 输出。")
    plan = validate_plan(data, raw_plan)
    plan_hash = sha256(canonical_plan_bytes(plan))
    filenames = {**FILENAMES}
    if review_bundle is None:
        filenames["report"] = AUDIT_REPORT_FILENAME
    files = {}
    for variant in variants:
        files[filenames[variant]] = (delivery_report(data, plan, review_bundle) if variant == "report"
                                      else render_docx(data, plan, variant))
    manifest = {
        "schema_version": 1,
        "source_sha256": sha256(data),
        "plan_sha256": plan_hash,
        "snapshot_revision": plan["revision"],
        "author": plan["author"],
        "created_at": plan["created_at"],
        "edit_ids": [edit["id"] for edit in plan["edits"]],
        "comment_ids": [comment["id"] for comment in plan["comments"]],
        "legal_review_included": review_bundle is not None and "report" in variants,
        "model_called_for_export": False,
        "files": [{"variant": v, "name": filenames[v], "sha256": sha256(files[filenames[v]])}
                  for v in variants],
    }
    return {"files": files, "manifest": manifest, "plan": plan, "source": inspect_source(data)}


def write_bundle(data: bytes, raw_plan: dict, output_dir, *, variants=None, review_bundle=None) -> dict:
    """Write a new directory only. A manifest appears last, after all files exist."""
    built = build_bundle(data, raw_plan, variants=variants, review_bundle=review_bundle)
    target = Path(output_dir).resolve()
    target.mkdir(parents=True, exist_ok=False)
    payloads = {"00-原稿.docx": data, **built["files"],
                "confirmed-plan.json": json.dumps(built["plan"], ensure_ascii=False, indent=2).encode("utf-8"),
                "source-snapshot.json": json.dumps(built["source"], ensure_ascii=False, indent=2).encode("utf-8")}
    for name, content in payloads.items():
        with (target / name).open("xb") as stream:
            stream.write(content)
    with (target / "manifest.json").open("x", encoding="utf-8") as stream:
        json.dump(built["manifest"], stream, ensure_ascii=False, indent=2)
    return built["manifest"]


def main():
    parser = argparse.ArgumentParser(description="按已确认的精确修改计划生成本机Word交付；不调用模型")
    commands = parser.add_subparsers(dest="command", required=True)
    inspect = commands.add_parser("inspect")
    inspect.add_argument("source")
    export = commands.add_parser("export")
    export.add_argument("source")
    export.add_argument("plan")
    export.add_argument("output_dir")
    export.add_argument("--variants", nargs="+", choices=list(FILENAMES), default=list(FILENAMES))
    export.add_argument("--review-bundle", help="可选：同一合同文字的既有审核结果bundle JSON")
    args = parser.parse_args()
    data = Path(args.source).read_bytes()
    if args.command == "inspect":
        print(json.dumps(inspect_source(data), ensure_ascii=False, indent=2))
    else:
        plan = json.loads(Path(args.plan).read_text(encoding="utf-8-sig"))
        review = json.loads(Path(args.review_bundle).read_text(encoding="utf-8-sig")) if args.review_bundle else None
        result = write_bundle(data, plan, args.output_dir, variants=args.variants, review_bundle=review)
        print(json.dumps(result, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
