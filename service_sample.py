"""Reproducible synthetic input and two instructions; no model answers."""
from __future__ import annotations
import argparse
import json
from pathlib import Path
from docx import Document
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

PARAGRAPHS = [
"小程序开发服务合同（合成样例）",
"本文件仅用于产品调试，主体、项目及交易均为虚构，未经针对具体交易的法律审查，不作为对外签署模板。",
"甲方：澄岸商贸（演示）有限公司；项目联系人：陈样；联系邮箱：client@example.com。",
"乙方：栖码软件（演示）有限公司；项目联系人：周样；联系邮箱：provider@example.com。",
"双方约定由乙方向甲方提供会员预约小程序开发服务，项目在上海实施，适用中华人民共和国大陆地区法律。",
"第一条 项目范围",
"1.1 本期功能包括会员手机号登录、服务项目展示、预约及取消、后台排期管理、预约记录查询与导出。支持一个经营主体、三家门店；不含在线支付、直播和多商户结算。",
"1.2 双方在项目启动后确认页面原型和功能清单，作为开发及验收的参考。项目群中甲方提出的合理调整由乙方配合完成，原有功能优化不另收费；新增独立模块由双方另行确认。",
"1.3 甲方负责提供合法使用的品牌素材、服务内容、测试账号及平台认证资料；乙方负责开发、测试和一次生产环境部署。双方各指定一名项目联系人，重要事项通过本合同邮箱或双方项目群确认。",
"第二条 进度与配合",
"2.1 自乙方收到首付款且甲方提供启动所需资料之次日起计算，乙方在10个自然日内提交页面原型、45个自然日内提交可测试版本、60个自然日内完成部署。",
"2.2 因甲方未提供资料或未确认原型影响进度的，乙方应及时列明待配合事项；双方沟通后调整相应进度。平台审核时间不计入乙方开发时间。",
"第三条 价款与支付",
"3.1 合同总价为人民币120000元（含税），包含约定开发、一次部署及本合同第五条维护。平台认证、短信、域名和云资源费用由甲方按第三方实际收费另行承担，采购前由乙方书面列明并经甲方确认。",
"3.2 甲方在合同签订后5个工作日内支付总价40%即48000元；乙方提交可测试版本后5个工作日内支付总价40%即48000元；最终验收通过后5个工作日内支付剩余20%即24000元。乙方应按约提供合法有效发票及收款信息。",
"第四条 交付与验收",
"4.1 乙方以邮件发送测试地址和操作说明，请甲方验收。甲方应在收到后7个自然日内一次性书面提出修改意见；逾期未反馈的，视为验收通过。甲方投入实际经营使用的，也视为验收通过。",
"4.2 乙方对属于本期功能范围的缺陷予以修复后再次提交。双方未就多次提交后的验收截止时间另作约定。",
"4.3 乙方在收到全部合同款后3个工作日内交付本项目源代码、部署说明和后台管理账号。甲方可将本项目成果用于自身经营及委托其他服务商维护；乙方已有通用组件不转让，但应提供本项目持续运行所需的使用权限。",
"第五条 维护与保密",
"5.1 验收通过后12个月内，乙方免费修复约定功能的程序缺陷。新增功能及第三方平台重大调整引起的改造，双方另行确认费用与工期；日常维护由项目联系人协调。",
"5.2 双方对履行本合同中获知的非公开经营资料承担保密义务，不向项目无关人员披露。演示、测试优先使用虚构或脱敏数据；确需使用真实业务数据的，双方另行确认数据范围、权限及处理安排。",
"第六条 违约责任",
"6.1 一方未按约付款或交付，每逾期一日按逾期部分金额的万分之三支付违约金，累计不超过合同总价10%；因对方未配合造成的相应延误不计入逾期。",
"6.2 因乙方服务不符合约定给甲方造成损失的，乙方应予赔偿。甲方应采取合理措施避免损失扩大；损失、原因及金额由提出请求的一方提供相应材料。",
"第七条 提前结束与交接",
"7.1 任一方可提前10个自然日书面通知对方结束项目。合同结束后，已支付款项不予退还，乙方将当时已经完成的成果按现状交付给甲方。",
"7.2 因一方违约解除合同的，守约方仍可依据第六条主张责任。双方应在项目结束后配合移交已形成的资料、账号及双方确认应交付的成果。",
"第八条 其他",
"8.1 本合同未尽事项由双方协商补充；书面补充约定与本合同具有同等效力。履行争议先协商，协商不成向甲方住所地有管辖权的人民法院起诉。",
"甲方签章：________    乙方签章：________",
"签订日期：2026年10月1日",
]
BRIEFS = {
 "client": {
  "title":"客户委托要求", "our_role":"client", "our_party":"甲方",
  "business_goal":"我方是甲方澄岸商贸，采购会员预约小程序，预算12万元，希望在2026年12月15日活动前稳定使用。可以接受分期付款，不要求供应商承担所有经营风险；但不希望测试地址一发就付到80%、系统不能正常使用却已被视为验收，也需要后续更换维护人员的实际能力。请优先看交付和验收条件、付款对应成果以及中途结束如何结算，给出能谈成的修改，不要只改成对乙方最苛刻的版本。",
 },
 "provider": {
  "title":"服务商委托要求", "our_role":"provider", "our_party":"乙方",
  "business_goal":"我方是乙方栖码软件，三人团队承接固定总价12万元的小程序项目，报价只含合同列明的本期功能和缺陷修复。需要阶段回款以覆盖人员投入，担心客户通过项目群不断追加修改、拖着不验收，以及项目取消后已完成工作和已投入成本无法结算。愿意修复自身缺陷并对自身违约负责，但不接受没有边界的免费新增工作或无法预估的责任。请在可执行范围内保护回款、工作范围和退出结算，不要建议我们可以任意拒绝交付。",
 }
}

def build_document():
 doc=Document()
 normal=doc.styles["Normal"];normal.font.name="宋体";normal.font.size=Pt(11)
 normal.element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"),"宋体")
 normal.paragraph_format.space_after=Pt(6)
 for section in doc.sections:
  section.top_margin=section.bottom_margin=Cm(2.1)
 for index,text in enumerate(PARAGRAPHS):
  if index==0:doc.add_heading(text,0)
  elif text.startswith("第") and "条 " in text:doc.add_heading(text,1)
  else:doc.add_paragraph(text)
 return doc

def main():
 parser=argparse.ArgumentParser(description="生成合成服务合同与两份委托，不生成审核答案")
 parser.add_argument("directory")
 args=parser.parse_args();root=Path(args.directory);root.mkdir(parents=True,exist_ok=True)
 targets=[root/"小程序开发服务合同（合成样例）.docx",root/"小程序开发服务合同（合成样例）.md",root/"委托要求-客户.md",root/"委托要求-服务商.md",root/"briefs.json"]
 if any(p.exists() for p in targets):raise FileExistsError("样例文件已存在，不覆盖")
 build_document().save(targets[0])
 targets[1].write_text("\n\n".join(PARAGRAPHS)+"\n",encoding="utf-8")
 for key,label in [("client","客户"),("provider","服务商")]:
  brief=BRIEFS[key]
  (root/f"委托要求-{label}.md").write_text(f"# {brief['title']}\n\n同一份《小程序开发服务合同（合成样例）》；我方称谓：{brief['our_party']}。\n\n{brief['business_goal']}\n",encoding="utf-8")
 contexts={key:{"contract_type":"software_service","our_role":value["our_role"],"our_party":value["our_party"],"business_goal":value["business_goal"],"review_date":"2026-09-22","region":"上海","input_kind":"拟签审阅"} for key,value in BRIEFS.items()}
 targets[-1].write_text(json.dumps(contexts,ensure_ascii=False,indent=2)+"\n",encoding="utf-8")
 print(targets[0].resolve())

if __name__=="__main__":main()
