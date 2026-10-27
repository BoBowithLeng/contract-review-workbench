"""Developer-only synthetic contract and canned model answers for offline tests.

The app never imports or bundles these fixtures. Public desktop builds start
with an empty workbench; these inputs do not represent real model output.
"""
from __future__ import annotations

import datetime as dt
import io
import json
import zipfile

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

SYNTHETIC_NOTICE = "合成测试合同 / 非真实交易"

# Every value below is invented. There are no real people, locations or contacts.
SYNTHETIC_CONTRACT = {
    "name": "合成测试 软件开发合同.docx",
    "title": "软件开发合同 合成测试",
    "checklist_id": "builtin-software-client-v02",
    "clauses": [
        "1.1 甲方为示例客户公司，乙方为示例开发公司。本合同约定开发内部工单管理软件，不处理真实个人信息。",
        "1.2 乙方开发工单登记、状态查询和统计导出功能，详细需求由双方后续口头确认。",
        "2.1 项目总价为人民币30000元。甲方应在合同签署后支付全部价款。",
        "3.1 乙方应在启动后30日交付测试版本。甲方收到后3日未提出书面异议，即视为验收合格。",
        "4.1 乙方仅交付可运行程序，无须提供源代码、构建说明或者管理员账号。",
        "5.1 本项目全部新开发成果归乙方所有，甲方仅可在合同有效期内使用。",
        "6.1 乙方提供30日缺陷修复支持。超出支持期的服务由双方另行协商。",
        "7.1 双方通过指定项目负责人确认履行事项。争议由双方先行协商解决。",
    ],
    # (checklist item, clause index, title, risk, reason, replacement)
    "opinions": [
        ("C03", 1, "需求确认缺少可核对的依据", "中", "口头需求缺乏版本和范围边界。", "1.2 乙方开发工单登记、状态查询和统计导出功能，功能范围及完成标准以双方确认的需求附件为准。"),
        ("C09", 2, "付款早于交付", "高", "全部价款先付。", "2.1 项目总价为人民币30000元。甲方在验收合格后支付全部价款。"),
        ("C10", 3, "沉默视为验收的期限过短", "高", "3日不足以完成测试。", "3.1 乙方应在启动后30日交付测试版本。甲方在收到后10个工作日内测试并书面反馈。"),
        ("C12", 4, "缺少接手材料", "中", "仅交程序会依赖开发方。", "4.1 乙方应交付源代码、构建说明及管理员账号清单。"),
        ("C18", 5, "成果使用期限受限", "中", "合同结束即不能使用成果。", "5.1 甲方在付清价款后可永久内部使用、修改本项目新开发成果。"),
    ],
}


def contract_bytes(sample: dict = SYNTHETIC_CONTRACT) -> bytes:
    """A small, reproducible DOCX with explicit fonts and fixed package times."""
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    for name in ("Normal", "Title"):
        style = document.styles[name]
        style.font.name = "宋体" if name == "Normal" else "黑体"
        style.font.size = Pt(11 if name == "Normal" else 18)
        style.font.color.rgb = RGBColor(0, 0, 0)
        fonts = style.element.get_or_add_rPr().rFonts
        fonts.set(qn("w:eastAsia"), style.font.name)
        for attribute in list(fonts.attrib):
            if attribute.lower().endswith("theme"):
                del fonts.attrib[attribute]
    document.add_paragraph(sample["title"], "Title")
    document.add_paragraph(SYNTHETIC_NOTICE + "。全部当事人、金额及安排均为虚构，仅用于程序测试。")
    for text in sample["clauses"]:
        document.add_paragraph(text)
    props = document.core_properties
    props.author = props.last_modified_by = "合成测试"
    props.created = props.modified = dt.datetime(2026, 10, 7, tzinfo=dt.timezone.utc)
    raw = io.BytesIO()
    document.save(raw)
    output = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(raw.getvalue())) as source, zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as target:
        for name in source.namelist():
            info = zipfile.ZipInfo(name, (2026, 10, 7, 0, 0, 0))
            info.compress_type = zipfile.ZIP_DEFLATED
            target.writestr(info, source.read(name))
    return output.getvalue()


def fixture_runner(sample: dict = SYNTHETIC_CONTRACT):
    """Answer the preparation and review phases for the synthetic contract only."""
    def complete(messages, config):
        if config.get("review_phase") == "review_preparation":
            return json.dumps({"questions": [], "requirements": [], "additional_checks": []})
        packet = json.loads(messages[1]["content"])
        pieces = {piece["text"]: piece["id"] for piece in packet["contract"]}
        findings, by_check = [], {}
        for index, (cid, clause_index, title, risk, reason, new_text) in enumerate(sample["opinions"], 1):
            fid = f"F{index:02d}"
            text = sample["clauses"][clause_index]
            pid = pieces[text]
            findings.append({
                "id": fid, "rule_id": f"G{index:02d}", "theme": "合同安排", "title": title,
                "issue_kind": "商业", "status": "发现风险", "risk": risk, "evidence_status": "原文可核",
                "refs": [{"id": pid, "quote": text}], "location_note": "", "source_ids": [],
                "reason": reason, "suggestion": "采用下列替换文字。", "questions": [],
                "draft_plan": {"status": "ready", "edits": [{"piece_id": pid, "new_text": new_text}],
                               "insertions": [], "paragraph_actions": [{"piece_id": pid, "action": "modify", "reason": reason}],
                               "questions": [], "comment": title},
            })
            by_check[cid] = [fid]
        rows = []
        for check in packet["review_preparation"]["checks"]:
            ids = by_check.get(check["id"], [])
            # Only the first group carries the linked opinion; the rest stay open.
            groups = [{"group_id": group["group_id"], "title": group["title"],
                       "status": "发现风险" if ids and i == 0 else "需要法务判断",
                       "reason": "合成测试。", "finding_ids": ids if i == 0 else []}
                      for i, group in enumerate(check["groups"])]
            rows.append({"check_id": check["id"], "title": check["title"],
                         "status": "发现风险" if ids else "需要法务判断", "reason": "合成测试。",
                         "finding_ids": ids, "groups": groups})
        checks = [{"rule_id": f"G{i:02d}", "title": finding["title"], "status": "发现风险", "reason": SYNTHETIC_NOTICE}
                  for i, finding in enumerate(findings, 1)]
        return json.dumps({"findings": findings, "checks": checks, "options": [], "uncovered": [],
                           "limitations": [SYNTHETIC_NOTICE], "checklist_results": rows}, ensure_ascii=False)
    return complete
