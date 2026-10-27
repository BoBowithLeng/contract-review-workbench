import io
import unittest

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from contract_core import parse_docx
from document_view import docx_headings, docx_navigation, index_document, reader_markup


def element(tag, **attrs):
    node = OxmlElement('w:' + tag)
    for name, value in attrs.items():
        node.set(qn('w:' + name), str(value))
    return node


def list_definition(doc, formats=('decimalHalfWidth',) * 3, starts=None, restarts=None):
    root = doc.part.numbering_part.element
    abstract = element('abstractNum', abstractNumId=99)
    for depth, fmt in enumerate(formats):
        level = element('lvl', ilvl=depth)
        level.append(element('start', val=(starts or {}).get(depth, 1)))
        level.append(element('numFmt', val=fmt))
        level.append(element('lvlText', val='.'.join('%' + str(i + 1) for i in range(depth + 1)) + '.'))
        if depth in (restarts or {}):
            level.append(element('lvlRestart', val=restarts[depth]))
        abstract.append(level)
    root.append(abstract)
    num = element('num', numId=100)
    num.append(element('abstractNumId', val=99))
    root.append(num)
    return abstract, num


def number(paragraph, depth=0, num_id=100):
    ppr = paragraph._p.get_or_add_pPr()
    prop = element('numPr')
    prop.append(element('ilvl', val=depth))
    prop.append(element('numId', val=num_id))
    ppr.append(prop)
    return paragraph


def saved(doc):
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def indexed(doc):
    data = saved(doc)
    contract = parse_docx(data).model_dump()
    original = parse_docx(data).model_dump()
    contract.update(docx_navigation(data))
    result = index_document(contract)
    assert contract['pieces'] == original['pieces']
    assert contract['fingerprint'] == original['fingerprint']
    assert [(p['id'], p['text']) for p in result] == [(p['id'], p['text']) for p in original['pieces']]
    return result, contract


class WordNumberingTests(unittest.TestCase):
    def test_numbered_headings_nested_clauses_and_attachment(self):
        doc = Document()
        list_definition(doc)
        doc.add_paragraph('开发合同', style='Heading 1')
        number(doc.add_paragraph('开发内容', style='Heading 3'))
        number(doc.add_paragraph('甲方委托乙方开发软件。'), 1)
        number(doc.add_paragraph('交付源代码。'), 2)
        number(doc.add_paragraph('乙方还应交付说明书。'), 1)
        number(doc.add_paragraph('费用支付', style='Heading 3'))
        number(doc.add_paragraph('先付款。'), 1)
        doc.add_paragraph('附件一：项目确认书', style='Heading 2')
        doc.add_paragraph('附件内容')
        paragraphs, contract = indexed(doc)
        self.assertEqual([p['numbering_label'] for p in paragraphs],
                         ['', '1.', '1.1.', '1.1.1.', '1.2.', '2.', '2.1.', '', ''])
        self.assertEqual([p['level'] for p in paragraphs], [1, 3, 0, 0, 0, 3, 0, 2, 0])
        self.assertEqual(paragraphs[1]['heading_parent'], 'P000')
        self.assertEqual(paragraphs[5]['heading_parent'], 'P000')
        self.assertEqual(paragraphs[7]['heading_parent'], 'P000')
        self.assertEqual(paragraphs[3]['location'], '开发合同 / 1. 开发内容 / 1.1. / 1.1.1.')
        self.assertEqual(paragraphs[8]['location'], '开发合同 / 附件一：项目确认书')
        rendered = reader_markup(contract, [{'id': 'F1', 'refs': [{'id': 'P002', 'quote': '委托乙方'}]}], 'F1')
        self.assertIn('1.1. 甲方<mark>委托乙方</mark>开发软件。', rendered['html'])
        self.assertEqual([s['text'] for s in rendered['sections']],
                         ['开发合同', '1. 开发内容', '2. 费用支付', '附件一：项目确认书'])
        self.assertEqual(docx_navigation(saved(doc))['headings'], docx_headings(saved(doc)))

    def test_start_override_independent_instance_and_restart(self):
        doc = Document()
        _, num = list_definition(doc, starts={1: 3})
        override = element('lvlOverride', ilvl=0)
        override.append(element('startOverride', val=5))
        num.append(override)
        replacement = element('lvlOverride', ilvl=1)
        level = element('lvl', ilvl=1)
        level.append(element('start', val=2))
        level.append(element('numFmt', val='lowerLetter'))
        level.append(element('lvlText', val='%1.(%2)'))
        replacement.append(level)
        num.append(replacement)
        num2 = element('num', numId=101)
        num2.append(element('abstractNumId', val=99))
        doc.part.numbering_part.element.append(num2)
        for depth in (0, 1, 1, 0, 1):
            number(doc.add_paragraph('内容'), depth)
        number(doc.add_paragraph('另一列表'), num_id=101)
        paragraphs, _ = indexed(doc)
        self.assertEqual([p['numbering_label'] for p in paragraphs],
                         ['5.', '5.(b)', '5.(c)', '6.', '6.(b)', '1.'])

    def test_style_inheritance_partial_direct_numpr_and_cancellation(self):
        doc = Document()
        list_definition(doc)
        parent = doc.styles.add_style('NumberedBase', WD_STYLE_TYPE.PARAGRAPH)
        prop = element('numPr')
        prop.append(element('numId', val=100))
        prop.append(element('ilvl', val=0))
        parent.element.get_or_add_pPr().append(prop)
        child = doc.styles.add_style('NumberedChild', WD_STYLE_TYPE.PARAGRAPH)
        child.base_style = parent
        doc.add_paragraph('父级', style=child)
        paragraph = doc.add_paragraph('子级', style=child)
        prop = element('numPr')
        prop.append(element('ilvl', val=1))
        paragraph._p.get_or_add_pPr().append(prop)
        number(doc.add_paragraph('取消编号', style=child), num_id=0)
        doc.add_paragraph('下一父级', style=child)
        paragraphs, _ = indexed(doc)
        self.assertEqual([p['numbering_label'] for p in paragraphs], ['1.', '1.1.', '', '2.'])

    def test_default_style_numbering_and_omitted_start(self):
        doc = Document()
        abstract, _ = list_definition(doc, formats=('decimal',))
        level = abstract.find(qn('w:lvl'))
        level.remove(level.find(qn('w:start')))
        prop = element('numPr')
        prop.append(element('numId', val=100))
        doc.styles['Normal'].element.get_or_add_pPr().append(prop)
        doc.add_paragraph('默认样式第一项')
        doc.add_paragraph('默认样式第二项')
        self.assertEqual([p['numbering_label'] for p in indexed(doc)[0]], ['0.', '1.'])

    def test_explicit_outline_values_override_heading_names(self):
        doc = Document()
        title = doc.add_paragraph('不在导航中的标题', style='Heading 1')
        title.style.element.pPr.find(qn('w:outlineLvl')).set(qn('w:val'), '9')
        root = doc.add_paragraph('附件', style='Heading 2')
        root.style.element.pPr.find(qn('w:outlineLvl')).set(qn('w:val'), '0')
        child = doc.add_paragraph('附件子标题', style='Heading 3')
        child.style.element.pPr.find(qn('w:outlineLvl')).set(qn('w:val'), '1')
        paragraphs, _ = indexed(doc)
        self.assertEqual([p['level'] for p in paragraphs], [0, 1, 2])
        self.assertEqual(paragraphs[2]['heading_parent'], 'P001')

    def test_never_restart_and_explicit_ancestor_restart(self):
        doc = Document()
        list_definition(doc, restarts={1: 0, 2: 1})
        for depth in (0, 1, 2, 1, 2, 0, 1, 2):
            number(doc.add_paragraph('内容'), depth)
        paragraphs, _ = indexed(doc)
        self.assertEqual([p['numbering_label'] for p in paragraphs],
                         ['1.', '1.1.', '1.1.1.', '1.2.', '1.2.2.', '2.', '2.3.', '2.3.1.'])

    def test_table_and_empty_paragraph_count_in_source_order(self):
        doc = Document()
        list_definition(doc)
        number(doc.add_paragraph('开始'))
        table = doc.add_table(rows=1, cols=1)
        number(table.cell(0, 0).paragraphs[0])  # Empty but actually numbered.
        number(table.cell(0, 0).add_paragraph('表格中的条款'))
        number(doc.add_paragraph('最后'))
        paragraphs, _ = indexed(doc)
        self.assertEqual([p['id'] for p in paragraphs], ['P000', 'T0R0C0P0', 'T0R0C0P1', 'P001'])
        self.assertEqual([p['numbering_label'] for p in paragraphs], ['1.', '2.', '3.', '4.'])

    def test_unsupported_and_missing_ancestors_do_not_fabricate_labels(self):
        doc = Document()
        list_definition(doc, formats=('arabicAbjad', 'decimal'))
        number(doc.add_paragraph('未知格式'))
        number(doc.add_paragraph('祖先格式未知'), 1)
        number(doc.add_paragraph('未知编号实例'), num_id=999)
        paragraphs, _ = indexed(doc)
        self.assertEqual([p['numbering_label'] for p in paragraphs], ['', '', ''])
        self.assertEqual([p['level'] for p in paragraphs], [0, 0, 0])
        doc = Document()
        list_definition(doc)
        number(doc.add_paragraph('缺少父级'), 1)
        self.assertEqual(indexed(doc)[0][0]['numbering_label'], '')

    def test_common_formats_and_legal_decimal(self):
        for fmt, label in (('upperRoman', 'IV.'), ('lowerLetter', 'd.'),
                           ('chineseCounting', '四.'), ('decimalZero', '04.'),
                           ('decimalFullWidth', '４.'), ('ideographDigital', '四.')):
            with self.subTest(fmt=fmt):
                doc = Document()
                list_definition(doc, formats=(fmt,), starts={0: 4})
                number(doc.add_paragraph('条款'))
                self.assertEqual(indexed(doc)[0][0]['numbering_label'], label)
        doc = Document()
        abstract, _ = list_definition(doc, formats=('upperRoman', 'lowerLetter'))
        abstract.findall(qn('w:lvl'))[1].append(element('isLgl'))
        number(doc.add_paragraph('父级'))
        number(doc.add_paragraph('法律编号'), 1)
        self.assertEqual(indexed(doc)[0][1]['numbering_label'], '1.1.')

    def test_unsupported_extended_restart_and_letter_range_are_unlabelled(self):
        doc = Document()
        abstract, _ = list_definition(doc)
        abstract.set('{http://schemas.microsoft.com/office/word/2012/wordml}restartNumberingAfterBreak', '1')
        number(doc.add_paragraph('跨节重启定义'))
        self.assertEqual(indexed(doc)[0][0]['numbering_label'], '')
        doc = Document()
        list_definition(doc, formats=('upperLetter',), starts={0: 27})
        number(doc.add_paragraph('超出已支持字母范围'))
        self.assertEqual(indexed(doc)[0][0]['numbering_label'], '')

    def test_linked_numbering_style_and_picture_bullet_suppression(self):
        doc = Document()
        list_definition(doc)
        style = doc.styles.add_style('LinkedNumbering', WD_STYLE_TYPE.LIST)
        prop = element('numPr')
        prop.append(element('numId', val=100))
        style.element.append(element('pPr'))
        style.element.find(qn('w:pPr')).append(prop)
        abstract = element('abstractNum', abstractNumId=98)
        abstract.append(element('numStyleLink', val=style.style_id))
        doc.part.numbering_part.element.append(abstract)
        num = element('num', numId=102)
        num.append(element('abstractNumId', val=98))
        doc.part.numbering_part.element.append(num)
        number(doc.add_paragraph('链接列表'), num_id=102)
        number(doc.add_paragraph('链接列表下一项'), num_id=102)
        self.assertEqual([p['numbering_label'] for p in indexed(doc)[0]], ['1.', '2.'])
        root = doc.part.numbering_part.element
        level = root.findall(qn('w:abstractNum'))[-2].find(qn('w:lvl'))
        level.append(element('lvlPicBulletId', val=1))
        self.assertEqual([p['numbering_label'] for p in indexed(doc)[0]], ['', ''])


if __name__ == '__main__':
    unittest.main()
