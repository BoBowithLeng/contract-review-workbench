from __future__ import annotations

import copy
import hashlib
import io
import json
import tempfile
import unittest
from pathlib import Path

from docx import Document

from contract_core import parse_docx
from delivery_bundle import AUDIT_REPORT_FILENAME, FILENAMES, build_bundle, write_bundle
from delivery_demo import demo_plan, source_document


class DeliveryBundleTests(unittest.TestCase):
    def setUp(self):
        self.data = source_document()
        self.plan = demo_plan(self.data)

    def test_all_outputs_share_one_source_and_confirmation_snapshot(self):
        before = copy.deepcopy(self.plan)
        bundle = build_bundle(self.data, self.plan)
        self.assertEqual(self.plan, before)
        self.assertEqual(set(bundle["files"]), (set(FILENAMES.values()) - {FILENAMES["report"]}) | {AUDIT_REPORT_FILENAME})
        manifest = bundle["manifest"]
        self.assertEqual(manifest["source_sha256"], hashlib.sha256(self.data).hexdigest())
        self.assertEqual(manifest["snapshot_revision"], 1)
        self.assertFalse(manifest["model_called_for_export"])
        self.assertFalse(manifest["legal_review_included"])
        for record in manifest["files"]:
            self.assertEqual(record["sha256"], hashlib.sha256(bundle["files"][record["name"]]).hexdigest())
        report = Document(io.BytesIO(bundle["files"][AUDIT_REPORT_FILENAME]))
        text = "\n".join(p.text for p in report.paragraphs)
        self.assertIn("没有运行法律审核", text)
        self.assertIn(manifest["source_sha256"], text)
        self.assertIn(manifest["plan_sha256"], text)
        self.assertIn("陈示例", text)
        self.assertIn("请业务部门核对实际联系人", text)

    def test_output_selection_and_new_directory_only(self):
        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp) / "delivery"
            result = write_bundle(self.data, self.plan, path, variants=["clean", "report"])
            self.assertEqual({r["variant"] for r in result["files"]}, {"clean", "report"})
            self.assertEqual((path / "00-原稿.docx").read_bytes(), self.data)
            before = {p.name: p.read_bytes() for p in path.iterdir()}
            with self.assertRaises(FileExistsError):
                write_bundle(self.data, self.plan, path, variants=["clean"])
            self.assertEqual(before, {p.name: p.read_bytes() for p in path.iterdir()})
            manifest = json.loads((path / "manifest.json").read_text(encoding="utf-8"))
            self.assertEqual(result, manifest)

    def test_invalid_plan_creates_no_output_directory(self):
        plan = copy.deepcopy(self.plan)
        plan["edits"][0]["confirmed"] = False
        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp) / "not-created"
            with self.assertRaises(ValueError):
                write_bundle(self.data, plan, path)
            self.assertFalse(path.exists())
        for variants in ([], ["clean", "clean"], ["pdf"]):
            with self.assertRaises(ValueError):
                build_bundle(self.data, self.plan, variants=variants)

    def test_unrelated_review_cannot_be_attached(self):
        with self.assertRaisesRegex(ValueError, "合同文字"):
            build_bundle(self.data, self.plan, review_bundle={"contract": {"fingerprint": "wrong"}})

    def test_same_text_review_can_include_the_delivery_snapshot(self):
        contract = parse_docx(self.data).model_dump()
        review = {
            "contract": contract, "sources": [],
            "report": {"fingerprint": contract["fingerprint"], "mode": "synthetic_test",
                       "context": {"region": "合成测试", "review_date": "2026-09-20"},
                       "findings": [], "checks": [], "options": [], "uncovered": ["测试夹具"]},
        }
        bundle = build_bundle(self.data, self.plan, variants=["report"], review_bundle=review)
        self.assertTrue(bundle["manifest"]["legal_review_included"])
        text = "\n".join(p.text for p in Document(io.BytesIO(bundle["files"][FILENAMES["report"]])).paragraphs)
        self.assertIn("审查报告", text)
        self.assertIn("陈示例", text)
        self.assertIn("请业务部门核对实际联系人", text)
        self.assertNotIn(bundle["manifest"]["plan_sha256"], text)
        self.assertNotIn(bundle["manifest"]["source_sha256"], text)
        self.assertNotIn("SHA-256", text)


if __name__ == "__main__":
    unittest.main()
