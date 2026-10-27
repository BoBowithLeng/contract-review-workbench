import io
import unittest
import zipfile

from openpyxl import Workbook, load_workbook

from checklist_excel import make_template, parse_workbook


def workbook_bytes(rows, headers=None):
    wb = Workbook()
    ws = wb.active
    ws.title = "清单"
    ws.append(headers or ["检查事项", "审核要求", "什么时候检查", "可能需要的信息", "处理建议", "备注"])
    for row in rows:
        ws.append(row)
    other = wb.create_sheet("说明与示例")
    other.append(["不是清单项", "不导入"])
    output = io.BytesIO()
    wb.save(output)
    wb.close()
    return output.getvalue()


class ChecklistExcelTests(unittest.TestCase):
    def test_template_real_xlsx_blank_import_sheet(self):
        data = make_template()
        wb = load_workbook(io.BytesIO(data))
        self.assertEqual(wb.sheetnames, ["清单", "说明与示例"])
        self.assertEqual(wb["清单"].max_row, 1)
        self.assertEqual([c.value for c in wb["清单"][1]], ["检查事项", "检查内容", "适用情形", "备注"])
        self.assertEqual(wb["清单"].auto_filter.ref, "A1:D1")
        wb.close()
        self.assertEqual(parse_workbook(data), {"ok": True, "rows": [], "errors": [], "warnings": []})

    def test_only_checklist_sheet_empty_optional_and_physical_row(self):
        result = parse_workbook(workbook_bytes([[], ["事项", "要求", None, "信息", None, 123]]))
        self.assertTrue(result["ok"])
        self.assertEqual(len(result["rows"]), 1)
        row = result["rows"][0]
        self.assertEqual(row["row_number"], 3)
        self.assertEqual(row["when"], "")
        self.assertEqual(row["notes"], "123")

    def test_row_errors_preview_retained_and_formula_type_only(self):
        data = workbook_bytes([["事项", "=1+2"], ["", "要求"], ["+普通文字", "-这也是正常要求"]])
        result = parse_workbook(data)
        self.assertTrue(result["ok"])
        self.assertEqual(len(result["rows"]), 3)
        self.assertIn("formula", [e["code"] for e in result["errors"]])
        self.assertIn("required", [e["code"] for e in result["errors"]])
        self.assertEqual(result["rows"][2]["title"], "+普通文字")
        wb = load_workbook(io.BytesIO(workbook_bytes([["事项", "文字"]])))
        wb["清单"]["B2"] = "=普通文字"
        wb["清单"]["B2"].data_type = "s"
        output = io.BytesIO()
        wb.save(output)
        wb.close()
        self.assertEqual(parse_workbook(output.getvalue())["errors"], [])

    def test_legacy_columns_warning_and_residual_contents(self):
        headers = ["检查事项", "核查要求", "适用条件", "判定要点", "通用目的", "类型细化与覆盖边界", "参考依据"]
        result = parse_workbook(workbook_bytes([["事项", "要求", "条件", "建议", "目的", "细化", "A PDF 12"]], headers))
        row = result["rows"][0]
        self.assertEqual(row["requirement"], "要求\n\n### 原可能需要的信息\n\n条件\n\n### 原处理建议\n\n建议")
        self.assertEqual(row["when"], "")
        self.assertEqual(row["facts"], "")
        self.assertEqual(row["guidance"], "")
        self.assertNotIn("sources", row)
        self.assertNotIn("PDF", str(row))
        self.assertEqual([w["value"] for w in result["warnings"] if "value" in w], ["目的", "细化", "A PDF 12"])

    def test_new_columns_markdown_content_and_no_legacy_fields(self):
        content = "### 判断范围\n- 完整条件。\n\n### 例外\n- 没有该场景时不强加要求。"
        result = parse_workbook(workbook_bytes([["事项", content, "涉及代采时", "说明"]],
                                              ["检查事项", "检查内容", "适用情形", "备注"]))
        self.assertEqual(result["errors"], [])
        self.assertEqual(result["warnings"], [])
        self.assertEqual(result["rows"][0]["requirement"], content)
        self.assertEqual(result["rows"][0]["when"], "涉及代采时")
        self.assertEqual(result["rows"][0]["notes"], "说明")
        self.assertEqual(result["rows"][0]["facts"], "")
        self.assertEqual(result["rows"][0]["guidance"], "")

    def test_old_six_columns_preserve_all_useful_text_without_deduplication(self):
        result = parse_workbook(workbook_bytes([["事项", "共同文字", "适用条件", "原事实\n第二行", "共同文字", "备注"]]))
        row = result["rows"][0]
        self.assertEqual(row["requirement"],
                         "共同文字\n\n### 原可能需要的信息\n\n原事实\n第二行\n\n### 原处理建议\n\n共同文字")
        self.assertEqual(row["when"], "适用条件")
        self.assertEqual(row["notes"], "备注")
        self.assertEqual(row["facts"], row["guidance"], "")
        self.assertEqual(result["errors"], [])
        self.assertTrue(any("未做语义去重" in w["message"] for w in result["warnings"]))

    def test_legacy_merge_over_limit_retains_rejected_preview_content(self):
        result = parse_workbook(workbook_bytes([["事项", "a" * 7990, "", "b" * 100, "原建议", "备注"]]))
        self.assertIn("legacy_merge", [e["code"] for e in result["errors"]])
        row = result["rows"][0]
        self.assertEqual(row["requirement"], "a" * 7990)
        self.assertEqual(row["facts"], "b" * 100)
        self.assertEqual(row["guidance"], "原建议")

    def test_formula_in_old_fields_is_rejected_not_migrated(self):
        result = parse_workbook(workbook_bytes([["事项", "要求", "", "=1+2", "原建议"]]))
        self.assertIn("formula", [e["code"] for e in result["errors"]])
        self.assertNotIn("=1+2", result["rows"][0]["requirement"])
        self.assertIn("原建议", result["rows"][0]["requirement"])

    def test_external_legacy_condition_remains_trigger_with_confirmation(self):
        result = parse_workbook(workbook_bytes([["事项", "要求", "涉及代采时"]], ["检查事项", "核查要求", "适用条件"]))
        self.assertEqual(result["rows"][0]["when"], "涉及代采时")
        self.assertEqual(result["rows"][0]["facts"], "")
        self.assertTrue(any(w["field"] == "when" and "确认" in w["message"] for w in result["warnings"]))

    def test_bad_package_missing_sheet_and_bad_headers(self):
        for data in (b"", b"broken", b"PKnotzip"):
            self.assertFalse(parse_workbook(data)["ok"])
        for headers in (["检查事项"], ["检查事项", "审核要求", "核查要求"], ["检查事项", "审核要求", "审核要求"], ["=1", "审核要求"]):
            self.assertFalse(parse_workbook(workbook_bytes([], headers))["ok"])
        wb = Workbook()
        output = io.BytesIO()
        wb.save(output)
        self.assertFalse(parse_workbook(output.getvalue())["ok"])

    def test_quotas(self):
        self.assertFalse(parse_workbook(b"a" * (2 * 1024 * 1024 + 1))["ok"])
        self.assertFalse(parse_workbook(workbook_bytes([["事项", "要求"]] * 201))["ok"])
        result = parse_workbook(workbook_bytes([["事项", "a" * 8001]]))
        self.assertEqual(result["errors"][0]["code"], "too_long")
        output = io.BytesIO()
        with zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as z:
            z.writestr("huge", b"a" * (20 * 1024 * 1024 + 1))
        result = parse_workbook(output.getvalue())
        self.assertFalse(result["ok"])
        self.assertEqual(result["errors"][0]["code"], "limit")

    def test_unnamed_and_unknown_columns_not_silently_lost(self):
        result = parse_workbook(workbook_bytes([["事项", "要求", "多余", "残余"]], ["检查事项", "审核要求", None, "未知"] ))
        self.assertEqual(len([w for w in result["warnings"] if w["code"] == "unmapped_content"]), 2)


if __name__ == "__main__":
    unittest.main()
