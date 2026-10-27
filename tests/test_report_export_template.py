"""The agreed report hierarchy, source positions and business-content contract."""
import copy
import io
import unittest

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from contract_core import parse_docx, parse_text
from document_view import docx_navigation, index_document
from report_export import export_docx
from report_presentation import source_locations
from tests.test_word_numbering import list_definition, number, saved


def report(findings=None, **extra):
    return {'context': {'contract_type': 'general', 'contract_type_label': '软件委托开发合同',
                        'our_party': '甲方', 'our_role': '委托方', 'review_date': '2026-10-06'},
            'findings': findings or [], 'options': [], 'checks': [], **extra}


def finding(fid='F01', **extra):
    return {'id': fid, 'rule_id': 'G01', 'title': '验收衔接不清', 'risk': '中',
            'status': '发现风险', 'reason': '验收阶段不明确。', 'suggestion': '写明各阶段触发条件。',
            'refs': [], **extra}


def read(output):
    document = Document(io.BytesIO(output))
    paragraphs = '\n'.join(p.text for p in document.paragraphs)
    all_text = paragraphs + '\n' + '\n'.join(c.text for t in document.tables for row in t.rows for c in row.cells)
    return document, paragraphs, all_text


def acceptance_source():
    """Mirror the original's titled 5.1.1 and unnumbered 5.1.2 continuation."""
    doc = Document()
    list_definition(doc, starts={0: 5})
    number(doc.add_paragraph('验收', style='Heading 2'), 0)
    number(doc.add_paragraph('验收'), 1)
    number(doc.add_paragraph('上线前验收'), 2)
    doc.add_paragraph('正式上线前，甲乙双方应进行上线前验收。')
    number(doc.add_paragraph('上线后30天为试运行期间。'), 2)
    doc.add_paragraph('甲乙双方应根据试运行期间情况进行最终验收。')
    number(doc.add_paragraph('验收标准'), 1)
    doc.add_paragraph('验收标准详见附件。')
    number(doc.add_paragraph('验收期限'), 1)
    number(doc.add_paragraph('通知验收后的5个工作日内完成验收。'), 2)
    number(doc.add_paragraph('异议以书面形式通知。'), 2)
    number(doc.add_paragraph('未签署亦未提出异议的，视为验收合格。'), 2)
    return saved(doc)


class ReportTemplateTests(unittest.TestCase):
    def test_source_numbering_recovered_without_mutating_contract(self):
        data = acceptance_source()
        contract = parse_docx(data).model_dump()
        original = copy.deepcopy(contract)
        refs = [{'id': pid, 'quote': next(p['text'] for p in contract['pieces'] if p['id'] == pid)}
                for pid in ('P003', 'P004', 'P005', 'P009', 'P011')]
        plan = {'status': 'ready', 'edits': [{'piece_id': ref['id'], 'new_text': '修改后的' + ref['quote']} for ref in refs],
                'insertions': [], 'paragraph_actions': [], 'questions': [], 'comment': '与P009衔接。'}
        item = finding(refs=refs, draft_plan=plan, clause_draft='不应采用拼接旧稿', reason='P003、P004、P005及P009需与F01一起处理。')
        result = report([item])
        original_report = copy.deepcopy(result)
        _, text, all_text = read(export_docx(result, [], contract, source_docx=data))
        for label in ['第5.1.1条「上线前验收」', '第5.1.2条', '第5.1.2条续段', '第5.3.1条', '第5.3.3条']:
            self.assertGreaterEqual(text.count(label), 2, label)
        self.assertNotIn('第5.1.3条', text)
        self.assertNotIn('不应采用拼接旧稿', text)
        self.assertNotRegex(all_text, r'\b[PFG]\d{2,}\b')
        self.assertNotIn('尚未采用', text)
        self.assertEqual(contract, original)
        self.assertEqual(result, original_report)

    def test_typography_hierarchy_table_and_no_theme_font_override(self):
        contract = parse_text('软件委托开发合同\n合同正文。').model_dump()
        document, text, _ = read(export_docx(report([finding(refs=[{'id': 'P001', 'quote': '合同正文。'}], clause_draft='修改正文。')]), [], contract))
        self.assertIn('一、审查意见概览', text)
        self.assertIn('二、详细审查意见', text)
        self.assertIn('（一）验收衔接不清', text)
        for number, label in enumerate(['合同原文', '为什么需要修改', '建议怎么改', '建议修改文字'], 1):
            self.assertIn(f'{number}. {label}', text)
        self.assertEqual(document.paragraphs[0].style.name, 'Title')
        for name, font, size, bold in [('Title', '黑体', 16, False), ('Heading 1', '黑体', 14, False),
                                       ('Heading 2', '宋体', 12, True), ('Heading 3', '宋体', 12, True), ('Normal', '宋体', 12, False)]:
            style = document.styles[name]
            self.assertEqual((style.font.name, style.font.size.pt, style.font.bold), (font, size, bold))
            self.assertEqual(style.element.rPr.rFonts.get(qn('w:eastAsia')), font)
        self.assertEqual(document.styles['Normal'].paragraph_format.line_spacing, 1.5)
        self.assertEqual(document.styles['Normal'].paragraph_format.first_line_indent.pt, 24)
        for node in document.styles.element.iter(qn('w:rFonts')):
            self.assertFalse(any(key.lower().endswith('theme') for key in node.attrib))
        for name in ['Title', 'Heading 1', 'Heading 2', 'Heading 3']:
            self.assertEqual(list(document.styles[name].element.iter(qn('w:pBdr'))), [])
        section = document.sections[0]
        self.assertAlmostEqual(section.footer_distance.cm, 0.76, places=2)
        footer = section.footer.paragraphs[0]
        self.assertEqual(footer.paragraph_format.first_line_indent.pt, 0)
        self.assertIsNotNone(footer._p.find(qn('w:fldSimple')))
        self.assertEqual(footer._p.find(qn('w:fldSimple')).get(qn('w:instr')), 'PAGE')
        self.assertTrue(all(run.font.size.pt == 9 for run in footer.runs))
        self.assertTrue(document.tables[0].rows[0]._tr.trPr.find(qn('w:tblHeader')) is not None)
        self.assertEqual([c.text for c in document.tables[0].rows[0].cells], ['序号', '需要关注的问题', '风险', '处理情况'])
        self.assertEqual(document.styles['Normal'].paragraph_format.space_after.pt, 3)
        self.assertEqual(document.styles['Heading 3'].paragraph_format.space_before.pt, 5)
        self.assertEqual(document.styles['Heading 3'].paragraph_format.space_after.pt, 3)
        self.assertFalse(any(p.paragraph_format.page_break_before for p in document.paragraphs))

    def test_old_data_does_not_invent_clause_numbers(self):
        contract = parse_text('合同正文\n第二段正文').model_dump()
        _, text, _ = read(export_docx(report([finding(refs=[{'id': 'P001', 'quote': '第二段正文'}], clause_draft='旧数据候选改稿')]), [], contract))
        self.assertIn('正文第2段：“第二段正文”', text)
        self.assertIn('旧数据候选改稿', text)
        self.assertNotIn('第2条', text)
        self.assertNotIn('可选择的处理办法', text)
        self.assertNotIn('本次合同的补充信息', text)
        self.assertNotIn('已确认的修改与批注', text)

    def test_heading_without_number_and_attachment_cell_have_honest_positions(self):
        doc = Document()
        doc.add_paragraph('软件委托开发合同')
        doc.add_paragraph('附件 项目说明', style='Heading 2')
        doc.add_paragraph('附件内容。')
        doc.add_table(rows=1, cols=1).cell(0, 0).text = '表格内的交付要求'
        data = saved(doc)
        contract = parse_docx(data).model_dump()
        contract.update(docx_navigation(data))
        positions = source_locations(index_document(contract))
        self.assertEqual(positions['P002'], '「附件 项目说明」下第1段')
        self.assertIn('表格1第1行第1列第1段', positions['T0R0C0P0'])
        _, text, _ = read(export_docx(report([finding(refs=[{'id': 'T0R0C0P0', 'quote': '表格内的交付要求'}])]), [], contract, source_docx=data))
        self.assertIn('表格1第1行第1列第1段', text)
        self.assertNotIn('T0R0C0P0', text)

    def test_checklist_is_appendix_and_preserves_each_group_status_and_reason(self):
        contract = parse_text('正文').model_dump()
        result = report([finding()], checklist_results=[{'check_id': 'C01', 'title': '自动续约', 'status': '处理失败',
                          'reason': '此检查未完成。', 'finding_ids': ['F01'],
                          'groups': [{'group_id': 'C01:G01', 'title': '通知时限', 'status': '信息不足',
                                      'reason': '需确认P000通知期限。', 'finding_ids': ['F01']}]}])
        _, text, _ = read(export_docx(result, [], contract))
        self.assertLess(text.index('风险等级：'), text.index('附录 清单逐项检查结果'))
        for value in ['自动续约：尚未完成检查', '此检查未完成。', '通知时限：待确认', '通知期限。', '对应审查意见：意见一']:
            self.assertIn(value, text)
        self.assertNotRegex(text, r'\b(?:F01|P000|G01)\b')

    def test_checklist_statuses_use_the_same_words_as_the_reader(self):
        contract = parse_text('正文').model_dump()
        rows = [{'check_id': f'C0{i}', 'title': f'检查{i}', 'status': status, 'reason': '说明。', 'finding_ids': [],
                 'groups': [{'group_id': f'C0{i}:G01', 'title': f'小组{i}', 'status': status, 'reason': '说明。', 'finding_ids': []}]}
                for i, status in enumerate(['未检出', '发现风险', '需要法务判断', '信息不足', '不适用', '处理失败'], 1)]
        _, text, _ = read(export_docx(report([], checklist_results=rows), [], contract))
        for i, label in enumerate(['未发现问题', '发现问题', '待确认', '待确认', '不适用', '尚未完成检查'], 1):
            self.assertIn(f'检查{i}：{label}', text)
            self.assertIn(f'小组{i}：{label}', text)
        for raw in ['未检出', '发现风险', '需要法务判断', '信息不足', '处理失败']:
            self.assertNotIn('：' + raw, text)

    def test_general_review_has_no_placeholder_basics_appendix(self):
        contract = parse_text('正文').model_dump()
        placeholder = {'description': '通用初审；合同基本要素由逐项审查意见说明。', 'source_ids': []}
        _, text, _ = read(export_docx(report([finding()], basic_info=placeholder), [], contract))
        self.assertNotIn('基本要素检查', text)
        self.assertNotIn('合同基本要素由逐项审查意见说明', text)
        labor = report([finding()], basic_info={'description': '劳动合同必备条款检查。', 'source_ids': []},
                       completeness=[{'title': '工作地点', 'status': '已找到', 'note': '', 'refs': []}])
        labor['context'] = {**labor['context'], 'contract_type': 'labor'}
        _, text, _ = read(export_docx(labor, [], contract))
        self.assertIn('附录 基本要素检查', text)
        self.assertIn('工作地点：已找到', text)

    def test_confirmed_changes_read_as_whole_paragraph_before_and_after(self):
        contract = parse_text('合同\n1.2 详细需求由双方后续口头确认，变更另行商定。').model_dump()
        source = contract['pieces'][1]['text']
        first = source.index('后续口头')
        second = source.index('另行商定')
        plan = {'edits': [
            {'id': 'e1', 'piece_id': 'P001', 'start': first, 'end': first + 4, 'old_text': '后续口头',
             'new_text': '以需求附件书面', 'confirmed': True, 'finding_ids': ['F01']},
            {'id': 'e2', 'piece_id': 'P001', 'start': second, 'end': second + 4, 'old_text': '另行商定',
             'new_text': '经双方负责人书面确认', 'confirmed': True, 'finding_ids': ['F01']}], 'comments': []}
        _, text, _ = read(export_docx(report([finding()]), [], contract, confirmed_plan=plan))
        section = text[text.index('已确认的修改与批注'):]
        self.assertIn('修改前：“1.2 详细需求由双方后续口头确认，变更另行商定。”', section)
        self.assertIn('修改后：“1.2 详细需求由双方以需求附件书面确认，变更经双方负责人书面确认。”', section)
        self.assertEqual(section.count('（一）'), 1)
        self.assertNotIn('（二）', section)
        self.assertEqual(section.count('对应审查意见：意见一'), 1)
        # A range that no longer matches the paragraph keeps the exact recorded fragments.
        plan['edits'][1]['old_text'] = '不存在的文字'
        _, text, _ = read(export_docx(report([finding()]), [], contract, confirmed_plan=plan))
        section = text[text.index('已确认的修改与批注'):]
        self.assertNotIn('修改后：', section)
        self.assertIn('原文：后续口头', section)

    def test_suggestions_insertions_unresolved_facts_sources_and_actions_remain_distinct(self):
        contract = parse_text('第九条 数据处理\n保密内容。').model_dump()
        plan = {'status': 'needs_input', 'edits': [], 'insertions': [], 'paragraph_actions': [], 'questions': ['保存位置在哪里？'], 'comment': '待确认保存位置。'}
        item = finding(refs=[{'id': 'P001', 'quote': '保密内容。'}], draft_plan=plan, clause_draft='过时的拼接草稿', source_ids=['WEB001'])
        facts = {'facts': [{'text': '需返还数据。', 'active': True}, {'text': '已失效事实', 'active': False}],
                 'questions': [{'question': '由谁删除数据？', 'status': 'unknown'}]}
        source = {'id': 'WEB001', 'title': '原始法源名称 LCL30_50', 'verified': False, 'url': 'https://example.test/source'}
        document, text, all_text = read(export_docx(report([item], confirmed_facts=facts, limitations=['适用仍待核。']), [source], contract,
                        reader_mode=True, reader_actions={'F01': {'label': '不采用修改建议'}}))
        for value in ['风险等级：中风险', '不采用修改建议不表示原风险已解决',
                      '修改建议状态：待补充信息', '保存位置在哪里？', '由谁删除数据？', '需返还数据。',
                      '待核资料：原始法源名称 LCL30_50', '适用仍待核。']:
            self.assertIn(value, text)
        self.assertIn('https://example.test/source', [r.target_ref for r in document.part.rels.values() if r.is_external])
        self.assertEqual(document.tables[0].rows[1].cells[3].text, '不采用修改建议')
        self.assertNotIn('处理状态：', text)
        self.assertNotIn('已失效事实', text)
        self.assertNotIn('过时的拼接草稿', text)
        self.assertNotIn('4. 建议修改文字', text)
        item['draft_plan'] = {'status': 'ready', 'edits': [], 'insertions': [{'after_piece_id': 'P001', 'new_text': '合同结束后返还数据。'}], 'comment': '新增返还条款。'}
        _, text, _ = read(export_docx(report([item]), [], contract))
        self.assertIn('之后新增段落：合同结束后返还数据。', text)

    def test_confirmed_changes_and_linked_comments_have_no_technical_ids(self):
        contract = parse_text('第一段正文。\n第二段正文。').model_dump()
        plan = {'source_sha256': '秘密指纹', 'revision': 2, 'author': '确认人', 'edits': [
            {'id': 'edit-secret-1', 'piece_id': 'P000', 'old_text': '第一段正文。', 'new_text': '确认替换。', 'confirmed': True, 'finding_ids': ['F01']},
            {'id': 'edit-secret-2', 'piece_id': 'P001', 'old_text': '', 'new_text': '确认新增。', 'kind': 'insert_paragraph_after', 'confirmed': True}],
            'comments': [{'id': 'comment-secret-1', 'edit_id': 'edit-secret-1', 'text': '确认批注全文。', 'confirmed': True}]}
        _, text, _ = read(export_docx(report([finding()]), [], contract, confirmed_plan=plan))
        for value in ['已确认的修改与批注', '确认替换。', '之后新增段落：确认新增。', '确认批注全文。', '对应审查意见：意见一']:
            self.assertIn(value, text)
        for value in ['秘密指纹', 'edit-secret', 'comment-secret', 'P000', 'F01', 'SHA-256']:
            self.assertNotIn(value, text)
        self.assertGreater(text.index('已确认的修改与批注'), text.index('详细审查意见'))

    def test_empty_or_unconfirmed_plan_does_not_claim_confirmed_actions(self):
        contract = parse_text('原文').model_dump()
        for plan in [{'edits': [], 'comments': []}, {'edits': [{'new_text': '建议', 'old_text': '原文', 'confirmed': False}], 'comments': []}]:
            _, text, _ = read(export_docx(report([finding()]), [], contract, confirmed_plan=plan))
            self.assertNotIn('已确认的修改与批注', text)

    def test_attachment_number_restart_keeps_its_scope_and_body_table_is_not_attachment(self):
        doc = Document()
        list_definition(doc)
        number(doc.add_paragraph('开发内容', style='Heading 2'), 0)
        number(doc.add_paragraph('正文交付要求。'), 1)
        doc.add_table(rows=1, cols=1).cell(0, 0).text = '正文价格表'
        doc.add_paragraph('第三方与开发方组件清单', style='Heading 1')
        # Independent list instance really restarts, rather than an invented ID.
        num = OxmlElement('w:num')
        num.set(qn('w:numId'), '101')
        abstract = OxmlElement('w:abstractNumId')
        abstract.set(qn('w:val'), '99')
        num.append(abstract)
        doc.part.numbering_part.element.append(num)
        number(doc.add_paragraph('第三方软件'), 0, num_id=101)
        number(doc.add_paragraph('应当取得第三方许可。'), 1, num_id=101)
        data = saved(doc)
        contract = parse_docx(data).model_dump()
        contract.update(docx_navigation(data))
        positions = source_locations(index_document(contract))
        self.assertTrue(positions['P001'].startswith('第1.1条'))
        self.assertTrue(positions['P004'].startswith('「第三方与开发方组件清单」第1.1条'))
        self.assertNotIn('附件', positions['T0R0C0P0'])
        self.assertIn('开发内容', positions['T0R0C0P0'])
        _, narrative, _ = read(export_docx(report([finding(reason='参照T0系列中的约定。')]), [], contract))
        self.assertNotIn('T0系列', narrative)
        self.assertIn('表格1系列', narrative)

    def test_information_failure_status_is_preserved_in_risk_line_only(self):
        contract = parse_text('正文').model_dump()
        items = [finding(status='信息不足', risk=None, issue_kind='信息', evidence_status='资料不足'),
                 finding('F02', status='处理失败', risk=None)]
        _, text, _ = read(export_docx(report(items), [], contract, reader_mode=True))
        self.assertIn('风险等级：待判断；信息不足', text)
        self.assertIn('风险等级：待判断；处理失败', text)
        self.assertNotIn('检查状态：', text)
        self.assertNotIn('问题类型：', text)

    def test_unsupported_numbering_is_not_presented_as_preceding_article(self):
        doc = Document()
        list_definition(doc, formats=('decimal', 'arabicAbjad'))
        number(doc.add_paragraph('验收', style='Heading 2'), 0)
        number(doc.add_paragraph('未知条号对应的正文。'), 1)
        data = saved(doc)
        contract = parse_docx(data).model_dump()
        contract.update(docx_navigation(data))
        positions = source_locations(index_document(contract))
        self.assertIn('条号未识别', positions['P001'])
        self.assertNotIn('第1条', positions['P001'])
        self.assertNotIn('续段', positions['P001'])

    def test_appendix_omits_only_proven_duplicate_summaries(self):
        contract = parse_text('正文').model_dump()
        groups = [{'title': '保存位置', 'status': '信息不足', 'reason': '保存位置需要确认。'}]
        rows = [{'title': '数据处理', 'status': '信息不足', 'reason': '保存位置需要确认。', 'groups': groups},
                {'title': '独有结论', 'status': '发现风险', 'reason': '顶层独有的协商结论。', 'groups': groups}]
        _, text, _ = read(export_docx(report([finding()], checklist_results=rows), [], contract))
        self.assertEqual(text.count('保存位置需要确认。'), 2)
        self.assertIn('顶层独有的协商结论。', text)
        self.assertEqual(text.count('数据处理：待确认'), 1)

    def test_heading_scope_and_readable_checklist_references(self):
        doc = Document()
        doc.add_paragraph('附件：软件开发项目确认书', style='Heading 1')
        doc.add_paragraph('运行环境要求', style='Heading 2')
        doc.add_paragraph('正式环境由甲方提供。')
        data = saved(doc)
        contract = parse_docx(data).model_dump()
        item = finding(reason='结合C19一并核对。', refs=[{'id':'P002', 'quote':'正式环境由甲方提供。'}])
        checks = [{'check_id':'C19', 'title':'许可安排', 'status':'信息不足'}]
        document, text, _ = read(export_docx(report([item], checklist_results=checks), [], contract, source_docx=data))
        self.assertIn('「附件：软件开发项目确认书 / 运行环境要求」下第1段', text)
        self.assertIn('结合清单第一项一并核对', text)
        risk = next(p for p in document.paragraphs if p.text.startswith('风险等级：'))
        self.assertTrue(risk.paragraph_format.keep_with_next)
        self.assertEqual(document.tables[0].cell(0,0).paragraphs[0].runs[0].font.size.pt, 10.5)

    def test_option_reference_is_readable_without_rewriting_real_product_codes(self):
        contract = parse_text('正文').model_dump()
        options = [{'id':'O1','title':'参照既有改法','clause_draft':'以F01项下draft_plan的修改稿为准。'},
                   {'id':'O2','title':'设备交付','clause_draft':'F01设备须安装P001模块。'},
                   {'id':'O3','title':'新增条款','clause_draft':'见 F01 新增段落'}]
        _, text, _ = read(export_docx(report([finding()], options=options), [], contract))
        self.assertIn('以意见一中的建议修改文字为准', text)
        self.assertNotIn('draft_plan', text)
        self.assertIn('F01设备须安装P001模块。', text)
        self.assertIn('见 意见一 新增段落', text)

    def test_sources_keep_full_hyperlink_without_printing_long_query_or_markup(self):
        contract = parse_text('正文').model_dump()
        url = 'https://example.test/law.pdf?download=' + 'abc' * 100
        source = {'id':'WEB001','title':'<p>**公开法规**</p>', 'url':url,'verified':False}
        document, text, _ = read(export_docx(report([finding(source_ids=['WEB001'])]), [source], contract))
        self.assertIn('公开法规', text)
        self.assertNotIn('<p>', text)
        self.assertNotIn('?download=', text)
        self.assertIn(url, [rel.target_ref for rel in document.part.rels.values() if rel.is_external])
        self.assertTrue(any('查看原文' in ''.join(node.itertext()) for node in document.element.iter(qn('w:hyperlink'))))

    def test_exact_business_codes_in_quotes_and_drafts_are_not_internal_refs(self):
        contract = parse_text('交付F01设备和P001模块。').model_dump()
        item = finding(refs=[{'id': 'P000', 'quote': '交付F01设备和P001模块。'}],
                       clause_draft='应当交付F01设备和P001模块。')
        options = [{'id': 'O1', 'title': '交付设备', 'clause_draft': '双方确认使用F01设备。',
                    'relationship': 'O1可以单独使用。'}]
        _, text, _ = read(export_docx(report([item], options=options), [], contract))
        self.assertIn('“交付F01设备和P001模块。”', text)
        self.assertIn('应当交付F01设备和P001模块。', text)
        self.assertIn('建议条款：双方确认使用F01设备。', text)
        self.assertIn('（一）O1 交付设备', text)


if __name__ == '__main__':
    unittest.main()


class NarrativeLocationTests(unittest.TestCase):
    def test_range_inside_one_clause_names_the_clause_once(self):
        from report_presentation import ReportWriter
        writer = ReportWriter(Document(), {'P093': '第9.2条续段「保密」', 'P103': '第9.2条「保密」',
                                           'P104': '第10.1条「违约责任」'}, {}, {})
        self.assertEqual(writer.clean('P093—P103已新增'), '第9.2条已新增')
        self.assertEqual(writer.clean('P093至P104'), '第9.2条续段至第10.1条')
        self.assertEqual(writer.clean('按draft_plan修改P104：'), '按建议改稿修改第10.1条：')
        self.assertEqual(writer.clean('research_memo 与 WEB 来源'), '检索记录 与 WEB 来源')

    def test_prose_uses_short_positions_and_collapses_ranges(self):
        doc = Document()
        list_definition(doc)
        number(doc.add_paragraph('开发软件', style='Heading 2'), 0)
        number(doc.add_paragraph('软件名称：'), 1)
        doc.add_paragraph('用途：')
        doc.add_paragraph('功能：')
        number(doc.add_paragraph('交付范围以附件为准。'), 1)
        data = saved(doc)
        contract = parse_docx(data).model_dump()
        ids = [p['id'] for p in contract['pieces']]
        reason = f'{ids[2]}-{ids[3]}均为空白；{ids[1]}—{ids[4]}需一并填写；另见{ids[2]}。'
        item = finding(refs=[{'id': ids[2], 'quote': '用途：'}], reason=reason)
        _, text, _ = read(export_docx(report([item]), [], contract, source_docx=data))
        prose = next(line for line in text.splitlines() if '均为空白' in line)
        self.assertNotIn('「开发软件」', prose)
        self.assertNotRegex(prose, r'(第[\d.]+条续段)\s*[-—]\s*\1')
        self.assertIn('至', prose)
        self.assertNotRegex(prose, r'\bP\d+\b')
        # The quoted original keeps the full position with its clause title.
        self.assertIn('「开发软件」', text[text.index('1. 合同原文'):text.index('2. 为什么需要修改')])
