from __future__ import annotations

import copy
import io
import json
import tempfile
import time
import unittest
import uuid
from pathlib import Path
from unittest.mock import patch

from docx import Document

from clause_discussion import ClauseDiscussionManager
from contract_core import parse_docx
from docx_delivery import inspect_source
from opinion_revisions import OpinionRevisionManager
from reader_delivery import DeliveryManager
from reader_server import _default_state
from tests.test_clause_discussion import FakeDelivery, request_body, wait_done
from tests.test_opinion_revisions import full_loaded, body as revision_body, answer as revision_answer


class LinkedDiscussionContextTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        doc = Document()
        for text in ("4.1 甲方应在七日内验收。", "4.2 验收后付款。", "应删除的旧条款。", "保留末段。"):
            doc.add_paragraph(text)
        doc.add_table(rows=1, cols=1).cell(0, 0).text = "附件价格表。"
        stream = io.BytesIO()
        doc.save(stream)
        self.data = stream.getvalue()
        self.source = self.root / "original.docx"
        self.source.write_bytes(self.data)
        contract = parse_docx(self.data).model_dump()
        self.loaded = full_loaded()
        self.loaded["contract"] = contract
        self.loaded["report"]["fingerprint"] = contract["fingerprint"]
        self.loaded["paragraphs"] = [{"id": p["id"], "location": p["id"]} for p in contract["pieces"]]
        self.loaded["report"]["findings"][0]["refs"] = [{"id": "P000", "quote": contract["pieces"][0]["text"]}]
        self.loaded["report"]["findings"][1]["refs"] = [{"id": "P001", "quote": contract["pieces"][1]["text"]}]
        self.delivery = DeliveryManager(self.root / "delivery", {self.loaded["id"]: self.source})
        self.review = _default_state(self.loaded["report"])
        self.calls = []
        self.manager = ClauseDiscussionManager(self.root / "discussion", self.delivery,
            completion=lambda _, messages: self.calls.append(copy.deepcopy(messages)) or "讨论答复",
            review_state_provider=lambda _: copy.deepcopy(self.review))

    def edit(self, piece_id, new_text, *, finding_id="F02", kind="replace", start=0, end=None):
        piece = next(p for p in inspect_source(self.data)["pieces"] if p["id"] == piece_id)
        if kind == "insert_paragraph_after":
            start = end = len(piece["text"])
        if end is None:
            end = len(piece["text"])
        return {"id": f"edit-{finding_id}-{piece_id}-{kind}-{start}", "piece_id": piece_id,
                "paragraph_sha256": piece["paragraph_sha256"], "start": start, "end": end,
                "old_text": piece["text"][start:end], "new_text": new_text,
                "confirmed": True, "kind": kind, "finding_ids": [finding_id]}

    def save(self, edits):
        sha = inspect_source(self.data)["source_sha256"]
        current = self.delivery.describe(self.loaded)["state"]["revision"]
        return self.delivery.save(self.loaded, {"expected_revision": current, "source_sha256": sha,
            "author": "测试确认人", "edits": edits, "comments": []})

    def context(self):
        return self.manager._source_context(self.loaded, self.loaded["report"]["findings"][0])

    def test_entire_current_contract_matches_clean_word_and_keeps_original_anchors(self):
        edits = [self.edit("P000", "十日", start=9, end=11),
                 self.edit("P001", "4.2 付款已独立确认。"), self.edit("P002", ""),
                 self.edit("P001", "4.3 第一新增段。\r\n4.4 第二新增段。", kind="insert_paragraph_after"),
                 self.edit("P003", "独立换行\r\n含制表\t文本。")]
        self.save(edits)
        source = self.context()
        actual = parse_docx(self.delivery.rereview_snapshot(self.loaded)["data"])
        self.assertEqual([p["text"] for p in source["current_contract"]], [p.text for p in actual.pieces])
        self.assertEqual([p["id"] for p in source["original_contract"]], [p["id"] for p in self.loaded["contract"]["pieces"]])
        self.assertEqual(source["original_contract"][2]["text"], "应删除的旧条款。")
        self.assertEqual(source["relevant_paragraphs"][0]["current_text"], source["current_contract"][0]["text"])
        self.assertIn(edits[0]["id"], source["relevant_paragraphs"][0]["confirmed_edit_ids"])
        inserted = [p for p in source["current_contract"] if p["source_piece_id"] is None]
        self.assertEqual(len(inserted), 2)
        self.assertTrue(all(p["id"].startswith("inserted:") for p in inserted))
        self.assertTrue(set(p["id"] for p in source["current_contract"]).isdisjoint(p["id"] for p in source["original_contract"]))
        self.assertEqual(source["current_contract"][-1]["text"], "附件价格表。")

    def test_relevant_insertion_is_separate_not_appended_and_pending_drafts_not_applied(self):
        self.loaded["report"]["findings"][1]["clause_draft"] = "尚未应用的模型建议"
        self.review["reviews"]["F01"].update(decision="修改", note="只存建议", clause_draft="尚未应用的人工建议")
        self.save([self.edit("P000", "已确认新增条款。", kind="insert_paragraph_after")])
        source = self.context()
        relevant = source["relevant_paragraphs"][0]
        self.assertEqual(relevant["current_text"], self.loaded["contract"]["pieces"][0]["text"])
        self.assertEqual(relevant["inserted_paragraphs"][0]["text"], "已确认新增条款。")
        self.assertEqual(relevant["confirmed_edit_ids"], [])
        text = "\n".join(p["text"] for p in source["current_contract"])
        self.assertNotIn("尚未应用", text)
        row = source["report_opinions"][1]
        self.assertTrue(row["document_action"]["modified"])

    def test_all_current_opinions_human_review_and_stale_version_not_active(self):
        self.review["revision"] = 3
        self.review["reviews"]["F02"].update(decision="驳回", note="愿意承担付款风险", risk="高")
        row = self.context()["report_opinions"][1]
        self.assertEqual(row["human_review"]["note"], "愿意承担付款风险")
        self.assertEqual(row["current_opinion"]["status"], "发现风险")
        self.assertTrue(row["document_action"]["ignored"])
        self.loaded["opinion_revisions"] = {"F01": 0, "F02": 1}
        self.loaded["report"]["findings"][1].update(opinion_revision=1, reason="已接受的新意见判断")
        row = self.context()["report_opinions"][1]
        self.assertEqual(row["current_opinion"]["reason"], "已接受的新意见判断")
        self.assertFalse(row["human_review"]["applies_to_current_opinion"])
        self.assertEqual(row["human_review"]["decision"], "未复核")
        self.assertEqual(row["human_review"]["note"], "")
        self.assertFalse(row["document_action"]["ignored"])
        self.review["reviews"]["F02"]["opinion_revision"] = 1
        self.assertTrue(self.context()["report_opinions"][1]["human_review"]["applies_to_current_opinion"])

    def test_source_reads_never_render_or_write_and_snapshot_versions_match(self):
        # A pristine GET must not create even a source-binding snapshot.
        self.manager.get(self.loaded, "F01")
        self.assertFalse(self.delivery.root.exists())
        self.save([self.edit("P001", "4.2 当前确认文字。")])
        before = {p: p.read_bytes() for p in self.root.rglob("*") if p.is_file()}
        with patch("reader_delivery.build_bundle", side_effect=AssertionError("no render")):
            for _ in range(3):
                self.manager.get(self.loaded, "F01")
        after = {p: p.read_bytes() for p in self.root.rglob("*") if p.is_file()}
        self.assertEqual(before, after)
        source = self.context()
        self.assertEqual(source["current_state_versions"]["delivery_revision"], 1)
        self.assertTrue(source["report_opinions"][1]["document_action"]["modified"])

    def test_overlay_binding_check_does_not_build_linked_context(self):
        with patch.object(self.delivery, "action_context", side_effect=AssertionError("no linked projection")):
            self.manager.review_state_provider = lambda _: (_ for _ in ()).throw(AssertionError("no review read"))
            source = self.manager._source_context(self.loaded, self.loaded["report"]["findings"][0], binding_only=True)
        self.assertEqual(set(source), {"source_binding", "delivery_revision"})
        self.assertEqual(source["delivery_revision"], 0)
        self.assertFalse(self.delivery.root.exists())

    def test_global_active_facts_and_only_selected_session_history(self):
        self.manager.facts_provider = lambda _: {"revision": 4, "facts": [
            {"id": "fact-" + "a" * 32, "text": "已更正旧事实", "active": False},
            {"id": "fact-" + "b" * 32, "text": "全合同当前事实", "active": True}]}
        self.manager.start(self.loaded, "F02", request_body("其他意见会话秘密"))
        wait_done(self.manager, self.loaded, "F02")
        self.manager.start(self.loaded, "F01", request_body("上一会话秘密"))
        wait_done(self.manager, self.loaded)
        new, _ = self.manager.new_session(self.loaded, "F01", {"expected_sessions_revision": 0,
            "expected_delivery_revision": 0, "expected_opinion_revision": 0, "request_id": str(uuid.uuid4())})
        request = {**request_body("当前会话问题"), "session_id": new["session_id"]}
        self.manager.start(self.loaded, "F01", request)
        wait_done(self.manager, self.loaded, session_id=new["session_id"])
        messages = self.calls[-1]
        self.assertEqual([m["role"] for m in messages], ["system", "user", "user"])
        bound = json.loads(messages[1]["content"].split("\n", 1)[1])
        self.assertEqual([f["text"] for f in bound["CONFIRMED_FACTS"]], ["全合同当前事实"])
        self.assertEqual(len(bound["REPORT_OPINIONS"]), 2)
        self.assertNotIn("会话秘密", json.dumps(messages, ensure_ascii=False))
        self.assertEqual(self.manager.get(self.loaded, "F01", session_id="initial")["turns"][0]["user"], "上一会话秘密")

    def test_size_guard_fails_before_call_or_write_no_truncation(self):
        with patch("clause_discussion.MAX_INPUT_BYTES", 100):
            with self.assertRaisesRegex(ValueError, "输入上限"):
                self.manager.start(self.loaded, "F01", request_body("讨论"))
        self.assertEqual(self.calls, [])
        self.assertFalse(self.manager.root.exists())

    def test_revision_candidate_not_promoted_and_accepted_other_opinion_visible(self):
        loaded = full_loaded()
        discussion = ClauseDiscussionManager(self.root / "fake-discussion", FakeDelivery(), lambda *_: "答复")
        def completion(_, messages):
            result = json.loads(revision_answer(messages))
            result["finding"]["reason"] = "候选更新后的付款意见"
            result.update(change_summary="付款更新", related_finding_ids=["F01"])
            return json.dumps(result, ensure_ascii=False)
        revisions = OpinionRevisionManager(self.root / "versions", discussion, completion)
        initial = revisions.get(loaded, "F02")
        revisions.start(loaded, "F02", revision_body(initial))
        deadline = time.time() + 3
        while time.time() < deadline:
            done = revisions.get(loaded, "F02")
            if done["pending_candidate_id"] is None:
                break
            time.sleep(.01)
        self.assertEqual(done["candidates"][0]["status"], "completed", done["candidates"][0])
        row = discussion._source_context(revisions.apply(loaded), loaded["report"]["findings"][0])["report_opinions"][1]
        self.assertEqual(row["current_opinion"]["reason"], loaded["report"]["findings"][1]["reason"])
        revisions.accept(loaded, "F02", {"expected_revision": done["revision"], "expected_opinion_revision": 0,
            "expected_discussion_revision": 0, "expected_delivery_revision": 0, "candidate_id": done["candidates"][0]["id"]})
        applied = revisions.apply(loaded)
        source = discussion._source_context(applied, applied["report"]["findings"][0])
        row = source["report_opinions"][1]
        self.assertEqual(row["opinion_revision"], 1)
        self.assertEqual(row["current_opinion"]["reason"], "候选更新后的付款意见")
        source.update(turns=[], other_findings=[], contract_fingerprint=loaded["contract"]["fingerprint"])
        messages, _ = revisions._messages(source, applied["report"]["findings"][0], loaded["report"]["findings"][0], [])
        self.assertEqual(json.loads(messages[1]["content"])["CURRENT_CONTRACT"], source["current_contract"])


if __name__ == "__main__":
    unittest.main()
