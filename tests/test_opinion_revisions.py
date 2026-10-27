from __future__ import annotations

import copy
import io
import json
import tempfile
import threading
import time
import unittest
import uuid
from pathlib import Path
from unittest.mock import patch
from docx import Document
from docx.oxml import OxmlElement
from contract_core import parse_docx
from reader_delivery import DeliveryManager

from clause_discussion import ClauseDiscussionManager
from opinion_revisions import OpinionRevisionConflict, OpinionRevisionCorrupt, OpinionRevisionManager
from clause_discussion import _canonical, _sha256
from general_review import response_schema
from tests.test_clause_discussion import FakeDelivery, make_loaded, request_body, wait_done


def full_loaded():
    loaded = make_loaded()
    loaded["sources"][0]["checked_at"] = "2026-09-20"
    loaded["report"]["context"].update(contract_type="general", contract_type_label="服务合同", input_kind="拟签审阅")
    for index, finding in enumerate(loaded["report"]["findings"], 1):
        finding.update(rule_id=f"G{index:02d}", theme=finding["title"], issue_kind="商业", status="发现风险",
                       risk="中", evidence_status="原文可核", location_note="", source_ids=[], questions=[])
    loaded["report"]["checks"] = [{"rule_id": "G01", "title": "验收期", "status": "发现风险", "reason": "需核对"}]
    loaded["report"]["options"] = [{"id": "O1", "title": "同步验收与付款", "finding_ids": ["F01", "F02"], "steps": ["协调"],
                                      "preconditions": [], "cost": "", "remaining_risk": "", "source_ids": [], "clause_draft": "", "relationship": ""}]
    return loaded


def body(state, *, request_id=None):
    return {"expected_revision": state["revision"], "expected_opinion_revision": state["opinion_revision"],
            "expected_discussion_revision": state["context"]["discussion_revision"],
            "expected_delivery_revision": state["context"]["delivery_revision"],
            "request_id": request_id or str(uuid.uuid4()),
            "config": {"base_url": "http://127.0.0.1:9999/v1", "api_key": "secret-api-key", "model": "fixture"},
            "consent": True}


def answer(messages, *, action="revise", source_ids=None, bad_quote=False):
    data = json.loads(messages[1]["content"])
    finding = copy.deepcopy(data["CURRENT_FINDING"])
    finding["reason"] = "讨论后需将验收期延长。"
    if data.get("draft_plan_version") == 1:
        finding.pop("clause_draft", None)
        finding["draft_plan"] = {"status": "not_needed", "edits": [], "insertions": [], "questions": [],
            "comment": "本轮只更新内部判断，现合同文字保留。", "paragraph_actions": [
                {"piece_id": ref["id"], "action": "retain", "reason": "本轮只调整内部判断，保留相关原文。"} for ref in finding["refs"]]}
    else:
        finding["clause_draft"] = "十个工作日内验收。"
    finding["source_ids"] = source_ids or []
    if bad_quote:
        finding["refs"][0]["quote"] = "模型杜撰的引文"
    if action == "withdraw":
        finding.update(status="未检出", risk=None, reason="补充事实证明原判断不成立。")
    return json.dumps({"action": action, "finding": finding,
                       "change_summary": "依据讨论补充事实调整判断。",
                       "related_finding_ids": ["F02"]}, ensure_ascii=False)


def wait_revision(manager, loaded, timeout=3):
    deadline = time.time() + timeout
    while time.time() < deadline:
        state = manager.get(loaded, "F01")
        if state["pending_candidate_id"] is None:
            return state
        time.sleep(.01)
    raise AssertionError("revision did not finish")


class OpinionRevisionTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.loaded = full_loaded()
        self.delivery = FakeDelivery()
        self.discussions = ClauseDiscussionManager(self.root / "discussions", self.delivery, completion=lambda *_: "讨论回答")

    def test_generate_then_explicit_accept_and_apply_preserves_original(self):
        manager = OpinionRevisionManager(self.root / "revisions", self.discussions, completion=lambda _, msgs: answer(msgs))
        initial = manager.get(self.loaded, "F01")
        self.assertEqual(initial["revision"], 0)
        started, did_start = manager.start(self.loaded, "F01", body(initial))
        self.assertTrue(did_start)
        self.assertIsNotNone(started["pending_candidate_id"])
        completed = wait_revision(manager, self.loaded)
        self.assertEqual(completed["candidates"][0]["status"], "completed")
        self.assertEqual(manager.apply(self.loaded)["report"]["findings"][0]["reason"], self.loaded["report"]["findings"][0]["reason"])
        candidate = completed["candidates"][0]
        self.assertEqual(candidate["base_finding"], self.loaded["report"]["findings"][0])
        self.assertNotIn("raw_response", candidate)  # evidence stays in sidecar, not every browser poll
        disk = json.loads(manager._path(self.loaded, "F01").read_text(encoding="utf-8"))
        self.assertIn("CURRENT_FINDING", disk["candidates"][0]["prompt_snapshot"]["messages"][1]["content"])
        sent = json.loads(disk["candidates"][0]["prompt_snapshot"]["messages"][1]["content"])
        self.assertEqual(sent["OTHER_FINDINGS"][0]["id"], "F02")
        self.assertEqual(sent["OTHER_FINDINGS"][0]["title"], self.loaded["report"]["findings"][1]["title"])
        self.assertIn("第2条", initial["context"]["finding_labels"]["F02"])
        self.assertIn("讨论后需将验收期延长", disk["candidates"][0]["raw_response"])
        self.assertNotIn("secret-api-key", json.dumps(disk, ensure_ascii=False))
        accepted = manager.accept(self.loaded, "F01", {"expected_revision": completed["revision"],
                    "expected_opinion_revision": 0, "expected_discussion_revision": 0,
                    "expected_delivery_revision": 0, "candidate_id": candidate["id"]})
        self.assertEqual(accepted["opinion_revision"], 1)
        self.assertIsNotNone(accepted["candidates"][0]["accepted_at"])
        applied = manager.apply(self.loaded)
        self.assertEqual(applied["report"]["findings"][0]["reason"], "讨论后需将验收期延长。")
        self.assertEqual(applied["report"]["findings"][0]["opinion_revision"], 1)
        self.assertEqual(applied["report"]["findings"][0]["rule_id"], "G01")
        self.assertTrue(applied["report"]["options"][0]["need_review"])
        self.assertTrue(applied["report"]["checks"][0]["need_review"])
        self.assertEqual(self.loaded["report"]["findings"][0]["reason"], "验收时间可能不足。")
        self.assertEqual(manager.apply(applied)["report"]["findings"][0], applied["report"]["findings"][0])

    def test_request_id_and_restart_recovery_no_second_charge(self):
        gate = threading.Event()
        calls = []
        def completion(_, msgs):
            calls.append(1)
            gate.wait(2)
            return answer(msgs)
        manager = OpinionRevisionManager(self.root / "revisions", self.discussions, completion)
        original = manager.get(self.loaded, "F01")
        request = body(original)
        manager.start(self.loaded, "F01", request)
        restarted = OpinionRevisionManager(self.root / "revisions", self.discussions, completion)
        recovered = restarted.get(self.loaded, "F01")
        self.assertEqual(recovered["candidates"][0]["status"], "failed")
        cached, started = restarted.start(self.loaded, "F01", request)
        self.assertFalse(started)
        self.assertEqual(cached, recovered)
        changed_model = copy.deepcopy(request)
        changed_model["config"]["model"] = "another-model"
        with self.assertRaisesRegex(OpinionRevisionConflict, "请求的编号"):
            restarted.start(self.loaded, "F01", changed_model)
        gate.set()
        deadline = time.time() + 3
        while manager._active and time.time() < deadline:
            time.sleep(.01)
        self.assertEqual(len(calls), 1)

    def test_stale_discussion_and_delivery_reject_accept(self):
        manager = OpinionRevisionManager(self.root / "revisions", self.discussions, completion=lambda _, msgs: answer(msgs))
        state = manager.get(self.loaded, "F01")
        manager.start(self.loaded, "F01", body(state))
        done = wait_revision(manager, self.loaded)
        accept = {"expected_revision": done["revision"], "expected_opinion_revision": 0,
                  "expected_discussion_revision": 0, "expected_delivery_revision": 0,
                  "candidate_id": done["candidates"][0]["id"]}
        self.delivery.revision = 1
        with self.assertRaises(OpinionRevisionConflict):
            manager.accept(self.loaded, "F01", accept)
        self.delivery.revision = 0
        self.discussions.start(self.loaded, "F01", request_body("新事实"))
        wait_done(self.discussions, self.loaded)
        with self.assertRaises(OpinionRevisionConflict):
            manager.accept(self.loaded, "F01", accept)

    def test_invalid_paragraph_and_source_are_not_accepted(self):
        manager = OpinionRevisionManager(self.root / "revisions", self.discussions,
                                         completion=lambda _, msgs: answer(msgs).replace('"P001"', '"P999"').replace('"P000"', '"P999"'))
        manager.start(self.loaded, "F01", body(manager.get(self.loaded, "F01")))
        failed = wait_revision(manager, self.loaded)
        self.assertEqual(failed["candidates"][0]["status"], "failed")
        self.assertIn("原文引用", failed["candidates"][0]["error"])
        self.assertIsNone(failed["accepted_candidate_id"])
        manager2 = OpinionRevisionManager(self.root / "revisions", self.discussions,
                                          completion=lambda _, msgs: answer(msgs, source_ids=["FAKE-LINK"]))
        manager2.start(self.loaded, "F01", body(failed))
        second = wait_revision(manager2, self.loaded)
        self.assertEqual(second["candidates"][1]["status"], "failed")

    def test_real_search_source_stays_unverified(self):
        search = {"status": "searched", "queries": ["法律问题"],
                  "sources": [{"title": "实际工具结果", "url": "https://example.org/law"}], "request_count": 1}
        discussion = ClauseDiscussionManager(self.root / "discussions", self.delivery,
                    search_completion=lambda *_: {"text": "已检索。", "search": search})
        req = request_body("查证")
        req["config"]["base_url"] = "https://api.deepseek.com/v1"
        req["web_search"] = True
        discussion.start(self.loaded, "F01", req)
        wait_done(discussion, self.loaded)
        def completion(_, msgs):
            data = json.loads(msgs[1]["content"])
            sid = next(s["id"] for s in data["PROVIDED_SOURCES"] if s["id"].startswith("WEB-D"))
            return answer(msgs, source_ids=[sid])
        manager = OpinionRevisionManager(self.root / "revisions", discussion, completion)
        manager.start(self.loaded, "F01", body(manager.get(self.loaded, "F01")))
        state = wait_revision(manager, self.loaded)
        candidate = state["candidates"][0]
        self.assertEqual(candidate["status"], "completed")
        self.assertFalse(candidate["source_candidates"][0]["verified"])
        manager.accept(self.loaded, "F01", {"expected_revision": state["revision"], "expected_opinion_revision": 0,
                       "expected_discussion_revision": 2, "expected_delivery_revision": 0, "candidate_id": candidate["id"]})
        applied = manager.apply(self.loaded)
        self.assertEqual(applied["sources"][-1]["verified"], False)
        self.assertIn("实际工具结果", applied["sources"][-1]["title"])

    def test_withdraw_and_corrupt_preserved(self):
        manager = OpinionRevisionManager(self.root / "revisions", self.discussions,
                                         completion=lambda _, msgs: answer(msgs, action="withdraw"))
        manager.start(self.loaded, "F01", body(manager.get(self.loaded, "F01")))
        done = wait_revision(manager, self.loaded)
        manager.accept(self.loaded, "F01", {"expected_revision": done["revision"], "expected_opinion_revision": 0,
                       "expected_discussion_revision": 0, "expected_delivery_revision": 0,
                       "candidate_id": done["candidates"][0]["id"]})
        applied = manager.apply(self.loaded)
        finding = applied["report"]["findings"][0]
        self.assertEqual((finding["id"], finding["rule_id"], finding["status"], finding["risk"]),
                         ("F01", "G01", "未检出", None))
        self.assertEqual(len(applied["report"]["findings"]), 2)
        path = manager._path(self.loaded, "F01")
        path.write_text("{broken", encoding="utf-8")
        with self.assertRaises(OpinionRevisionCorrupt):
            manager.get(self.loaded, "F01")
        self.assertEqual(path.read_text(encoding="utf-8"), "{broken")

    def test_legacy_rule_id_and_discussion_after_accept(self):
        legacy = make_loaded()
        legacy["report"]["findings"][0].update(
            rule_id="SC-02", status="发现风险", risk="中", evidence_status="资料不足",
            source_ids=[], questions=[])
        legacy["report"]["findings"][1].update(
            rule_id="SC-03", status="需要法务判断", risk=None, evidence_status="资料不足",
            source_ids=[], questions=[])
        manager = OpinionRevisionManager(self.root / "revisions", self.discussions,
                                         completion=lambda _, msgs: answer(msgs))
        manager.start(legacy, "F01", body(manager.get(legacy, "F01")))
        done = wait_revision(manager, legacy)
        self.assertEqual(done["candidates"][0]["status"], "completed")
        self.assertEqual(done["candidates"][0]["finding"]["draft_plan_version"], 1)
        manager.accept(legacy, "F01", {"expected_revision": done["revision"], "expected_opinion_revision": 0,
                         "expected_discussion_revision": 0, "expected_delivery_revision": 0,
                         "candidate_id": done["candidates"][0]["id"]})
        applied = manager.apply(legacy)
        self.assertEqual(applied["report"]["findings"][0]["rule_id"], "SC-02")
        discussion = self.discussions.get(applied, "F01")
        self.assertEqual(discussion["opinion_revision"], 1)
        request = request_body("新版意见还有什么问题？")
        request["expected_opinion_revision"] = 1
        self.discussions.start(applied, "F01", request)
        wait_done(self.discussions, applied)
        with self.assertRaisesRegex(Exception, "意见版本已更新"):
            self.discussions.start(applied, "F01", {**request_body("请继续", revision=2),
                                                   "expected_opinion_revision": 0})

    def test_capacity_rejects_before_paid_call(self):
        calls = []
        manager = OpinionRevisionManager(self.root / "revisions", self.discussions,
                                         completion=lambda *_: calls.append(1) or "{}")
        with patch("opinion_revisions.MAX_STATE_BYTES", 1000):
            with self.assertRaisesRegex(ValueError, "容量上限"):
                manager.start(self.loaded, "F01", body(manager.get(self.loaded, "F01")))
        self.assertEqual(calls, [])
        self.assertFalse(manager._path(self.loaded, "F01").exists())

    def test_running_discussion_blocks_generation(self):
        gate = threading.Event()
        discussion = ClauseDiscussionManager(self.root / "discussions", self.delivery,
                                              completion=lambda *_: gate.wait(2) or "讨论回答")
        discussion.start(self.loaded, "F01", request_body("还在讨论"))
        calls = []
        manager = OpinionRevisionManager(self.root / "revisions", discussion,
                                         completion=lambda *_: calls.append(1) or "{}")
        try:
            state = manager.get(self.loaded, "F01")
            with self.assertRaisesRegex(OpinionRevisionConflict, "讨论还没有回答完"):
                manager.start(self.loaded, "F01", body(state))
            self.assertEqual(calls, [])
        finally:
            gate.set()
            wait_done(discussion, self.loaded)

    def test_second_accept_keeps_first_candidate_history(self):
        count = 0
        def completion(_, msgs):
            nonlocal count
            count += 1
            data = json.loads(answer(msgs))
            data["finding"]["reason"] = f"第{count}版判断。"
            return json.dumps(data, ensure_ascii=False)
        manager = OpinionRevisionManager(self.root / "revisions", self.discussions, completion)
        first = manager.get(self.loaded, "F01")
        manager.start(self.loaded, "F01", body(first))
        done = wait_revision(manager, self.loaded)
        first_id = done["candidates"][0]["id"]
        manager.accept(self.loaded, "F01", {"expected_revision": done["revision"], "expected_opinion_revision": 0,
                       "expected_discussion_revision": 0, "expected_delivery_revision": 0, "candidate_id": first_id})
        applied = manager.apply(self.loaded)
        current = manager.get(applied, "F01")
        manager.start(applied, "F01", body(current))
        second = wait_revision(manager, applied)
        self.assertEqual(second["candidates"][1]["base_finding"]["reason"], "第1版判断。")
        second_id = second["candidates"][1]["id"]
        manager.accept(applied, "F01", {"expected_revision": second["revision"], "expected_opinion_revision": 1,
                       "expected_discussion_revision": 0, "expected_delivery_revision": 0, "candidate_id": second_id})
        final = manager.get(manager.apply(self.loaded), "F01")
        self.assertEqual(final["opinion_revision"], 2)
        self.assertEqual([x["finding"]["reason"] for x in final["candidates"]], ["第1版判断。", "第2版判断。"])
        self.assertTrue(all(x["accepted_at"] for x in final["candidates"]))

    def test_ready_plan_on_legacy_general_has_one_text_source_and_readback_validation(self):
        self.delivery.configured = False  # This test deliberately covers text-only history.
        def ready_answer(_, msgs):
            data = json.loads(answer(msgs))
            value = data["finding"]
            value["draft_plan"] = {"status": "ready", "edits": [
                {"piece_id": "P001", "new_text": "4.1 甲方应在十个工作日内验收。"}],
                "insertions": [], "questions": [], "comment": "请核对十个工作日的验收期限。",
                "paragraph_actions": [{"piece_id": "P001", "action": "modify", "reason": "修改验收期限以覆盖测试工作量。"}]}
            return json.dumps(data, ensure_ascii=False)
        manager = OpinionRevisionManager(self.root / "revisions", self.discussions, ready_answer)
        manager.start(self.loaded, "F01", body(manager.get(self.loaded, "F01")))
        done = wait_revision(manager, self.loaded)
        candidate = done["candidates"][0]
        self.assertEqual(candidate["status"], "completed", candidate)
        self.assertNotIn("draft_plan_version", self.loaded["report"]["context"])
        self.assertEqual(candidate["finding"]["clause_draft"], "4.1 甲方应在十个工作日内验收。")
        self.assertEqual(candidate["finding"]["draft_plan_basis"]["facts_digest_version"], 1)
        path = manager._path(self.loaded, "F01")
        before = path.read_bytes()
        restarted = OpinionRevisionManager(self.root / "revisions", self.discussions, lambda *_: self.fail("read called model"))
        self.assertEqual(restarted.get(self.loaded, "F01")["candidates"][0], candidate)
        self.assertEqual(path.read_bytes(), before)
        saved = json.loads(before)
        saved["candidates"][0]["finding"]["draft_plan"]["edits"][0]["new_text"] = "4.1 被静默换成别的正文。"
        path.write_text(json.dumps(saved, ensure_ascii=False), encoding="utf-8")
        with self.assertRaises(OpinionRevisionCorrupt):
            restarted.get(self.loaded, "F01")

    def test_legacy_configured_docx_without_loaded_path_rejects_complex_ready_revision(self):
        document = Document()
        paragraph = document.add_paragraph("4.1 甲方应在七日内验收。")
        paragraph._p.append(OxmlElement("w:bookmarkStart"))
        document.add_paragraph("4.2 甲方应在验收后付款。")
        stream = io.BytesIO()
        document.save(stream)
        source_path = self.root / "legacy.docx"
        source_path.write_bytes(stream.getvalue())
        parsed = parse_docx(stream.getvalue())
        for general in (True, False):
            with self.subTest(general=general):
                loaded = full_loaded()
                loaded["contract"] = parsed.model_dump()
                loaded["report"]["fingerprint"] = parsed.fingerprint
                loaded["paragraphs"] = [{"id": p.id, "location": p.id} for p in parsed.pieces]
                for finding, piece in zip(loaded["report"]["findings"], parsed.pieces):
                    finding["refs"] = [{"id": piece.id, "quote": piece.text}]
                if not general:
                    loaded["report"]["context"]["contract_type"] = "software_service"
                    for index, finding in enumerate(loaded["report"]["findings"]):
                        finding["rule_id"] = f"SC-0{index + 2}"
                self.assertNotIn("source_path", loaded)
                root = self.root / ("general" if general else "service")
                delivery = DeliveryManager(root / "delivery", {loaded["id"]: source_path})
                discussions = ClauseDiscussionManager(root / "discussions", delivery, completion=lambda *_: "讨论回复")
                def ready(_, msgs):
                    data = json.loads(answer(msgs))
                    data["finding"]["draft_plan"] = {"status": "ready", "edits": [
                        {"piece_id": "P000", "new_text": "4.1 甲方应在十个工作日内验收。"}],
                        "insertions": [], "questions": [], "comment": "请核对验收期限的调整。",
                        "paragraph_actions": [{"piece_id": "P000", "action": "modify", "reason": "扩大期限以覆盖测试工作量。"}]}
                    return json.dumps(data, ensure_ascii=False)
                manager = OpinionRevisionManager(root / "revisions", discussions, ready)
                manager.start(loaded, "F01", body(manager.get(loaded, "F01")))
                done = wait_revision(manager, loaded)
                candidate = done["candidates"][0]
                self.assertEqual(candidate["status"], "failed", candidate)
                self.assertIn("复杂段落", candidate["error"])
                self.assertIsNone(candidate["finding"])
                self.assertIsNone(done["accepted_candidate_id"])
        self.assertEqual(source_path.read_bytes(), stream.getvalue())

    def test_saved_legacy_candidate_without_plan_version_reads_without_upgrade_or_write(self):
        manager = OpinionRevisionManager(self.root / "revisions", self.discussions, lambda _, msgs: answer(msgs))
        manager.start(self.loaded, "F01", body(manager.get(self.loaded, "F01")))
        wait_revision(manager, self.loaded)
        path = manager._path(self.loaded, "F01")
        saved = json.loads(path.read_text(encoding="utf-8"))
        item = saved["candidates"][0]
        item.pop("draft_plan_version")
        item.pop("draft_text_version")  # Legacy candidates predate both versions.
        for key in ("draft_plan", "draft_plan_version", "draft_plan_basis"):
            item["finding"].pop(key, None)
        item["finding"]["clause_draft"] = "旧候选只保存展示建议，尚无可执行plan。"
        snapshot = item["prompt_snapshot"]
        snapshot.pop("draft_plan_version")
        snapshot.pop("draft_text_version")
        snapshot.pop("draft_plan_insertion_anchors")
        packet = json.loads(snapshot["messages"][1]["content"])
        packet.pop("draft_plan_version")
        packet.pop("draft_plan_insertion_anchors")
        packet["CURRENT_FINDING"].pop("draft_plan", None)
        packet["CURRENT_FINDING"]["clause_draft"] = self.loaded["report"]["findings"][0]["clause_draft"]
        packet["required_schema"]["properties"]["finding"] = response_schema()["properties"]["findings"]["items"]
        snapshot["messages"][1]["content"] = json.dumps(packet, ensure_ascii=False)
        snapshot["messages_sha256"] = _sha256(_canonical(snapshot["messages"]))
        raw = json.loads(item["raw_response"])
        raw["finding"] = copy.deepcopy(item["finding"])
        item["raw_response"] = json.dumps(raw, ensure_ascii=False)
        path.write_text(json.dumps(saved, ensure_ascii=False), encoding="utf-8")
        before = path.read_bytes()
        restarted = OpinionRevisionManager(self.root / "revisions", self.discussions, lambda *_: self.fail("read called model"))
        state = restarted.get(self.loaded, "F01")
        self.assertNotIn("draft_plan_version", state["candidates"][0])
        self.assertNotIn("draft_plan", state["candidates"][0]["finding"])
        self.assertEqual(path.read_bytes(), before)

    def test_summary_nested_inside_finding_is_lifted_out(self):
        def nested(msgs):
            data = json.loads(answer(msgs))
            data["finding"]["change_summary"] = data.pop("change_summary")
            data["finding"]["related_finding_ids"] = data.pop("related_finding_ids")
            return json.dumps(data, ensure_ascii=False)
        manager = OpinionRevisionManager(self.root / "revisions", self.discussions, lambda _, msgs: nested(msgs))
        manager.start(self.loaded, "F01", body(manager.get(self.loaded, "F01")))
        state = wait_revision(manager, self.loaded)
        candidate = state["candidates"][0]
        self.assertEqual(candidate["status"], "completed", candidate.get("error"))
        self.assertNotIn("change_summary", candidate["finding"])
        # Reading it back revalidates the same raw response.
        self.assertEqual(manager.get(self.loaded, "F01")["candidates"][0]["status"], "completed")

    def test_new_candidates_ask_for_clean_comments_and_version_one_still_reads(self):
        from draft_plan import COMMENT_RULE, DRAFT_TEXT_VERSION
        manager = OpinionRevisionManager(self.root / "revisions", self.discussions, lambda _, msgs: answer(msgs))
        manager.start(self.loaded, "F01", body(manager.get(self.loaded, "F01")))
        wait_revision(manager, self.loaded)
        path = manager._path(self.loaded, "F01")
        saved = json.loads(path.read_text(encoding="utf-8"))
        item = saved["candidates"][0]
        self.assertEqual(item["draft_text_version"], DRAFT_TEXT_VERSION)
        self.assertIn(COMMENT_RULE, item["prompt_snapshot"]["messages"][0]["content"])
        # A candidate saved under version 1 keeps reading as it was.
        item["draft_text_version"] = item["prompt_snapshot"]["draft_text_version"] = 1
        path.write_text(json.dumps(saved, ensure_ascii=False), encoding="utf-8")
        before = path.read_bytes()
        restarted = OpinionRevisionManager(self.root / "revisions", self.discussions, lambda *_: self.fail("read called model"))
        state = restarted.get(self.loaded, "F01")
        self.assertEqual(state["candidates"][0]["status"], item["status"])
        self.assertIn("draft_plan", state["candidates"][0]["finding"])
        self.assertEqual(path.read_bytes(), before)

    def test_new_version_missing_plan_is_failed_not_business_needs_input(self):
        def missing(_, messages):
            data = json.loads(answer(messages))
            data["finding"].pop("draft_plan")
            return json.dumps(data)
        manager = OpinionRevisionManager(self.root / "revisions", self.discussions, missing)
        manager.start(self.loaded, "F01", body(manager.get(self.loaded, "F01")))
        done = wait_revision(manager, self.loaded)
        self.assertEqual(done["candidates"][0]["status"], "failed")
        self.assertIn("draft_plan", done["candidates"][0]["error"])


if __name__ == "__main__":
    unittest.main()
