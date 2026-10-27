from __future__ import annotations

import io
import json
import os
import sys
import unittest
import urllib.error
import zipfile
from unittest import mock

from docx import Document
from docx.oxml import OxmlElement


PROJECT_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
if PROJECT_DIR not in sys.path:
    sys.path.insert(0, PROJECT_DIR)

import contract_core
import model_client


def current_sources(checked_at: str = "2026-09-20") -> list[dict]:
    ids = ["LCL14", "LCL23_24", "LCL30_50", "SPC10_11", "SPC13_14", "SHW7", "NATIONALW9"]
    return [
        {
            "id": source_id,
            "title": source_id,
            "authority": "测试机关",
            "url": "https://example.invalid/source",
            "article": "测试条文",
            "text": f"{source_id} 的已核测试文本",
            "region": "上海" if source_id == "SHW7" else "全国",
            "effective_from": "2025-01-01",
            "effective_to": None,
            "checked_at": checked_at,
            "verified": True,
            "version_status": "已核官方文本",
            "notes": "",
        }
        for source_id in ids
    ]


def context(**overrides) -> dict:
    value = {
        "region": "上海",
        "review_date": "2026-09-20",
        "input_kind": "模板审阅",
        "payday": 10,
    }
    value.update(overrides)
    return value


def docx_bytes(paragraphs: list[str], table: list[list[str]] | None = None) -> bytes:
    document = Document()
    for text in paragraphs:
        document.add_paragraph(text)
    if table is not None:
        word_table = document.add_table(rows=len(table), cols=len(table[0]))
        for row_index, row in enumerate(table):
            for col_index, text in enumerate(row):
                word_table.cell(row_index, col_index).text = text
    stream = io.BytesIO()
    document.save(stream)
    return stream.getvalue()


def mutate_document_xml(data: bytes, mutator) -> bytes:
    source = zipfile.ZipFile(io.BytesIO(data))
    output = io.BytesIO()
    with source, zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as target:
        for info in source.infolist():
            member = source.read(info.filename)
            if info.filename == "word/document.xml":
                member = mutator(member)
            target.writestr(info, member)
    return output.getvalue()


class ParseTests(unittest.TestCase):
    def test_plain_text_preserves_duplicate_blank_lines_and_is_serializable(self):
        parsed = contract_core.parse_text("甲\n\n\n乙\n")
        self.assertEqual([piece.id for piece in parsed.pieces], ["P000", "P001", "P002", "P003", "P004"])
        self.assertEqual([piece.text for piece in parsed.pieces], ["甲", "", "", "乙", ""])
        json.dumps(parsed.model_dump(), ensure_ascii=False)
        self.assertEqual(parsed.fingerprint, contract_core.parse_text("甲\r\n\r\n\r\n乙\r\n").fingerprint)

    def test_plain_text_size_and_null_byte_are_rejected(self):
        with self.assertRaisesRegex(ValueError, "5 MB"):
            contract_core.parse_text("好" * (contract_core.MAX_TEXT_BYTES // 3 + 2))
        with self.assertRaisesRegex(ValueError, "空字节"):
            contract_core.parse_text("合同\x00正文")

    def test_docx_keeps_244_body_paragraph_numbers_and_table_reading_order(self):
        paragraphs = [f"第{i}段" if i not in {1, 2, 100} else "" for i in range(244)]
        parsed = contract_core.parse_docx(docx_bytes(paragraphs, [["甲", "乙"], ["", "丁"]]))
        body = [piece for piece in parsed.pieces if piece.id.startswith("P")]
        self.assertEqual(len(body), 244)
        self.assertEqual(body[0].id, "P000")
        self.assertEqual(body[1].text, "")
        self.assertEqual(body[2].text, "")
        self.assertEqual(body[-1].id, "P243")
        self.assertEqual(
            [(piece.id, piece.text) for piece in parsed.pieces[-4:]],
            [("T0R0C0P0", "甲"), ("T0R0C1P0", "乙"), ("T0R1C0P0", ""), ("T0R1C1P0", "丁")],
        )

    def test_docx_with_tracked_changes_is_rejected_instead_of_partially_reviewed(self):
        original = docx_bytes(["旧文本"])

        def add_revision(xml: bytes) -> bytes:
            return xml.replace(b"<w:r>", b'<w:ins w:id="1" w:author="x"><w:r>', 1).replace(b"</w:r>", b"</w:r></w:ins>", 1)

        with self.assertRaisesRegex(ValueError, "修订"):
            contract_core.parse_docx(mutate_document_xml(original, add_revision))

    def test_docx_with_textbox_is_rejected(self):
        original = docx_bytes(["正文"])

        def add_textbox(xml: bytes) -> bytes:
            marker = b"</w:body>"
            box = "<w:p><w:r><w:txbxContent><w:p><w:r><w:t>隐藏条款</w:t></w:r></w:p></w:txbxContent></w:r></w:p>".encode("utf-8")
            return xml.replace(marker, box + marker)

        with self.assertRaisesRegex(ValueError, "文本框"):
            contract_core.parse_docx(mutate_document_xml(original, add_textbox))

    def test_page_number_only_footer_is_ignored_with_warning(self):
        document = Document()
        document.add_paragraph("合同正文")
        paragraph = document.sections[0].footer.paragraphs[0]
        paragraph.add_run("第 ")
        field_begin = OxmlElement("w:fldChar")
        field_begin.set("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}fldCharType", "begin")
        instruction = OxmlElement("w:instrText")
        instruction.text = " PAGE "
        field_separate = OxmlElement("w:fldChar")
        field_separate.set("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}fldCharType", "separate")
        result = OxmlElement("w:t")
        result.text = "1"
        field_end = OxmlElement("w:fldChar")
        field_end.set("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}fldCharType", "end")
        for node in (field_begin, instruction, field_separate, result, field_end):
            paragraph._p.append(node)
        paragraph.add_run(" 页")
        stream = io.BytesIO()
        document.save(stream)
        parsed = contract_core.parse_docx(stream.getvalue())
        self.assertEqual(parsed.pieces[0].text, "合同正文")
        self.assertTrue(any("PAGE/NUMPAGES" in warning for warning in parsed.warnings))

    def test_substantive_header_is_rejected(self):
        document = Document()
        document.add_paragraph("合同正文")
        document.sections[0].header.paragraphs[0].text = "页眉中的补充义务"
        stream = io.BytesIO()
        document.save(stream)
        with self.assertRaisesRegex(ValueError, "页眉、页脚"):
            contract_core.parse_docx(stream.getvalue())

    def test_docx_xml_size_limit_is_enforced_before_parsing(self):
        stream = io.BytesIO()
        with zipfile.ZipFile(stream, "w", zipfile.ZIP_DEFLATED) as archive:
            archive.writestr("word/document.xml", b"<x>123456789</x>")
        with mock.patch.object(contract_core, "MAX_XML_PART_BYTES", 8):
            with self.assertRaisesRegex(ValueError, "XML 过大"):
                contract_core.parse_docx(stream.getvalue())


class LocalReviewTests(unittest.TestCase):
    def test_unknown_region_never_passes(self):
        report = contract_core.review_local(contract_core.parse_text("普通合同"), context(region="广东"), current_sources())
        self.assertEqual(report["mode"], "local_rules")
        self.assertEqual(len(report["findings"]), 3)
        self.assertTrue(all(item["status"] == "信息不足" for item in report["findings"]))
        self.assertTrue(all(item["risk"] is None for item in report["findings"]))

    def test_review_date_after_source_check_date_requires_recheck(self):
        report = contract_core.review_local(
            contract_core.parse_text("普通合同"),
            context(review_date="2026-09-21"),
            current_sources(checked_at="2026-09-20"),
        )
        self.assertTrue(all(item["status"] == "信息不足" for item in report["findings"]))
        self.assertTrue(any("版本待复核" in question for item in report["findings"] for question in item["questions"]))

    def test_lc01_explicit_waiver_is_high_risk_with_exact_quote(self):
        text = "员工未提出续签申请，视为员工放弃签订无固定期限劳动合同。"
        report = contract_core.review_local(contract_core.parse_text(text), context(), current_sources())
        finding = report["findings"][0]
        self.assertEqual((finding["rule_id"], finding["status"], finding["risk"]), ("LC-01", "发现风险", "高"))
        self.assertIn(finding["refs"][0]["quote"], text)
        self.assertEqual(finding["source_ids"], ["LCL14", "SPC10_11"])

    def test_lc01_written_application_waiver_wording_is_detected(self):
        text = "未书面提出申请的，视为乙方放弃续签无固定期限劳动合同。"
        finding = contract_core.review_local(contract_core.parse_text(text), context(), current_sources())["findings"][0]
        self.assertEqual((finding["status"], finding["risk"]), ("发现风险", "高"))
        self.assertIn(finding["refs"][0]["quote"], text)

    def test_lc01_short_sample_waiver_is_detected(self):
        text = "乙方应在合同期满前一个月书面申请订立无固定期限劳动合同，未申请视为放弃。"
        finding = contract_core.review_local(contract_core.parse_text(text), context(), current_sources())["findings"][0]
        self.assertEqual((finding["status"], finding["risk"]), ("发现风险", "高"))

    def test_lc01_prohibition_is_not_reversed_into_high_risk(self):
        text = "不得将未申请视为放弃无固定期限劳动合同。"
        finding = contract_core.review_local(contract_core.parse_text(text), context(), current_sources())["findings"][0]
        self.assertNotEqual(finding["status"], "发现风险")
        self.assertIsNone(finding["risk"])

    def test_lc01_negative_wording_is_not_misread_as_waiver(self):
        text = "员工未申请续签，不得视为放弃订立无固定期限劳动合同。"
        finding = contract_core.review_local(contract_core.parse_text(text), context(), current_sources())["findings"][0]
        self.assertNotEqual(finding["status"], "发现风险")

    def test_lc02_wage_withholding_and_compensation_are_not_conflated(self):
        risky = "员工未完成离职交接的，公司暂缓支付工资，且不构成拖欠工资。"
        finding = contract_core.review_local(contract_core.parse_text(risky), context(), current_sources())["findings"][1]
        self.assertEqual((finding["status"], finding["risk"]), ("发现风险", "高"))
        self.assertIn(finding["refs"][0]["quote"], risky)

        compensation_only = "员工未完成工作交接的，公司暂缓支付经济补偿金。"
        finding = contract_core.review_local(contract_core.parse_text(compensation_only), context(), current_sources())["findings"][1]
        self.assertEqual(finding["status"], "未检出")
        self.assertIsNone(finding["risk"])

    def test_lc02_complete_handover_before_wage_settlement_is_detected(self):
        text = (
            "未发放的工资必须要在乙方办理完工作交接以后才能够办理工资结算。"
            "乙方未办理完工作交接而导致工资迟发的，不构成拖欠或克扣工资。"
        )
        finding = contract_core.review_local(contract_core.parse_text(text), context(), current_sources())["findings"][1]
        self.assertEqual((finding["status"], finding["risk"]), ("发现风险", "高"))
        self.assertIn(finding["refs"][0]["quote"], text)

    def test_lc02_short_sample_nonpayment_is_detected(self):
        text = "乙方离职时未完成工作交接的，甲方不支付工资，且不构成拖欠。"
        finding = contract_core.review_local(contract_core.parse_text(text), context(), current_sources())["findings"][1]
        self.assertEqual((finding["status"], finding["risk"]), ("发现风险", "高"))

    def test_lc02_prohibition_is_not_reversed_into_high_risk(self):
        text = "不得因未完成交接而不支付工资。"
        finding = contract_core.review_local(contract_core.parse_text(text), context(), current_sources())["findings"][1]
        self.assertNotEqual(finding["status"], "发现风险")
        self.assertIsNone(finding["risk"])

    def test_lc02_high_risk_keeps_later_payday_and_wage_exception_context(self):
        text = (
            "未发放的工资必须要在乙方办理完工作交接以后才能够办理工资结算。\n"
            "乙方同意公司于下个工资支付日将离职工资一次性足额结算。\n"
            "公司可暂不支付经济补偿至交接完成为止（工资除外）。"
        )
        finding = contract_core.review_local(contract_core.parse_text(text), context(), current_sources())["findings"][1]
        self.assertEqual((finding["status"], finding["risk"]), ("发现风险", "高"))
        self.assertEqual([ref["id"] for ref in finding["refs"]], ["P000", "P001", "P002"])
        self.assertIn("不会自动消除前段文字的矛盾", finding["reason"])
        self.assertTrue(any("工资除外" in ref["quote"] for ref in finding["refs"]))

    def test_lc02_first_payday_is_left_for_legal_judgment(self):
        text = "劳动合同解除后，工资于离职后的首个发薪日结清。"
        finding = contract_core.review_local(contract_core.parse_text(text), context(payday=None), current_sources())["findings"][1]
        self.assertEqual(finding["status"], "需要法务判断")
        self.assertIsNone(finding["risk"])
        self.assertIn("请补充公司的固定发薪日。", finding["questions"])

    def test_confidentiality_alone_is_not_called_noncompete(self):
        text = "员工对在职期间知悉的商业秘密承担保密义务。"
        finding = contract_core.review_local(contract_core.parse_text(text), context(), current_sources())["findings"][2]
        self.assertEqual(finding["status"], "不适用")
        self.assertIsNone(finding["risk"])

    def test_post_employment_customer_restriction_keeps_monthly_compensation(self):
        text = (
            "员工离职后两年内不得接触或服务公司的现有客户。\n"
            "竞业限制期间，公司每月支付经济补偿。"
        )
        finding = contract_core.review_local(contract_core.parse_text(text), context(), current_sources())["findings"][2]
        self.assertEqual(finding["status"], "需要法务判断")
        self.assertIsNone(finding["risk"])
        self.assertIn("也检出了按月补偿", finding["reason"])
        self.assertEqual(len(finding["refs"]), 2)


def ai_finding(rule_id: str, quote: str = "", source_ids: list[str] | None = None) -> dict:
    refs = [{"id": "P000", "quote": quote}] if quote else []
    return {
        "rule_id": rule_id,
        "title": "模型标题",
        "status": "发现风险" if quote else "未检出",
        "risk": "中" if quote else None,
        "evidence_status": "依据已核",
        "refs": refs,
        "source_ids": source_ids or [],
        "reason": "模型理由",
        "suggestion": "模型建议",
        "clause_draft": "",
        "questions": [],
    }


class AiValidationTests(unittest.TestCase):
    def test_hallucinated_quote_and_unknown_source_fail_only_their_rules(self):
        parsed = contract_core.parse_text("合同原文")
        response = {
            "findings": [
                ai_finding("LC-01", "不存在的原文", ["LCL14"]),
                ai_finding("LC-02", "合同原文", ["FAKE_SOURCE"]),
                ai_finding("LC-03"),
            ],
            "uncovered": ["模型未覆盖附件"],
        }
        with mock.patch("contract_core.chat_completion", return_value=json.dumps(response, ensure_ascii=False)):
            report = contract_core.review_ai(
                parsed,
                context(),
                current_sources(),
                {"base_url": "https://api.example.com/v1", "api_key": "secret", "model": "demo"},
            )
        self.assertEqual([item["status"] for item in report["findings"]], ["处理失败", "处理失败", "未检出"])
        self.assertEqual(len(report["findings"]), 3)
        self.assertTrue(report["metadata"]["model_called"])
        self.assertEqual(len(report["metadata"]["sources_digest"]), 64)
        self.assertGreaterEqual(len(report["uncovered"]), 3)

    def test_missing_and_duplicate_rules_do_not_become_passes(self):
        response = {"findings": [ai_finding("LC-01"), ai_finding("LC-01"), ai_finding("LC-03")], "uncovered": []}
        with mock.patch("contract_core.chat_completion", return_value=json.dumps(response, ensure_ascii=False)):
            report = contract_core.review_ai(
                contract_core.parse_text("正文"),
                context(),
                current_sources(),
                {"base_url": "https://api.example.com/v1", "api_key": "secret", "model": "demo"},
            )
        by_rule = {item["rule_id"]: item for item in report["findings"]}
        self.assertEqual(by_rule["LC-01"]["status"], "处理失败")
        self.assertEqual(by_rule["LC-02"]["status"], "处理失败")
        self.assertEqual(by_rule["LC-03"]["status"], "未检出")

    def test_template_employee_history_guess_is_rejected(self):
        bad = ai_finding("LC-03")
        bad["reason"] = "该员工属于高级管理人员并且已经离职。"
        response = {"findings": [ai_finding("LC-01"), ai_finding("LC-02"), bad], "uncovered": []}
        with mock.patch("contract_core.chat_completion", return_value=json.dumps(response, ensure_ascii=False)):
            report = contract_core.review_ai(
                contract_core.parse_text("正文"),
                context(),
                current_sources(),
                {"base_url": "https://api.example.com/v1", "api_key": "secret", "model": "demo"},
            )
        self.assertEqual(report["findings"][2]["status"], "处理失败")

    def test_wrong_json_types_are_isolated_instead_of_crashing(self):
        wrong_status = ai_finding("LC-01")
        wrong_status["status"] = ["未检出"]
        wrong_risk = ai_finding("LC-02")
        wrong_risk["risk"] = {"level": "中"}
        response = {
            "findings": [wrong_status, wrong_risk, {"rule_id": ["LC-03"]}],
            "uncovered": [],
        }
        with mock.patch("contract_core.chat_completion", return_value=json.dumps(response, ensure_ascii=False)):
            report = contract_core.review_ai(
                contract_core.parse_text("正文"),
                context(),
                current_sources(),
                {"base_url": "https://api.example.com/v1", "api_key": "secret", "model": "demo"},
            )
        self.assertEqual([item["status"] for item in report["findings"]], ["处理失败", "处理失败", "处理失败"])

    def test_evidence_shortage_cannot_carry_medium_risk(self):
        bad = ai_finding("LC-02")
        bad["evidence_status"] = "资料不足"
        bad["risk"] = "中"
        response = {"findings": [ai_finding("LC-01"), bad, ai_finding("LC-03")], "uncovered": []}
        with mock.patch("contract_core.chat_completion", return_value=json.dumps(response, ensure_ascii=False)):
            report = contract_core.review_ai(
                contract_core.parse_text("正文"),
                context(),
                current_sources(),
                {"base_url": "https://api.example.com/v1", "api_key": "secret", "model": "demo"},
            )
        self.assertEqual(report["findings"][1]["status"], "处理失败")

    def test_ai_gate_report_says_model_was_not_called(self):
        with mock.patch("contract_core.chat_completion") as called:
            report = contract_core.review_ai(
                contract_core.parse_text("正文"),
                context(region="广东"),
                current_sources(),
                {"base_url": "https://api.example.com/v1", "api_key": "secret", "model": "demo"},
            )
        called.assert_not_called()
        self.assertFalse(report["metadata"]["model_called"])


class ModelClientTests(unittest.TestCase):
    def test_non_loopback_http_and_query_are_rejected(self):
        with self.assertRaisesRegex(ValueError, "HTTPS"):
            model_client.chat_completion(
                {"base_url": "http://api.example.com/v1", "api_key": "x", "model": "m"}, [{"role": "user", "content": "OK"}]
            )
        with self.assertRaisesRegex(ValueError, "查询参数"):
            model_client.chat_completion(
                {"base_url": "https://api.example.com/v1?x=1", "api_key": "x", "model": "m"}, [{"role": "user", "content": "OK"}]
            )

    def test_network_failure_does_not_leak_api_key(self):
        secret = "sk-super-private"
        fake_opener = mock.Mock()
        fake_opener.open.side_effect = urllib.error.URLError(f"failure with {secret}")
        with mock.patch("model_client.urllib.request.build_opener", return_value=fake_opener):
            with self.assertRaises(model_client.ModelClientError) as caught:
                model_client.chat_completion(
                    {"base_url": "https://api.example.com/v1", "api_key": secret, "model": "m"},
                    [{"role": "user", "content": "请只回复OK"}],
                )
        self.assertNotIn(secret, str(caught.exception))

    def test_http_error_body_and_key_are_not_exposed(self):
        secret = "sk-secret"
        fake_opener = mock.Mock()
        fake_opener.open.side_effect = urllib.error.HTTPError(
            "https://api.example.com/v1/chat/completions",
            401,
            "Unauthorized",
            {},
            io.BytesIO(f"echo {secret}".encode()),
        )
        with mock.patch("model_client.urllib.request.build_opener", return_value=fake_opener):
            with self.assertRaises(model_client.ModelClientError) as caught:
                model_client.chat_completion(
                    {"base_url": "https://api.example.com/v1", "api_key": secret, "model": "m"},
                    [{"role": "user", "content": "OK"}],
                )
        self.assertEqual(caught.exception.code, "http_error")
        self.assertEqual(caught.exception.http_status, 401)
        self.assertIn("HTTP 401", str(caught.exception))
        self.assertNotIn(secret, str(caught.exception))

    def test_success_uses_no_redirect_handler_and_requires_text(self):
        response = mock.MagicMock()
        response.__enter__.return_value.read.return_value = json.dumps(
            {"choices": [{"message": {"content": "OK"}}]}
        ).encode()
        opener = mock.Mock()
        opener.open.return_value = response
        with mock.patch("model_client.urllib.request.build_opener", return_value=opener) as build:
            result = model_client.chat_completion(
                {"base_url": "https://openrouter.ai/api/v1", "api_key": "secret", "model": "demo"},
                [{"role": "user", "content": "请只回复OK"}],
            )
        self.assertEqual(result, "OK")
        self.assertTrue(any(isinstance(handler, model_client._NoRedirect) for handler in build.call_args.args))
        request = opener.open.call_args.args[0]
        self.assertEqual(request.full_url, "https://openrouter.ai/api/v1/chat/completions")


if __name__ == "__main__":
    unittest.main()
