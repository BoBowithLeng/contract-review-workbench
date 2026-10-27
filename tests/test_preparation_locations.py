"""Preparation questions are shown before any report exists; positions must still read as clauses."""
from __future__ import annotations

import base64
import json
import re
import tempfile
import unittest
from pathlib import Path

from docx import Document

from review_jobs import ReviewJobManager
from tests.test_word_numbering import list_definition, number, saved

PROJECT = Path(__file__).resolve().parents[1]


class PreparationLocationTests(unittest.TestCase):
    def test_job_preview_maps_paragraph_ids_to_short_clause_numbers(self):
        doc = Document()
        list_definition(doc)
        number(doc.add_paragraph('开发软件', style='Heading 2'), 0)
        number(doc.add_paragraph('硬件设备、服务器及相关费用由甲方承担。'), 1)
        number(doc.add_paragraph('交付范围以附件为准。'), 1)
        with tempfile.TemporaryDirectory() as temp:
            catalog = json.loads((PROJECT / 'builtin_checklists.json').read_text(encoding='utf-8-sig'))
            manager = ReviewJobManager(Path(temp) / 'jobs', runner=None, engines=[],
                                       prompt_path=PROJECT / 'prompts' / 'initial_review.md')
            created = manager.create({'document_name': '测试.docx', 'docx_base64': base64.b64encode(saved(doc)).decode(),
                                      'context': {'contract_type': 'general', 'our_party': '甲方', 'our_role': '委托方'},
                                      'checklist_snapshot': catalog['checklists'][0]})
            for preview in (created['preview'], manager.get(created['job']['job_id'])['preview']):
                locations = preview['locations']
                self.assertEqual(locations['P001'], '第1.1条')
                self.assertEqual(locations['P002'], '第1.2条')
                self.assertFalse(any(re.search(r'\bP\d+\b|「', label) for label in locations.values()))


if __name__ == '__main__':
    unittest.main()
