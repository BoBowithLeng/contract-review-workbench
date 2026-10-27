import base64
import http.client
import io
import json
import tempfile
import threading
import unittest
from pathlib import Path
from unittest.mock import Mock, patch

from openpyxl import load_workbook
from reader_server import create_server, ReaderHandler
from contract_facts import ContractFactsConflict
from review_jobs import JobConflict


ROOT = Path(__file__).resolve().parents[1]


class ChecklistPreparationRoutesTests(unittest.TestCase):
    def test_facts_confirmation_uses_saved_candidate_not_browser_text(self):
        loaded={'id':'detailed'}
        candidate={'id':'fact-candidate-'+'a'*32,'text':'真实讨论原话','source_quote':'真实讨论原话'}
        self.server.discussions.get=Mock(return_value={'turns':[{'id':'turn-one','status':'completed','fact_candidates':[candidate]}]})
        self.server.facts.confirm=Mock(return_value={'revision':1,'facts':[candidate]})
        with patch.object(ReaderHandler,'_report',return_value=loaded):
            body={'expected_revision':0,'finding_id':'F01','session_id':'initial','turn_id':'turn-one',
                  'candidate_id':candidate['id'],'text':'不能从浏览器伪造的正文'}
            code,_,raw=self.request('POST','/api/reports/detailed/facts/confirm',body)
            self.assertEqual(code,200,raw)
            args=self.server.facts.confirm.call_args.args
            self.assertEqual(args[2],candidate)
            self.assertNotIn('text',args[1])
            self.server.facts.confirm.side_effect=ContractFactsConflict('已变更',current_revision=2)
            self.assertEqual(self.request('POST','/api/reports/detailed/facts/confirm',body)[0],409)
            self.assertEqual(self.request('POST','/api/reports/detailed/facts/confirm',{**body,'candidate_id':'missing'})[0],400)

    def test_manual_facts_routes_preserve_payload_and_report_binding(self):
        loaded = {"id": "detailed"}
        with patch.object(ReaderHandler, "_report", return_value=loaded):
            for operation, body in (
                ("save", {"expected_revision": 0, "text": "手动事实", "replace_id": "fact-old"}),
                ("withdraw", {"expected_revision": 1, "fact_id": "fact-old"}),
                ("answer", {"expected_revision": 2, "question_id": "question-one", "answer": "", "status": "unknown"}),
            ):
                method = Mock(return_value={"revision": 3, "questions": [], "facts": [], "has_updates_since_review": True})
                setattr(self.server.facts, operation, method)
                code, _, raw = self.request("POST", "/api/reports/detailed/facts/" + operation, body)
                self.assertEqual(code, 200, raw)
                method.assert_called_once_with(loaded, body)
                method.side_effect = ContractFactsConflict("并发更新", current_revision=4)
                code, _, raw = self.request("POST", "/api/reports/detailed/facts/" + operation, body)
                self.assertEqual(code, 409, raw)
                self.assertEqual(json.loads(raw)["current_revision"], 4)

    def test_rereview_routes_and_explicit_origin(self):
        self.server.rereviews.preflight=Mock(return_value={'available':True,'input_revision':'abc'})
        self.server.rereviews.start=Mock(return_value=({'job_id':'job-'+'a'*32,'status':'running'},True))
        with patch.object(ReaderHandler,'_report',return_value={'id':'detailed'}), patch.object(ReaderHandler,'_state_for',return_value=(None,{'revision':0})):
            self.assertEqual(self.request('GET','/api/reports/detailed/rereview')[0],200)
            self.assertEqual(self.request('GET','/api/reports/detailed/rereview',origin='https://elsewhere.test')[0],403)
            self.assertEqual(self.request('POST','/api/reports/detailed/rereview',{'consent':True})[0],202)
            self.server.rereviews.start.side_effect=JobConflict('输入已变')
            self.assertEqual(self.request('POST','/api/reports/detailed/rereview',{'consent':True})[0],409)

    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.directory = Path(self.temp.name)
        self.server = create_server(root=ROOT, state_dir=self.directory/'state', jobs_dir=self.directory/'jobs',
                                    source_paths={'detailed':None,'minimal':None,'table':None},
                                    review_runner=lambda _messages,_config:'{}')
        self.thread=threading.Thread(target=self.server.serve_forever,daemon=True)
        self.thread.start()
        self.addCleanup(self.stop_server)
        self.origin=f'http://127.0.0.1:{self.server.server_port}'

    def stop_server(self):
        self.server.shutdown(); self.server.server_close(); self.thread.join(timeout=2)

    def request(self,method,path,data=None,origin=None):
        body=None if data is None else json.dumps(data).encode('utf-8')
        headers={'Origin':origin or self.origin}
        if body is not None: headers['Content-Type']='application/json'
        connection=http.client.HTTPConnection('127.0.0.1',self.server.server_port,timeout=5)
        try:
            connection.request(method,path,body,headers)
            response=connection.getresponse(); raw=response.read()
            return response.status,dict(response.getheaders()),raw
        finally: connection.close()

    def test_list_read_only_then_save_conflict_and_delete(self):
        code,_,raw=self.request('GET','/api/checklists')
        self.assertEqual(code,200)
        initial=json.loads(raw)
        self.assertFalse((self.directory/'state'/'checklists.json').exists())
        payload={'expected_revision':initial['revision'],'checklist':{
            'id':'','name':'测试清单','contract_type':'服务合同','stance':'客户',
            'items':[{'id':'R01','title':'续约','requirement':'检查自动续约。','enabled':True}]}}
        code,_,raw=self.request('POST','/api/checklists/save',payload)
        self.assertEqual(code,200,raw)
        saved=json.loads(raw)
        self.assertEqual(self.request('POST','/api/checklists/save',payload)[0],409)
        code,_,raw=self.request('POST','/api/checklists/delete',{'id':saved['checklist']['id'],'expected_revision':saved['revision']})
        self.assertEqual(code,200,raw)

    def test_template_is_xlsx_and_empty_import_explains_error(self):
        code,headers,raw=self.request('GET','/api/checklists/template.xlsx')
        self.assertEqual(code,200)
        self.assertIn('attachment',headers['Content-Disposition'])
        workbook=load_workbook(io.BytesIO(raw))
        self.assertIn('清单',workbook.sheetnames)
        self.assertEqual(workbook['清单'].max_row,1)
        code,_,result=self.request('POST','/api/checklists/import',{'filename':'template.xlsx','xlsx_base64':base64.b64encode(raw).decode()})
        self.assertIn(code,(200,400))
        self.assertEqual(json.loads(result)['rows'],[])

    def test_cross_origin_denied(self):
        self.assertEqual(self.request('GET','/api/checklists',origin='https://example.com')[0],403)
        self.assertEqual(self.request('POST','/api/checklists/save',{},origin='https://example.com')[0],403)

    def test_prepare_route_and_answer_conflict(self):
        job_id='job-'+'a'*32
        self.server.jobs.prepare=Mock(return_value=({'job_id':job_id,'status':'preparing'},True))
        code,_,raw=self.request('POST',f'/api/jobs/{job_id}/prepare',{'engine_key':'injected','consent':True})
        self.assertEqual(code,202,raw)
        self.server.jobs.prepare.assert_called_once()
        self.server.jobs.answer_preparation=Mock(side_effect=JobConflict('回答版本已变化，请刷新后核对。'))
        code,_,raw=self.request('POST',f'/api/jobs/{job_id}/answers',{'expected_revision':1,'answers':[]})
        self.assertEqual(code,409,raw)
        self.assertIn('版本已变化',json.loads(raw)['error'])

    def test_bad_base64_is_readable(self):
        code,_,raw=self.request('POST','/api/checklists/import',{'filename':'test.xlsx','xlsx_base64':'!invalid!'})
        self.assertEqual(code,400)
        self.assertIn('编码无效',json.loads(raw)['error'])


if __name__=='__main__': unittest.main()
