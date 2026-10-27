import copy
import io
import unittest
import zipfile

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from lxml import etree

from docx_delivery import ValidationError, inspect_source, render_docx, validate_plan


W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
W = f"{{{W_NS}}}"


def synthetic_docx() -> bytes:
    doc = Document()
    p = doc.add_paragraph()
    p.paragraph_format.keep_together = True
    p.add_run("前缀 ")
    bold = p.add_run("加粗旧")
    bold.bold = True
    italic = p.add_run("后缀 删除我 尾声")
    italic.italic = True
    p.add_run("\t换行前\n换行后")
    table = doc.add_table(rows=1, cols=1)
    table.cell(0, 0).paragraphs[0].add_run("表格旧条款")
    doc.core_properties.title = "synthetic-delivery"
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


def proofed_docx() -> bytes:
    doc = Document()
    first = doc.add_paragraph()
    first.paragraph_format.keep_together = True
    first.add_run("前文 ")
    marked = first.add_run("拼写目标")
    marked.bold = True
    first.add_run(" 后文")
    first._p.insert(first._p.index(marked._r), _proof_error("spellStart"))
    first._p.insert(first._p.index(marked._r) + 1, _proof_error("spellEnd"))

    second = doc.add_paragraph()
    second.add_run("保持 ")
    unchanged = second.add_run("原标记")
    unchanged.italic = True
    second._p.insert(second._p.index(unchanged._r), _proof_error("gramStart"))
    second._p.insert(second._p.index(unchanged._r) + 1, _proof_error("gramEnd"))
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


def _proof_error(value):
    node = OxmlElement("w:proofErr")
    node.set(qn("w:type"), value)
    return node


def piece(source, piece_id):
    return next(item for item in inspect_source(source)["pieces"] if item["id"] == piece_id)


def edit(source, edit_id, piece_id, old_text, new_text, *, occurrence=0, confirmed=True):
    item = piece(source, piece_id)
    starts = []
    cursor = 0
    while True:
        position = item["text"].find(old_text, cursor)
        if position < 0:
            break
        starts.append(position)
        cursor = position + max(1, len(old_text))
    start = starts[occurrence] if old_text else occurrence
    return {
        "id": edit_id,
        "piece_id": piece_id,
        "paragraph_sha256": item["paragraph_sha256"],
        "start": start,
        "end": start + len(old_text),
        "old_text": old_text,
        "new_text": new_text,
        "confirmed": confirmed,
    }


def plan(source, edits=None, comments=None, **updates):
    value = {
        "schema_version": 1,
        "source_sha256": inspect_source(source)["source_sha256"],
        "revision": 1,
        "author": "BoBo Review",
        "created_at": "2026-09-20T12:34:56Z",
        "edits": edits or [],
        "comments": comments or [],
    }
    value.update(updates)
    return value


def zip_parts(data):
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        return {name: archive.read(name) for name in archive.namelist()}


def document_root(data):
    return etree.fromstring(zip_parts(data)["word/document.xml"])


def paragraph_texts(data):
    root = document_root(data)
    result = []
    for p in root.iter(W + "p"):
        chunks = []
        for node in p.iter():
            if node.tag == W + "t":
                chunks.append(node.text or "")
            elif node.tag == W + "tab":
                chunks.append("\t")
            elif node.tag in {W + "br", W + "cr"}:
                chunks.append("\n")
        result.append("".join(chunks))
    return result


def project_revisions(data, *, accept):
    parts = zip_parts(data)
    root = etree.fromstring(parts["word/document.xml"])
    if not accept:
        for paragraph in list(root.iter(W + "p")):
            if paragraph.find(f"{W}pPr/{W}rPr/{W}ins") is not None:
                paragraph.getparent().remove(paragraph)
    remove_tag = W + ("del" if accept else "ins")
    unwrap_tag = W + ("ins" if accept else "del")
    for node in list(root.iter(remove_tag)):
        node.getparent().remove(node)
    for node in list(root.iter(unwrap_tag)):
        parent = node.getparent()
        index = parent.index(node)
        for child in list(node):
            if not accept:
                for deleted_text in child.iter(W + "delText"):
                    deleted_text.tag = W + "t"
            parent.insert(index, child)
            index += 1
        parent.remove(node)
    parts["word/document.xml"] = etree.tostring(
        root, xml_declaration=True, encoding="UTF-8", standalone=True
    )
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as archive:
        for name, payload in parts.items():
            archive.writestr(name, payload)
    return out.getvalue()


class DocxDeliveryTests(unittest.TestCase):
    def setUp(self):
        self.source = synthetic_docx()

    def test_inspect_source_includes_body_and_table_hashes(self):
        inspected = inspect_source(self.source)
        self.assertEqual(inspected["parser_version"], "docx-delivery/1")
        self.assertRegex(inspected["source_sha256"], r"^[0-9a-f]{64}$")
        ids = [item["id"] for item in inspected["pieces"]]
        self.assertIn("P000", ids)
        self.assertIn("T0R0C0P0", ids)
        self.assertTrue(all(len(item["paragraph_sha256"]) == 64 for item in inspected["pieces"]))
        self.assertTrue(piece(self.source, "P000")["supports_paragraph_insertion"])
        self.assertFalse(piece(self.source, "T0R0C0P0")["supports_paragraph_insertion"])

    def test_standalone_clause_addition_is_native_paragraph_and_rejects_to_original(self):
        anchor = piece(self.source, "P000")
        addition = edit(self.source, "ADD1", "P000", "", "8.2 指定邮箱发送成功视为送达。\n8.3 地址变更应通知。",
                        occurrence=len(anchor["text"]))
        addition["kind"] = "insert_paragraph_after"
        value = plan(self.source, [addition], [
            {"id": "C1", "text": "请确认新增的送达机制。", "confirmed": True, "edit_id": "ADD1"}
        ])
        self.assertEqual(validate_plan(self.source, value)["edits"][0]["kind"], "insert_paragraph_after")
        clean = render_docx(self.source, value, "clean")
        clean_root = document_root(clean)
        body_paragraphs = list(clean_root.find(W + "body").iterchildren(W + "p"))
        self.assertEqual([_xml_visible_text(p) for p in body_paragraphs][-2:],
                         ["8.2 指定邮箱发送成功视为送达。", "8.3 地址变更应通知。"])
        self.assertFalse(list(body_paragraphs[-2].iter(W + "br")))
        original_root = document_root(self.source)
        original_anchor = original_root.find(W + "body").findall(W + "p")[0]
        self.assertEqual(etree.tostring(body_paragraphs[0]), etree.tostring(original_anchor),
                         "inserting after an anchor must not flatten or rewrite its runs/proof markers")
        for variant in ("tracked", "tracked_comments"):
            tracked = render_docx(self.source, value, variant)
            root = document_root(tracked)
            inserted_paragraphs = [p for p in root.iter(W + "p") if p.find(f"{W}pPr/{W}rPr/{W}ins") is not None]
            self.assertEqual(len(inserted_paragraphs), 2)
            self.assertTrue(all(p.find(W + "ins") is not None for p in inserted_paragraphs))
            revision_ids = [node.get(W + "id") for node in root.iter(W + "ins")]
            self.assertEqual(len(revision_ids), len(set(revision_ids)))
            self.assertEqual(paragraph_texts(project_revisions(tracked, accept=True)), paragraph_texts(clean))
            self.assertEqual(paragraph_texts(project_revisions(tracked, accept=False)), paragraph_texts(self.source))
        self.assertEqual(paragraph_texts(render_docx(self.source, value, "comments")), paragraph_texts(self.source))
        for name, payload in zip_parts(self.source).items():
            if name != "word/document.xml":
                self.assertEqual(zip_parts(clean)[name], payload, name)

    def test_addition_validation_rejects_tables_stale_anchor_and_ambiguous_order(self):
        for piece_id in ("P000", "T0R0C0P0"):
            anchor = piece(self.source, piece_id)
            addition = edit(self.source, "ADD1", piece_id, "", "8.2 新条款。", occurrence=len(anchor["text"]))
            addition["kind"] = "insert_paragraph_after"
            if piece_id.startswith("T"):
                with self.assertRaises(ValidationError):
                    validate_plan(self.source, plan(self.source, [addition]))
                continue
            for updates in ({"paragraph_sha256": "0" * 64}, {"start": 0, "end": 0},
                            {"new_text": "8.2 新条款。\n\n8.3 另一条。"}, {"kind": "unbound_addition"}):
                with self.subTest(updates=updates), self.assertRaises(ValidationError):
                    validate_plan(self.source, plan(self.source, [dict(addition, **updates)]))
            with self.assertRaises(ValidationError):
                validate_plan(self.source, plan(self.source, [addition, dict(addition, id="ADD2")]))
            replacement = edit(self.source, "E1", piece_id, "加粗旧", "更新")
            self.assertEqual(len(validate_plan(self.source, plan(self.source, [replacement, addition]))["edits"]), 2)

    def test_clean_mixed_format_cross_run_and_table_edit(self):
        first = piece(self.source, "P000")["text"]
        cross = edit(self.source, "E1", "P000", "加粗旧后缀", "新\t值\n续 ")
        deletion = edit(self.source, "E2", "P000", "删除我", "")
        insertion_pos = first.index("尾声") + len("尾声")
        insertion = edit(self.source, "E3", "P000", "", "（新增）", occurrence=insertion_pos)
        table_edit = edit(self.source, "E4", "T0R0C0P0", "旧条款", "新条款")
        value = plan(self.source, [cross, deletion, insertion, table_edit])

        clean = render_docx(self.source, value, "clean")
        texts = paragraph_texts(clean)
        self.assertIn("前缀 新\t值\n续   尾声（新增）\t换行前\n换行后", texts)
        self.assertIn("表格新条款", texts)

        root = document_root(clean)
        inserted_run = next(
            run for run in root.iter(W + "r") if "新\t值\n续 " in _run_visible_text(run)
        )
        self.assertIsNotNone(inserted_run.find(f"{W}rPr/{W}b"))
        first_para = next(root.iter(W + "p"))
        self.assertIsNotNone(first_para.find(f"{W}pPr/{W}keepLines"))

    def test_tracked_accept_and_reject_project_to_clean_and_original(self):
        value = plan(
            self.source,
            [
                edit(self.source, "E1", "P000", "加粗旧后缀", "替换条款"),
                edit(self.source, "E2", "P000", "删除我", ""),
                edit(self.source, "E3", "T0R0C0P0", "", "新增：", occurrence=0),
            ],
        )
        clean = render_docx(self.source, value, "clean")
        tracked = render_docx(self.source, value, "tracked")
        root = document_root(tracked)
        self.assertTrue(list(root.iter(W + "ins")))
        self.assertTrue(list(root.iter(W + "del")))
        self.assertTrue(list(root.iter(W + "delText")))
        self.assertEqual(
            [item["text"] for item in inspect_source(project_revisions(tracked, accept=True))["pieces"]],
            [item["text"] for item in inspect_source(clean)["pieces"]],
        )
        self.assertEqual(
            [item["text"] for item in inspect_source(project_revisions(tracked, accept=False))["pieces"]],
            [item["text"] for item in inspect_source(self.source)["pieces"]],
        )

    def test_comments_and_tracked_comments_have_valid_references(self):
        replacement = edit(self.source, "E1", "P000", "加粗旧", "明确新条款")
        deletion = edit(self.source, "E2", "P000", "删除我", "")
        table_piece = piece(self.source, "T0R0C0P0")
        start = table_piece["text"].index("旧条款")
        comments = [
            {"id": "C1", "text": "对外批注一", "confirmed": True, "edit_id": "E1"},
            {"id": "C2", "text": "删除说明", "confirmed": True, "edit_id": "E2"},
            {
                "id": "C3",
                "text": "表格批注",
                "confirmed": True,
                "piece_id": "T0R0C0P0",
                "paragraph_sha256": table_piece["paragraph_sha256"],
                "start": start,
                "end": start + len("旧条款"),
                "old_text": "旧条款",
            },
        ]
        value = plan(self.source, [replacement, deletion], comments)
        for variant in ("comments", "tracked_comments"):
            output = render_docx(self.source, value, variant)
            parts = zip_parts(output)
            root = etree.fromstring(parts["word/document.xml"])
            comment_root = etree.fromstring(parts["word/comments.xml"])
            starts = [n.get(W + "id") for n in root.iter(W + "commentRangeStart")]
            ends = [n.get(W + "id") for n in root.iter(W + "commentRangeEnd")]
            refs = [n.get(W + "id") for n in root.iter(W + "commentReference")]
            bodies = [n.get(W + "id") for n in comment_root.iter(W + "comment")]
            self.assertEqual(sorted(starts), sorted(ends))
            self.assertEqual(sorted(starts), sorted(refs))
            self.assertEqual(sorted(starts), sorted(bodies))
            body_text = "".join(n.text or "" for n in comment_root.iter(W + "t"))
            self.assertEqual(body_text, "对外批注一删除说明表格批注")
        tracked_root = document_root(render_docx(self.source, value, "tracked_comments"))
        insertion = next(
            node for node in tracked_root.iter(W + "ins") if "明确新条款" in _xml_visible_text(node)
        )
        siblings = list(insertion.getparent())
        idx = siblings.index(insertion)
        self.assertEqual(siblings[idx - 1].tag, W + "commentRangeStart")
        self.assertEqual(siblings[idx + 1].tag, W + "commentRangeEnd")

    def test_unrelated_zip_parts_are_byte_identical(self):
        replacement = edit(self.source, "E1", "P000", "加粗旧", "新条款")
        value = plan(
            self.source,
            [replacement],
            [{"id": "C1", "text": "公开意见", "confirmed": True, "edit_id": "E1"}],
        )
        original = zip_parts(self.source)
        clean = zip_parts(render_docx(self.source, value, "clean"))
        for name, payload in original.items():
            if name != "word/document.xml":
                self.assertEqual(clean[name], payload, name)
        commented = zip_parts(render_docx(self.source, value, "tracked_comments"))
        allowed = {
            "word/document.xml",
            "word/comments.xml",
            "word/_rels/document.xml.rels",
            "[Content_Types].xml",
        }
        for name, payload in original.items():
            if name not in allowed:
                self.assertEqual(commented[name], payload, name)

    def test_source_hash_paragraph_hash_and_offsets_must_match(self):
        base = edit(self.source, "E1", "P000", "加粗旧", "新")
        cases = []
        wrong_source = plan(self.source, [base])
        wrong_source["source_sha256"] = "0" * 64
        cases.append(wrong_source)
        wrong_paragraph = plan(self.source, [copy.deepcopy(base)])
        wrong_paragraph["edits"][0]["paragraph_sha256"] = "0" * 64
        cases.append(wrong_paragraph)
        wrong_old = plan(self.source, [copy.deepcopy(base)])
        wrong_old["edits"][0]["old_text"] = "错字"
        cases.append(wrong_old)
        for value in cases:
            with self.subTest(value=value):
                with self.assertRaises(ValidationError):
                    validate_plan(self.source, value)

    def test_overlap_same_point_insert_and_unconfirmed_are_rejected(self):
        first = edit(self.source, "E1", "P000", "加粗旧", "A")
        text = piece(self.source, "P000")["text"]
        overlap = edit(self.source, "E2", "P000", "粗旧后", "B")
        insert_at_start = edit(
            self.source, "E3", "P000", "", "C", occurrence=text.index("加粗旧")
        )
        for edits in ([first, overlap], [first, insert_at_start]):
            with self.assertRaisesRegex(ValidationError, "重叠|含糊"):
                validate_plan(self.source, plan(self.source, edits))
        one = edit(self.source, "E4", "P000", "", "A", occurrence=0)
        two = edit(self.source, "E5", "P000", "", "B", occurrence=0)
        with self.assertRaisesRegex(ValidationError, "重叠|含糊"):
            validate_plan(self.source, plan(self.source, [one, two]))
        first["confirmed"] = False
        with self.assertRaisesRegex(ValidationError, "未被明确确认"):
            validate_plan(self.source, plan(self.source, [first]))

    def test_unconfirmed_comment_and_unknown_internal_field_are_rejected(self):
        replacement = edit(self.source, "E1", "P000", "加粗旧", "新")
        comment = {"id": "C1", "text": "公开", "confirmed": False, "edit_id": "E1"}
        with self.assertRaisesRegex(ValidationError, "未被明确确认"):
            validate_plan(self.source, plan(self.source, [replacement], [comment]))
        value = plan(self.source, [replacement])
        value["internal_notes"] = "绝不能进批注"
        with self.assertRaisesRegex(ValidationError, "未知字段"):
            validate_plan(self.source, value)

        non_integer_schema = plan(self.source, [replacement])
        non_integer_schema["schema_version"] = 1.0
        with self.assertRaisesRegex(ValidationError, "schema_version"):
            validate_plan(self.source, non_integer_schema)

    def test_complex_target_paragraph_is_rejected(self):
        doc = Document(io.BytesIO(self.source))
        paragraph = doc.paragraphs[0]
        hyperlink = OxmlElement("w:hyperlink")
        hyperlink.set(qn("w:anchor"), "target")
        run = OxmlElement("w:r")
        text_node = OxmlElement("w:t")
        text_node.text = "复杂"
        run.append(text_node)
        hyperlink.append(run)
        paragraph._p.append(hyperlink)
        out = io.BytesIO()
        doc.save(out)
        source = out.getvalue()
        target = edit(source, "E1", "P000", "前缀", "替换")
        with self.assertRaisesRegex(ValidationError, "复杂段落节点 hyperlink"):
            validate_plan(source, plan(source, [target]))

    def test_text_edit_drops_stale_proof_errors_only_in_changed_paragraph(self):
        source = proofed_docx()
        target = edit(source, "E1", "P000", "拼写目标", "明确条款")
        output = render_docx(source, plan(source, [target]), "clean")
        before = list(document_root(source).iter(W + "p"))
        after = list(document_root(output).iter(W + "p"))

        self.assertEqual([node.get(W + "type") for node in before[0].iter(W + "proofErr")],
                         ["spellStart", "spellEnd"])
        self.assertFalse(list(after[0].iter(W + "proofErr")))
        self.assertEqual(
            [node.get(W + "type") for node in after[1].iter(W + "proofErr")],
            ["gramStart", "gramEnd"],
        )
        self.assertEqual(etree.tostring(before[1]), etree.tostring(after[1]))
        inserted = next(run for run in after[0].iter(W + "r") if _run_visible_text(run) == "明确条款")
        self.assertIsNotNone(inserted.find(f"{W}rPr/{W}b"))
        self.assertIsNotNone(after[0].find(f"{W}pPr/{W}keepLines"))

    def test_comments_only_preserve_proof_errors_and_original_formatting(self):
        source = proofed_docx()
        source_piece = piece(source, "P000")
        start = source_piece["text"].index("拼写目标")
        comment = {
            "id": "C1",
            "text": "公开批注",
            "confirmed": True,
            "piece_id": "P000",
            "paragraph_sha256": source_piece["paragraph_sha256"],
            "start": start,
            "end": start + len("拼写目标"),
            "old_text": "拼写目标",
        }
        output = render_docx(source, plan(source, comments=[comment]), "comments")
        before = next(document_root(source).iter(W + "p"))
        after = next(document_root(output).iter(W + "p"))

        self.assertEqual(
            [node.get(W + "type") for node in after.iter(W + "proofErr")],
            ["spellStart", "spellEnd"],
        )
        self.assertEqual(
            etree.tostring(before.find(W + "pPr")), etree.tostring(after.find(W + "pPr"))
        )
        original_bold = next(run for run in before.iter(W + "r") if _run_visible_text(run) == "拼写目标")
        output_bold = next(run for run in after.iter(W + "r") if _run_visible_text(run) == "拼写目标")
        self.assertEqual(
            etree.tostring(original_bold.find(W + "rPr")),
            etree.tostring(output_bold.find(W + "rPr")),
        )
        self.assertEqual(paragraph_texts(source), paragraph_texts(output))


def _run_visible_text(run):
    return _xml_visible_text(run)


def _xml_visible_text(element):
    result = []
    for node in element.iter():
        if node.tag == W + "t":
            result.append(node.text or "")
        elif node.tag == W + "tab":
            result.append("\t")
        elif node.tag in {W + "br", W + "cr"}:
            result.append("\n")
    return "".join(result)


if __name__ == "__main__":
    unittest.main()
