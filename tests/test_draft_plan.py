"""Deterministic protocol tests; all contracts and model responses are synthetic."""
import copy
import io
import json
import unittest

from docx import Document
from docx.oxml import OxmlElement
from contract_core import parse_docx, parse_text
from draft_plan import insertion_capabilities, validate_draft_plan, render_clause_draft, validate_docx_plan
from general_review import build_messages, parse_response, response_schema
from review_runner import output_schema
from tests.test_general_review import context, finding, payload


def plan(status="ready", *, edit=True, insert=False):
    return {"status": status, "edits": [{"piece_id": "P000", "new_text": "1.1 押金按约定返还。"}] if edit else [],
            "insertions": [{"after_piece_id": "P001", "new_text": "1.3 新增独立结算机制。"}] if insert else [],
            "paragraph_actions": [
                {"piece_id": "P000", "action": "modify" if edit else "retain", "reason": "明确押金结算后返还的安排。"},
                {"piece_id": "P001", "action": "retain", "reason": "此处仅为相关背景，原有租期不变。"}],
            "questions": ["押金返还期限希望定为多少日？"] if status == "needs_input" else [],
            "comment": "请核对押金返还安排，租期相关条款保留。"}


class DraftPlanTests(unittest.TestCase):
    def setUp(self):
        self.pieces = {"P000": "1.1 押金不退。", "P001": "1.2 租期一年。"}
        self.refs = [{"id": pid, "quote": text} for pid, text in self.pieces.items()]

    def validate(self, value):
        return validate_draft_plan(value, pieces=self.pieces, refs=self.refs, allowed_insertion_ids=["P001"])

    def test_ready_modifies_only_one_related_reference_and_renders_exact_text(self):
        result = self.validate(plan(insert=True))
        self.assertEqual([r["action"] for r in result["paragraph_actions"]], ["modify", "retain"])
        self.assertEqual(render_clause_draft(result), "1.1 押金按约定返还。\n1.3 新增独立结算机制。")

    def test_needs_input_and_not_needed_never_render_drafts(self):
        for status in ("needs_input", "not_needed"):
            value = plan(status, edit=False)
            self.assertEqual(self.validate(value)["status"], status)
            self.assertEqual(render_clause_draft(value), "")
        value = plan("needs_input", edit=False)
        value["questions"] = []
        with self.assertRaises(ValueError):
            self.validate(value)

    def test_unlocated_missing_clause_can_insert_at_legal_anchor(self):
        value = plan(edit=False, insert=True)
        value["paragraph_actions"] = []
        self.assertEqual(validate_draft_plan(value, pieces=self.pieces, refs=[], allowed_insertion_ids=["P001"])["status"], "ready")

    def test_malformed_or_unsafe_plans_reject_instead_of_becoming_needs_input(self):
        cases = [None, {}, {**plan(), "status": []}, {**plan(), "paragraph_actions": []},
                 {**plan(), "edits": []}, {**plan(), "questions": ["尚未明确"]},
                 {**plan(), "edits": [{"piece_id": "P000", "new_text": self.pieces["P000"]}]},
                 {**plan(), "edits": [{"piece_id": "P000", "new_text": "2.1 押金返还。"}]},
                 {**plan(), "edits": [{"piece_id": "P000", "new_text": "1.1 修改为：押金返还。"}]},
                 {**plan(), "edits": [{"piece_id": "P000", "new_text": "1.1 押金返还。\n\n说明。"}]},
                 {**plan(), "edits": [{"piece_id": "outside", "new_text": "押金返还。"}]},
                 {**plan(), "insertions": [{"after_piece_id": "P000", "new_text": "新增条款。"}]},
                 {**plan(), "insertions": [{"after_piece_id": "P001", "new_text": "1.2 重复条号。"}]},
                 {**plan(), "paragraph_actions": [*plan()["paragraph_actions"], plan()["paragraph_actions"][0]]},
                 plan("not_needed"), plan("needs_input")]
        for value in cases:
            with self.subTest(value=value), self.assertRaises(ValueError):
                self.validate(value)

    def test_partial_old_reference_cannot_broaden_edit(self):
        refs = [{"id": "P000", "quote": "押金不退"}, self.refs[1]]
        valid = plan()
        valid["edits"][0]["new_text"] = "1.1 押金应返还。"
        validate_draft_plan(valid, pieces=self.pieces, refs=refs, allowed_insertion_ids=[])
        valid["edits"][0]["new_text"] = "1.1 押金应返还；同时改变句尾"
        with self.assertRaises(ValueError):
            validate_draft_plan(valid, pieces=self.pieces, refs=refs, allowed_insertion_ids=[])

    def test_actual_docx_insertion_excludes_table_and_complex_paragraph(self):
        document = Document()
        document.add_paragraph("正文可新增。")
        document.add_table(rows=1, cols=1).cell(0, 0).text = "表格只作修改。"
        stream = io.BytesIO()
        document.save(stream)
        parsed = parse_docx(stream.getvalue())
        anchors = insertion_capabilities(parsed, stream.getvalue())
        self.assertEqual([a["text"] for a in anchors], ["正文可新增。"])
        self.assertEqual(insertion_capabilities(parse_text("纯文本")), [])

    def test_ready_is_rejected_before_publication_when_docx_paragraph_is_complex(self):
        document = Document()
        paragraph = document.add_paragraph("1.1 押金不退。")
        paragraph._p.append(OxmlElement("w:bookmarkStart"))
        document.add_paragraph("1.2 租期一年。")
        stream = io.BytesIO()
        document.save(stream)
        data = stream.getvalue()
        self.assertEqual([a["piece_id"] for a in insertion_capabilities(parse_docx(data), data)], ["P001"])
        with self.assertRaisesRegex(ValueError, "复杂段落"):
            validate_docx_plan(self.validate(plan()), data)
        row = finding(refs=[{"id": "P000"}, {"id": "P001"}])
        row.pop("clause_draft")
        row.pop("evidence_status")
        row["draft_plan"] = plan()
        result = parse_response(json.dumps(payload([row])), parse_docx(data), context(draft_plan_version=1), [],
                                draft_plan_source_bytes=data)
        self.assertEqual(result["findings"][0]["status"], "处理失败")
        self.assertNotIn("draft_plan", result["findings"][0])

    def test_new_and_old_schemas_and_model_input_are_distinct(self):
        new = response_schema(draft_plan_version=1)["properties"]["findings"]["items"]
        old = response_schema()["properties"]["findings"]["items"]
        self.assertIn("draft_plan", new["required"])
        self.assertNotIn("clause_draft", new["properties"])
        self.assertIn("clause_draft", old["required"])
        self.assertIn("clause_draft", response_schema(draft_plan_version=1)["properties"]["options"]["items"]["properties"])
        for phase in ("review", "post_review_information"):
            self.assertEqual(output_schema("general", phase=phase, draft_plan_version=1)["properties"]["findings"]["items"], new)
        ctx = context(draft_plan_version=1, draft_plan_insertion_anchors=[{"piece_id": "P001", "text": self.pieces["P001"]}])
        inputs = json.loads(build_messages(parse_text("\n".join(self.pieces.values())), ctx, [])[1]["content"])
        self.assertEqual(inputs["context"]["draft_plan_insertion_anchors"], ctx["draft_plan_insertion_anchors"])
        self.assertNotIn("draft_plan_version", json.loads(build_messages(parse_text("旧协议原文"), context(), [])[1]["content"])["context"])

    def test_missing_plan_is_processing_failure_but_business_questions_are_not(self):
        parsed = parse_text("\n".join(self.pieces.values()))
        ctx = context(draft_plan_version=1, draft_plan_insertion_anchors=[])
        raw_finding = finding(refs=[{"id": "P000"}, {"id": "P001"}])
        raw_finding.pop("clause_draft")
        raw_finding.pop("evidence_status")
        missing = parse_response(json.dumps(payload([raw_finding])), parsed, ctx, [])
        self.assertEqual(missing["findings"][0]["status"], "处理失败")
        self.assertNotIn("draft_plan", missing["findings"][0])
        raw_finding["draft_plan"] = plan("needs_input", edit=False)
        result = parse_response(json.dumps(payload([raw_finding])), parsed, ctx, [])
        self.assertNotEqual(result["findings"][0]["status"], "处理失败")
        self.assertEqual(result["findings"][0]["clause_draft"], "")
        self.assertEqual(result["findings"][0]["questions"], raw_finding["draft_plan"]["questions"])


class BlankLineParagraphTests(unittest.TestCase):
    """A blank line between proposed clauses separates paragraphs from draft text version 1 on."""
    def setUp(self):
        self.pieces = {"P000": "1.1 押金不退。", "P001": "1.2 租期一年。"}
        self.parsed = parse_text("\n".join(self.pieces.values()))
        row = finding(refs=[{"id": "P000"}, {"id": "P001"}])
        row.pop("clause_draft")
        row.pop("evidence_status")
        row["draft_plan"] = plan(insert=True)
        row["draft_plan"]["insertions"][0]["new_text"] = "1.3 押金结算\n\n（一）退租交接后核算费用。\n\n（二）七日内返还余额。"
        self.row = row

    def parse(self, **updates):
        ctx = context(draft_plan_version=1, draft_plan_insertion_anchors=[{"piece_id": "P001", "text": self.pieces["P001"]}], **updates)
        return parse_response(json.dumps(payload([self.row])), self.parsed, ctx, [])["findings"][0]

    def test_versioned_task_reads_blank_lines_as_paragraph_breaks(self):
        result = self.parse(draft_text_version=1)
        self.assertNotEqual(result["status"], "处理失败")
        self.assertEqual(result["draft_plan"]["insertions"][0]["new_text"],
                         "1.3 押金结算\n（一）退租交接后核算费用。\n（二）七日内返还余额。")
        self.assertEqual(result["draft_plan"]["edits"][0]["new_text"], "1.1 押金按约定返还。")

    def test_version_two_also_reads_blank_lines_as_paragraph_breaks(self):
        result = self.parse(draft_text_version=2)
        self.assertNotEqual(result["status"], "处理失败")
        self.assertEqual(result["draft_plan"]["insertions"][0]["new_text"],
                         "1.3 押金结算\n（一）退租交接后核算费用。\n（二）七日内返还余额。")

    def test_version_two_drops_a_blank_stray_field_but_not_one_with_content(self):
        self.row["draft_plan"]["paragraph_actions"][0]["reason_note"] = ""
        self.row["draft_plan"]["edits"][0]["note"] = None
        kept = self.parse(draft_text_version=2)
        self.assertNotEqual(kept["status"], "处理失败")
        self.assertNotIn("reason_note", kept["draft_plan"]["paragraph_actions"][0])
        # Saved version-1 tasks keep failing exactly as they did.
        self.assertEqual(self.parse(draft_text_version=1)["status"], "处理失败")
        self.row["draft_plan"]["paragraph_actions"][0]["reason_note"] = "另需核对附件"
        self.assertEqual(self.parse(draft_text_version=2)["status"], "处理失败")

    def test_only_version_two_asks_comments_to_avoid_internal_ids(self):
        from draft_plan import COMMENT_RULE
        anchors = [{"piece_id": "P001", "text": self.pieces["P001"]}]
        def system(**updates):
            ctx = context(draft_plan_version=1, draft_plan_insertion_anchors=anchors, **updates)
            return build_messages(self.parsed, ctx, [])[0]["content"]
        self.assertIn(COMMENT_RULE, system(draft_text_version=2))
        # Saved tasks re-derive with the prompt they were sent; their text must not change.
        self.assertNotIn(COMMENT_RULE, system(draft_text_version=1))
        self.assertNotIn(COMMENT_RULE, system())
        self.assertNotIn(COMMENT_RULE, system(draft_text_version=99))

    def test_version_three_adds_clause_number_cross_references(self):
        from draft_plan import COMMENT_RULE, CROSS_REFERENCE_RULE
        anchors = [{"piece_id": "P001", "text": self.pieces["P001"]}]
        def system(**updates):
            ctx = context(draft_plan_version=1, draft_plan_insertion_anchors=anchors, **updates)
            return build_messages(self.parsed, ctx, [])[0]["content"]
        self.assertIn(CROSS_REFERENCE_RULE, system(draft_text_version=3))
        self.assertIn(COMMENT_RULE, system(draft_text_version=3))
        # Tasks frozen at 2 keep the exact prompt they were sent.
        self.assertNotIn(CROSS_REFERENCE_RULE, system(draft_text_version=2))

    def test_task_frozen_before_the_version_keeps_its_original_strict_reading(self):
        result = self.parse()
        self.assertEqual(result["status"], "处理失败")
        self.assertNotIn("draft_plan", result)
