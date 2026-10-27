import io
import unittest
from docx import Document
from contract_core import parse_text, parse_docx
from document_view import index_document, reader_markup, docx_headings, sort_findings_by_document
from demo_contracts import MINIMAL_TEXT, table_docx
from report_export import export_docx
from completeness import check_completeness


class DocumentTests(unittest.TestCase):
    def test_article_heading_and_decimal_clause_location(self):
        c=parse_text('第三条 价款与支付\n3.2 '+'付款安排'*20+'\n第四条 交付与验收\n4.1 '+'验收约定'*20)
        index=index_document(c)
        self.assertEqual([p['level'] for p in index], [0, 0, 0, 0])
        self.assertEqual(index[1]['location'],'第三条 / 3.2')
        self.assertEqual(index[3]['location'],'第四条 / 4.1')

    def test_sections_follow_original_text(self):
        c=parse_text('六、劳动报酬\n（五）'+'原文工资约定'*15+'\n十、劳动合同的变更、解除、终止\n（一）'+'离职条款'*20)
        index=index_document(c)
        self.assertEqual(index[1]['location'],'六、劳动报酬 / （五）')
        self.assertEqual(index[3]['location'],'十、劳动合同的变更、解除、终止 / （一）')

    def test_no_headings_and_table_fallback(self):
        c=parse_text('这是没有标题的正文\n另一段')
        self.assertEqual(index_document(c)[1]['location'],'正文第2段')
        t=parse_docx(table_docx())
        self.assertTrue(any('第1行第1列' in p['location'] for p in index_document(t)))

    def test_exact_highlight_and_escape(self):
        c=parse_text('原文<svg onload=evil()>不支付工资。')
        items=[{'id':'F1','refs':[{'id':'P000','quote':'不支付工资'}]}]
        result=reader_markup(c,items,'F1')
        self.assertIn('<mark>不支付工资</mark>',result['html'])
        self.assertNotIn('<svg',result['html'])
        self.assertEqual(result['target'],'p-P000')
        self.assertIn('data-issue="F1"',result['html'])

    def test_word_report_has_locations_quotes_options_and_human_decision(self):
        c=parse_text('六、劳动报酬\n（五）公司暂不支付工资。')
        r={'mode':'ai','context':{'region':'上海','review_date':'2026-09-20'},'findings':[
            {'id':'F1','rule_id':'LC-02','title':'工资条件','status':'发现风险','risk':'高','refs':[{'id':'P001','quote':'公司暂不支付工资。'}],
             'reason':'测试意见','source_ids':[],'clause_draft':'测试条款'}],
             'options':[{'id':'O1','title':'方案甲','finding_ids':['F1'],'preconditions':['条件甲'],'steps':['步骤甲'],'cost':'人力','remaining_risk':'需确认','relationship':'可并用','clause_draft':'方案稿'}]}
        b=export_docx(r,[],c,{'F1':{'decision':'修改','risk':'中','note':'已核','clause_draft':'人工条款'}},['O1'])
        doc=Document(io.BytesIO(b)); text='\n'.join(p.text for p in doc.paragraphs)
        for needle in ('六、劳动报酬','公司暂不支付工资。','方案甲（已选）','人工条款'):
            self.assertIn(needle,text)

    def test_two_synthetic_formats_read_and_check(self):
        for c in (parse_text(MINIMAL_TEXT),parse_docx(table_docx())):
            basic=check_completeness(c,{'input_kind':'模板审阅'})
            self.assertEqual(len(basic),8)
            self.assertTrue(all(p['quote'] in next(x.text for x in c.pieces if x.id==p['id']) for item in basic for p in item['refs']))

    def test_word_heading_without_number_is_used(self):
        data=table_docx(); c=parse_docx(data).model_dump();c['headings']=docx_headings(data)
        self.assertTrue(any(p['text']=='离职安排' and p['level']==1 for p in index_document(c)))

    def test_explicit_word_headings_only_and_depth_preserved(self):
        doc=Document()
        doc.add_paragraph('总则', style='Heading 1')
        doc.add_paragraph('第一节 适用范围', style='Heading 2')
        doc.add_paragraph('材料说明', style='Heading 3')
        doc.add_paragraph('1.甲方为依法登记的企业，具有签约资格。')
        doc.add_paragraph('5.乙方不存在劳动争议。')
        doc.add_paragraph('（一）经双方提出并同意后，相关约定继续履行。')
        buf=io.BytesIO();doc.save(buf);data=buf.getvalue()
        contract=parse_docx(data).model_dump()
        contract['headings']=docx_headings(data)
        indexed=index_document(contract)
        self.assertEqual([p['level'] for p in indexed], [1,2,3,0,0,0])
        self.assertEqual(indexed[3]['location'], '总则 / 第一节 适用范围 / 材料说明 / 1.')
        self.assertEqual([s['text'] for s in reader_markup(contract,[],None)['sections']],
                         ['总则','第一节 适用范围','材料说明'])

    def test_short_numbered_sentences_and_list_items_are_not_navigation(self):
        c=parse_text('1.甲方为依法登记的企业。\n5.乙方不存在劳动争议。\n（一）经双方提出并同意后解除。\n第十条 甲方应按时付款。')
        indexed=index_document(c)
        self.assertEqual([p['level'] for p in indexed], [0,0,0,0])
        self.assertEqual(reader_markup(c,[],None)['sections'], [])
        self.assertEqual(indexed[-1]['location'],'第十条')

    def test_only_standalone_chapter_or_section_is_inferred(self):
        c=parse_text('第一章 总则\n甲方应履行义务。\n第一节 付款安排\n1.付款应于验收后完成。\n第十条 付款期限')
        indexed=index_document(c)
        self.assertEqual([p['level'] for p in indexed], [1,0,2,0,0])
        self.assertEqual(indexed[3]['location'],'第一章 总则 / 第一节 付款安排 / 1.')

    def test_sort_and_export_keep_original_ids_and_quotes(self):
        c=parse_text('第一章 总则\n第十条 付款期限\n甲方于验收后付款。\n第二章 交付\n乙方于七日内交付。')
        findings=[{'id':'F02','refs':[{'id':'P004','quote':'七日内'}]},
                  {'id':'F01','refs':[{'id':'P002','quote':'验收后'}]}]
        ordered=sort_findings_by_document(findings,c)
        self.assertEqual([f['id'] for f in ordered], ['F01','F02'])
        self.assertEqual(c.pieces[2].text,'甲方于验收后付款。')
        self.assertIn('第二章 交付', index_document(c)[4]['location'])


if __name__=='__main__':unittest.main()
