"""Real local HTTP workflow with disposable DOCX and injected model replies."""
import base64
import copy
import http.client
import io
import json
import tempfile
import threading
import time
import unittest
import uuid
from pathlib import Path

from docx import Document
from reader_server import create_server


ROOT = Path(__file__).resolve().parents[1]


class ChecklistWorkspaceHTTPTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.base = Path(self.temp.name)
        self.calls = []
        self.omit_group = False

        def runner(messages, config):
            self.calls.append(copy.deepcopy(config))
            if config.get('review_phase') == 'review_preparation':
                return json.dumps({'questions': [], 'requirements': [], 'additional_checks': []})
            prep = json.loads(messages[1]['content'])['review_preparation']
            rows = []
            for check in prep['checks']:
                groups = [{'group_id': g['group_id'], 'title': g['title'], 'status': '未检出',
                           'reason': '注入回复仅验证流程。', 'finding_ids': []} for g in check['groups']]
                if self.omit_group:
                    groups = groups[:-1]
                rows.append({'check_id': check['id'], 'title': check['title'], 'status': '未检出',
                             'reason': '注入回复仅验证流程。', 'finding_ids': [], 'groups': groups})
            return json.dumps({'findings': [], 'checks': [], 'options': [], 'uncovered': [],
                               'limitations': [], 'checklist_results': rows}, ensure_ascii=False)

        self.server = create_server(root=ROOT, state_dir=self.base / 'state', jobs_dir=self.base / 'jobs',
                                    review_runner=runner,
                                    review_engines=[{'key': 'injected', 'available': True, 'label': '测试桩'}])
        self.thread = threading.Thread(target=self.server.serve_forever, daemon=True)
        self.thread.start()
        self.addCleanup(self.stop)
        doc = Document()
        doc.add_heading('合成采购合同', level=1)
        doc.add_heading('付款与交付', level=2)
        doc.add_paragraph('甲方在验收后支付货款。乙方应按约定交付设备。')
        stream = io.BytesIO()
        doc.save(stream)
        self.source = stream.getvalue()
        self.payload = {'engine_key': 'injected', 'config': {}, 'consent': True}

    def stop(self):
        self.server.shutdown()
        self.server.server_close()
        self.thread.join(3)

    def request(self, path, body=None):
        conn = http.client.HTTPConnection('127.0.0.1', self.server.server_port, timeout=10)
        headers = {'Origin': f'http://127.0.0.1:{self.server.server_port}', 'Content-Type': 'application/json'}
        conn.request('GET' if body is None else 'POST', path,
                     None if body is None else json.dumps(body).encode(), headers)
        response = conn.getresponse()
        result = response.status, json.loads(response.read())
        conn.close()
        self.assertLess(result[0], 400, result)
        return result[1]

    def create(self):
        return self.request('/api/jobs', {
            'document_name': '工作区合成验收.docx', 'docx_base64': base64.b64encode(self.source).decode(),
            'context': {'contract_type': 'general', 'our_party': '甲方', 'our_role': '采购方',
                        'region': '中国大陆', 'review_date': '2026-10-04', 'input_kind': '拟签审阅'},
            'checklist_snapshot': {'id': 'fixture-checklist', 'name': '合成检查', 'version': 1,
                'builtin': False, 'contract_type': '采购', 'stance': '甲方', 'items': [
                    {'id': 'C01', 'title': '付款与交付', 'requirement':
                     '### 付款安排\n- 检查付款条件。\n\n### 交付安排\n- 检查交付期限。',
                     'when': '', 'notes': '', 'enabled': True}]}})['job']['job_id']

    def wait_status(self, jid, expected):
        deadline = time.monotonic() + 6
        while time.monotonic() < deadline:
            job = self.request('/api/jobs/' + jid)['job']
            if job['status'] == expected:
                return job
            if job['status'] == 'failed' and expected != 'failed':
                self.fail(str(job.get('error')))
            time.sleep(.02)
        self.fail('Timed out waiting for ' + expected)

    def run_initial(self, expected='completed'):
        jid = self.create()
        self.request('/api/jobs/' + jid + '/prepare', self.payload)
        ready = self.wait_status(jid, 'prepared')
        self.assertEqual(ready['preparation']['checklist_group_protocol_version'], 1)
        self.assertEqual(len(ready['preparation']['checks'][0]['groups']), 2)
        self.request('/api/jobs/' + jid + '/start', self.payload)
        self.wait_status(jid, expected)
        return jid, self.request('/api/reports/' + jid)

    def test_initial_frozen_groups_and_rereview_through_http(self):
        jid, report = self.run_initial()
        groups = report['report']['checklist_results'][0]['groups']
        self.assertEqual([g['title'] for g in groups], ['付款安排', '交付安排'])
        self.assertTrue(all(g['status'] == '未检出' for g in groups))
        before = (self.base / 'jobs' / jid / 'result.json').read_bytes()
        preflight = self.request('/api/reports/' + jid + '/rereview')
        self.assertTrue(preflight['available'])
        child = self.request('/api/reports/' + jid + '/rereview', {
            **self.payload, 'input_revision': preflight['input_revision'], 'request_id': str(uuid.uuid4())})['job']['job_id']
        self.wait_status(child, 'completed')
        rereview = self.request('/api/reports/' + child)
        self.assertEqual(rereview['report']['preparation']['checks'][0]['groups'],
                         report['report']['preparation']['checks'][0]['groups'])
        self.assertEqual(len(rereview['report']['checklist_results'][0]['groups']), 2)
        self.assertEqual(before, (self.base / 'jobs' / jid / 'result.json').read_bytes())
        self.assertEqual(len(self.calls), 3)

    def test_partial_group_report_is_readable_and_no_automatic_retry(self):
        self.omit_group = True
        jid, report = self.run_initial('failed')
        groups = report['report']['checklist_results'][0]['groups']
        self.assertEqual(len(groups), 2)
        self.assertEqual(groups[1]['status'], '处理失败')
        self.assertNotEqual(groups[1]['status'], '未检出')
        self.request('/api/reports/' + jid)
        self.assertEqual(len(self.calls), 2)


if __name__ == '__main__':
    unittest.main()
