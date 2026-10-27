"""HTTP and Word integration. Model responses here are deterministic test fixtures."""
from __future__ import annotations

import copy
import io
import json
import time
import unittest
import uuid

from docx import Document
from tests import test_reader_server as reader_fixtures


class OpinionRevisionRouteTests(unittest.TestCase):
    def setUp(self):
        self.fixture = reader_fixtures.ReaderServerTests("test_reads_only_three_fixed_reports")
        self.fixture.setUp()
        self.addCleanup(self.fixture.doCleanups)
        self.api = self.fixture.json_request
        self.path = "/api/reports/detailed/discussions/F01"
        self.calls = []
        self.action = "revise"
        self.config = {"base_url": "https://model.example/v1", "api_key": "fixture-key", "model": "fixture"}
        self.fixture.server.discussions.completion = lambda config, messages: "根据补充事实，工资支付与交接责任分别约定。"
        self.fixture.server.revisions.completion = self.complete
        self.fixture.server.delivery.source_paths["detailed"] = None
        self.original = self.api("GET", "/api/reports/detailed")[2]
        self.original_finding = next(f for f in self.original["report"]["findings"] if f["id"] == "F01")
        self.raw_path = self.fixture.root / "runtime" / "sol-detailed-run-20260920" / "response.json"
        self.raw_before = self.raw_path.read_bytes()

    def complete(self, config, messages):
        self.calls.append(copy.deepcopy(messages))
        finding = copy.deepcopy(self.original_finding)
        finding.update({
            "title": "讨论后：分别约定工资支付与交接责任",
            "reason": "讨论确认需分别说明支付节点和交接损失；这是集成测试中的新版意见。",
            "suggestion": "分别写清两项安排。", "clause_draft": "公司按约定支付工资，工作交接责任另列。",
        })
        if self.action == "withdraw":
            finding.update(status="未检出", risk=None, clause_draft="", suggestion="", reason="补充事实后撤回此项提示。")
        packet = json.loads(messages[1]["content"])
        if packet.get("draft_plan_version") == 1:
            finding.pop("clause_draft", None)
            finding["draft_plan"] = {"status": "not_needed", "edits": [], "insertions": [], "questions": [],
                "comment": "本轮只更新内部意见，现合同条款保持不变。", "paragraph_actions": [
                    {"piece_id": ref["id"], "action": "retain", "reason": "本轮只调整内部判断，保留相关原文。"} for ref in finding["refs"]]}
        return json.dumps({"action": self.action, "finding": finding,
                           "change_summary": "依据刚才的补充事实调整建议。", "related_finding_ids": []}, ensure_ascii=False)

    def wait(self, suffix, pending_key):
        deadline = time.monotonic() + 5
        while time.monotonic() < deadline:
            status, _, value = self.api("GET", self.path + suffix)
            self.assertEqual(status, 200, value)
            if not value.get(pending_key):
                return value
            time.sleep(.01)
        self.fail("fixture operation did not finish")

    def discuss(self, message="工资与交接能不能分开写？", opinion_revision=0):
        status, _, state = self.api("GET", self.path)
        self.assertEqual(status, 200, state)
        status, _, value = self.api("POST", self.path, {
            "expected_revision": state["revision"], "expected_delivery_revision": state["context"]["delivery_revision"],
            "expected_opinion_revision": opinion_revision, "request_id": str(uuid.uuid4()),
            "message": message, "config": self.config, "consent": True,
        })
        self.assertIn(status, (200, 202), value)
        return self.wait("", "pending_turn_id")

    def generate(self):
        state = self.api("GET", self.path + "/revisions")[2]
        discussion = self.api("GET", self.path)[2]
        body = {
            "expected_revision": state["revision"], "expected_opinion_revision": state["opinion_revision"],
            "expected_discussion_revision": discussion["revision"],
            "expected_delivery_revision": discussion["context"]["delivery_revision"],
            "request_id": str(uuid.uuid4()), "config": self.config, "consent": True,
        }
        status, _, value = self.api("POST", self.path + "/revisions", body)
        self.assertIn(status, (200, 202), value)
        state = self.wait("/revisions", "pending_candidate_id")
        self.assertEqual(state["candidates"][-1]["status"], "completed", state["candidates"][-1])
        return state, body

    def accept_body(self, state):
        return {"expected_revision": state["revision"], "expected_opinion_revision": state["opinion_revision"],
                "expected_discussion_revision": state["context"]["discussion_revision"],
                "expected_delivery_revision": state["context"]["delivery_revision"],
                "candidate_id": state["candidates"][-1]["id"]}

    def test_generate_compare_accept_continue_and_word_use_adopted_version(self):
        thread = self.discuss()
        from clause_discussion import _validate_turn
        earlier = copy.deepcopy(thread["turns"][0])
        for key in ("system_text", "system_sha256", "messages_sha256"):
            earlier["prompt_snapshot"].pop(key, None)
        self.assertEqual(_validate_turn(earlier)["assistant"], earlier["assistant"])
        state, body = self.generate()
        # Creating a proposal is not adoption; a repeated request cannot charge twice.
        self.assertEqual(self.api("GET", "/api/reports/detailed")[2]["report"], self.original["report"])
        status, _, _ = self.api("POST", self.path + "/revisions", body)
        self.assertIn(status, (200, 202))
        self.assertEqual(len(self.calls), 1)
        self.assertEqual(state["candidates"][-1]["base_finding"]["title"], self.original_finding["title"])
        status, _, _ = self.api("POST", self.path + "/revisions/accept", self.accept_body(state))
        self.assertEqual(status, 200)
        updated = self.api("GET", "/api/reports/detailed")[2]
        finding = next(f for f in updated["report"]["findings"] if f["id"] == "F01")
        self.assertEqual(finding["opinion_revision"], 1)
        self.assertIn("讨论后", finding["title"])
        self.assertTrue(updated["report"]["options"][0]["need_review"])
        self.assertEqual(self.raw_path.read_bytes(), self.raw_before)
        thread = self.discuss("那就按这个思路。", opinion_revision=1)
        self.assertEqual(len(thread["turns"]), 2)
        self.assertEqual(thread["opinion_revision"], 1)
        status, _, data = self.fixture.request("GET", "/api/reports/detailed/export.docx")
        self.assertEqual(status, 200)
        text = "\n".join(p.text for p in Document(io.BytesIO(data)).paragraphs)
        self.assertIn(finding["reason"], text)
        self.assertIn("讨论更新", text)
        self.assertIn("需重新核对", text)

    def test_new_discussion_makes_unadopted_candidate_stale(self):
        self.discuss()
        state, _ = self.generate()
        self.discuss("还有新事实：此前提到的安排不能实施。")
        status, _, _ = self.api("POST", self.path + "/revisions/accept", self.accept_body(state))
        self.assertEqual(status, 409)
        self.assertEqual(self.api("GET", "/api/reports/detailed")[2]["report"], self.original["report"])

    def test_cross_origin_is_rejected_and_withdrawal_retains_identity(self):
        status, _, _ = self.api("POST", self.path + "/revisions", {}, origin=False)
        self.assertEqual(status, 403)
        self.assertFalse(self.calls)
        self.discuss()
        self.action = "withdraw"
        state, _ = self.generate()
        status, _, _ = self.api("POST", self.path + "/revisions/accept", self.accept_body(state))
        self.assertEqual(status, 200)
        report = self.api("GET", "/api/reports/detailed")[2]["report"]
        finding = next(f for f in report["findings"] if f["id"] == "F01")
        self.assertEqual(finding["revision_action"], "withdraw")
        self.assertEqual(finding["rule_id"], self.original_finding["rule_id"])
        self.assertIsNone(finding["risk"])
        self.assertEqual(len(report["findings"]), len(self.original["report"]["findings"]))
        status, _, data = self.fixture.request("GET", "/api/reports/detailed/export.docx")
        self.assertEqual(status, 200)
        self.assertIn("讨论后撤回", "\n".join(p.text for p in Document(io.BytesIO(data)).paragraphs))


if __name__ == "__main__":
    unittest.main()
