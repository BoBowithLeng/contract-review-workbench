from __future__ import annotations

from tests.source_fixtures import labor_sources, write_sources

import hashlib
import copy
import http.client
import io
import json
import tempfile
import threading
import unittest
import urllib.parse
import zipfile
from pathlib import Path

from docx import Document

from contract_core import ParsedContract
from reader_server import REPORTS, create_server, _load_report
from reader_delivery import DeliveryManager
from tool_review import prepare


PROJECT = Path(__file__).resolve().parents[1]


def source_docx(text="（五）未完成工作交接前不支付工资，工资迟发不构成拖欠。") -> bytes:
    document = Document()
    document.add_paragraph(text)
    stream = io.BytesIO()
    document.save(stream)
    return stream.getvalue()


def result_for(contract: ParsedContract) -> dict:
    paragraph = contract.pieces[0].text
    return {
        "findings": [
            {
                "id": "F01", "rule_id": "LC-02", "title": "工资与交接绑定",
                "status": "发现风险", "risk": "高", "evidence_status": "依据已核",
                "refs": [{"id": "P000", "quote": paragraph}],
                "source_ids": ["LCL30_50"], "reason": "需要修改。", "suggestion": "分别处理。",
                "clause_draft": "（五）公司依法支付工资；工作交接责任另行处理。", "questions": [],
            },
            {
                "id": "F02", "rule_id": "LC-02", "title": "重复词定位测试",
                "status": "需要法务判断", "risk": None, "evidence_status": "适用待判断",
                "refs": [{"id": "P000", "quote": "工资"}],
                "source_ids": ["LCL30_50"], "reason": "需人工选定具体范围。", "suggestion": "不要自动选。",
                "clause_draft": "", "questions": [],
            },
        ],
        "checks": [
            {"rule_id": "LC-01", "status": "未检出", "reason": "已审。"},
            {"rule_id": "LC-02", "status": "发现风险", "reason": "已审。"},
            {"rule_id": "LC-03", "status": "未检出", "reason": "已审。"},
        ],
        "options": [{
            "id": "O1", "title": "工资与交接分开", "finding_ids": ["F01"],
            "preconditions": ["移除支付前提"], "steps": ["明确支付节点"],
            "cost": "需单独核算。", "remaining_risk": "工资口径仍需核对。",
            "source_ids": ["LCL30_50"],
            "clause_draft": "（五）公司依法支付工资；工作交接责任另行处理。",
            "relationship": "可与提前交接并用。",
        }],
        "uncovered": [],
    }


class ReaderDeliveryTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        base = Path(self.temp.name)
        self.root = base / "project"
        (self.root / "reader_ui").mkdir(parents=True)
        for name in ("index.html", "style.css", "app.js", "workspace.js", "delivery.js"):
            (self.root / "reader_ui" / name).write_text(name, encoding="utf-8")
        write_sources(self.root / "sources.json", labor_sources())
        self.detailed_source = base / "detailed.docx"
        self.table_source = base / "table.docx"
        self.detailed_source.write_bytes(source_docx())
        self.table_source.write_bytes(source_docx("表格报告绑定的普通段落。"))

        for key, spec in REPORTS.items():
            run = self.root / "runtime" / spec["directory"]
            run.mkdir(parents=True)
            if key == "detailed":
                input_path = self.detailed_source
            elif key == "table":
                input_path = self.table_source
            else:
                input_path = run / "minimal.txt"
                input_path.write_text("纯文本短合同。", encoding="utf-8")
            packet = prepare(input_path, run / "request.json", detailed=True, sources=labor_sources())
            contract = ParsedContract.model_validate(packet["contract"])
            response = {
                "request_digest": packet["request_digest"],
                "model": "test-sol",
                "result": result_for(contract),
            }
            # The table/minimal fixtures do not need a meaningful issue; exact
            # quotes still keep the saved report valid.
            if key != "detailed":
                response["result"]["findings"] = response["result"]["findings"][:1]
                response["result"]["findings"][0]["refs"] = [{"id": "P000", "quote": contract.pieces[0].text}]
                response["result"]["options"][0]["finding_ids"] = ["F01"]
            (run / "response.json").write_text(json.dumps(response, ensure_ascii=False), encoding="utf-8")

        self.state_dir = base / "reader-state"
        self.server = create_server(
            root=self.root,
            state_dir=self.state_dir,
            source_paths={
                "detailed": self.detailed_source,
                "minimal": None,
                "table": self.table_source,
            },
        )
        self.thread = threading.Thread(target=self.server.serve_forever, daemon=True)
        self.thread.start()
        self.addCleanup(self._stop)
        self.port = self.server.server_port

    def _stop(self):
        self.server.shutdown()
        self.server.server_close()
        self.thread.join(timeout=3)

    def request(self, method, path, body=None):
        connection = http.client.HTTPConnection("127.0.0.1", self.port, timeout=15)
        headers = {
            "Host": f"127.0.0.1:{self.port}",
            "Origin": f"http://127.0.0.1:{self.port}",
        }
        payload = None
        if body is not None:
            payload = json.dumps(body, ensure_ascii=False).encode("utf-8")
            headers["Content-Type"] = "application/json"
        connection.request(method, path, body=payload, headers=headers)
        response = connection.getresponse()
        data = response.read()
        result = response.status, dict(response.getheaders()), data
        connection.close()
        return result

    def json_request(self, method, path, body=None):
        status, headers, data = self.request(method, path, body)
        return status, headers, json.loads(data.decode("utf-8"))

    def delivery(self):
        status, _, payload = self.json_request("GET", "/api/reports/detailed/delivery")
        self.assertEqual(status, 200)
        self.assertTrue(payload["available"])
        return payload

    def review_once(self):
        _, _, report = self.json_request("GET", "/api/reports/detailed")
        reviews = report["state"]["reviews"]
        reviews["F01"]["decision"] = "接受"
        status, _, saved = self.json_request(
            "POST", "/api/reports/detailed/state",
            {"expected_revision": report["state"]["revision"], "reviews": reviews, "selected_options": ["O1"]},
        )
        self.assertEqual(status, 200)
        return saved["state"]

    def test_availability_binding_targets_and_independent_state(self):
        review = self.review_once()
        self.assertEqual(review["revision"], 1)
        _, _, report_after_old_accept = self.json_request("GET", "/api/reports/detailed")
        self.assertEqual(report_after_old_accept["actions"]["F01"]["status"], "待处理")
        payload = self.delivery()
        self.assertEqual(payload["state"]["revision"], 0)
        self.assertEqual(payload["state"]["edits"], [])
        target = payload["targets"]["F01"][0]
        self.assertTrue(target["supported"])
        self.assertEqual(target["old_text"], source_docx_text(self.detailed_source.read_bytes()))
        duplicate = payload["targets"]["F02"][0]
        self.assertFalse(duplicate["supported"])
        self.assertIn("多次出现", duplicate["reason"])
        self.assertTrue((self.state_dir / "delivery").is_dir())

        status, _, minimal = self.json_request("GET", "/api/reports/minimal/delivery")
        self.assertEqual(status, 200)
        self.assertFalse(minimal["available"])
        self.assertIn("纯文本", minimal["reason"])

    def test_confirmed_edit_comment_export_and_undo(self):
        review = self.review_once()
        payload = self.delivery()
        target = payload["targets"]["F01"][0]
        edit_id = f"edit-F01-{target['piece_id']}-{target['start']}"
        edit = {
            "id": edit_id,
            "piece_id": target["piece_id"],
            "paragraph_sha256": target["paragraph_sha256"],
            "start": target["start"],
            "end": target["end"],
            "old_text": target["old_text"],
            "new_text": "（五）公司依法支付工资；工作交接责任另行处理。",
            "confirmed": True,
            "finding_ids": ["F01"],
            "option_ids": ["O1"],
        }
        comment = {"id": "comment-F01", "text": "请业务部门核对实际交接流程。", "confirmed": True, "edit_id": edit_id}
        save_body = {
            "expected_revision": 0,
            "source_sha256": payload["source"]["source_sha256"],
            "author": "波波",
            "edits": [edit],
            "comments": [comment],
        }
        status, _, saved = self.json_request("POST", "/api/reports/detailed/delivery", save_body)
        self.assertEqual(status, 200)
        self.assertEqual(saved["state"]["revision"], 1)
        self.assertEqual(saved["state"]["author"], "波波")
        self.assertRegex(saved["state"]["created_at"], r"Z$")
        _, _, action_report = self.json_request("GET", "/api/reports/detailed")
        self.assertEqual(action_report["actions"]["F01"]["status"], "已修改")
        self.assertEqual(action_report["actions"]["F01"]["label"], "已修改并批注")

        export_body = {
            "expected_revision": 1,
            "source_sha256": payload["source"]["source_sha256"],
            "review_revision": review["revision"],
            "variants": ["clean", "tracked_comments", "report"],
        }
        status, headers, archive_data = self.request(
            "POST", "/api/reports/detailed/delivery/export", export_body
        )
        self.assertEqual(status, 200)
        self.assertEqual(headers["Content-Type"], "application/zip")
        zip_name = urllib.parse.unquote(headers["Content-Disposition"].split("filename*=UTF-8''", 1)[1])
        self.assertTrue(zip_name.endswith("-交付文件.zip"))
        stem = zip_name[:-len("-交付文件.zip")]
        self.assertTrue(stem and stem != "合同")
        clean_name, report_name = stem + "-洁净版.docx", stem + "-审查报告.docx"
        with zipfile.ZipFile(io.BytesIO(archive_data)) as archive:
            names = set(archive.namelist())
            self.assertEqual(names, {clean_name, stem + "-修订与批注版.docx", report_name})
            self.assertNotIn("manifest.json", names)
            self.assertTrue(all(name.endswith(".docx") for name in names))
            self.assertNotIn("00-原稿.docx", names)
            self.assertNotIn("confirmed-plan.json", names)
            self.assertNotIn("source-snapshot.json", names)
            clean_text = source_docx_text(archive.read(clean_name))
            self.assertEqual(clean_text, edit["new_text"])
            report_text = "\n".join(
                p.text for p in Document(io.BytesIO(archive.read(report_name))).paragraphs
            )
            self.assertIn("请业务部门核对实际交接流程", report_text)
            self.assertNotIn("SHA-256", report_text)
            self.assertNotIn(edit["id"], report_text)
            self.assertNotIn(comment["id"], report_text)
            self.assertNotIn("补充案例", report_text)
            self.assertIn("合成测试资料 LCL30_50", report_text)

        status, _, blocked = self.json_request(
            "POST",
            "/api/reports/detailed/actions",
            {
                "action": "ignore",
                "finding_id": "F01",
                "expected_review_revision": review["revision"],
                "expected_delivery_revision": 1,
            },
        )
        self.assertEqual(status, 409)
        self.assertIn("先撤销实际文档操作", blocked["error"])

        # Repeated downloads preserve the same confirmed contract and report;
        # technical manifests remain internal rather than becoming deliverables.
        status, _, archive_again = self.request(
            "POST", "/api/reports/detailed/delivery/export", export_body
        )
        self.assertEqual(status, 200)
        with zipfile.ZipFile(io.BytesIO(archive_data)) as first, zipfile.ZipFile(io.BytesIO(archive_again)) as second:
            self.assertEqual(set(first.namelist()), set(second.namelist()))
            self.assertEqual(source_docx_text(first.read(clean_name)),
                             source_docx_text(second.read(clean_name)))
        _, _, reread_after_export = self.json_request("GET", "/api/reports/detailed/delivery")
        self.assertEqual(reread_after_export["state"]["created_at"], saved["state"]["created_at"])
        self.assertEqual(reread_after_export["state"]["revision"], saved["state"]["revision"])
        direct_status, _, direct_report = self.request("GET", "/api/reports/detailed/export.docx")
        self.assertEqual(direct_status, 200)
        direct_text = "\n".join(p.text for p in Document(io.BytesIO(direct_report)).paragraphs)
        self.assertIn("请业务部门核对实际交接流程", direct_text)
        self.assertNotIn("SHA-256", direct_text)

        status, _, undone = self.json_request(
            "POST", "/api/reports/detailed/delivery/undo",
            {"expected_revision": 1, "source_sha256": payload["source"]["source_sha256"]},
        )
        self.assertEqual(status, 200)
        self.assertEqual(undone["state"]["revision"], 2)
        self.assertEqual(undone["state"]["edits"], [])
        self.assertEqual(undone["state"]["comments"], [])
        self.assertEqual(undone["state"]["author"], "")
        self.assertIsNone(undone["state"]["created_at"])
        _, _, reread = self.json_request("GET", "/api/reports/detailed/delivery")
        self.assertEqual(reread["state"]["revision"], 2)
        self.assertEqual(reread["state"]["author"], "")
        self.assertIsNone(reread["state"]["created_at"])
        _, _, restored_actions = self.json_request("GET", "/api/reports/detailed")
        self.assertEqual(restored_actions["actions"]["F01"]["status"], "待处理")

    def test_plan_validation_conflict_and_source_change_are_rejected(self):
        payload = self.delivery()
        target = payload["targets"]["F01"][0]
        edit = {
            "id": "edit-bad", "piece_id": target["piece_id"],
            "paragraph_sha256": target["paragraph_sha256"], "start": target["start"], "end": target["end"],
            "old_text": "被篡改的原文", "new_text": "新文字", "confirmed": True,
        }
        body = {
            "expected_revision": 0, "source_sha256": payload["source"]["source_sha256"],
            "author": "波波", "edits": [edit], "comments": [],
        }
        status, _, error = self.json_request("POST", "/api/reports/detailed/delivery", body)
        self.assertEqual(status, 400)
        self.assertIn("old_text", error["error"])

        body["edits"] = []
        status, _, saved = self.json_request("POST", "/api/reports/detailed/delivery", body)
        self.assertEqual(status, 200)
        status, _, conflict = self.json_request("POST", "/api/reports/detailed/delivery", body)
        self.assertEqual(status, 409)
        self.assertEqual(conflict["current_revision"], 1)

        self.detailed_source.write_bytes(source_docx("原稿后来发生变化。"))
        status, _, unavailable = self.json_request("GET", "/api/reports/detailed/delivery")
        self.assertEqual(status, 200)
        self.assertFalse(unavailable["available"])
        self.assertIn("合同文字不一致", unavailable["reason"])
        body["expected_revision"] = saved["state"]["revision"]
        status, _, _ = self.json_request("POST", "/api/reports/detailed/delivery", body)
        self.assertEqual(status, 409)

    def test_revision_zero_can_export_an_explicit_empty_snapshot(self):
        payload = self.delivery()
        status, headers, data = self.request(
            "POST",
            "/api/reports/detailed/delivery/export",
            {
                "expected_revision": 0,
                "source_sha256": payload["source"]["source_sha256"],
                "review_revision": 0,
                "variants": ["clean", "report"],
            },
        )
        self.assertEqual(status, 200)
        stem = urllib.parse.unquote(headers["Content-Disposition"].split("filename*=UTF-8''", 1)[1])[:-len("-交付文件.zip")]
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            self.assertEqual(set(archive.namelist()), {stem + "-洁净版.docx", stem + "-审查报告.docx"})
            self.assertEqual(source_docx_text(archive.read(stem + "-洁净版.docx")),
                             source_docx_text(self.detailed_source.read_bytes()))
        _, _, unchanged = self.json_request("GET", "/api/reports/detailed/delivery")
        self.assertEqual(unchanged["state"]["revision"], 0)
        self.assertEqual(unchanged["state"]["edits"], [])
        self.assertEqual(unchanged["state"]["comments"], [])

    def test_legacy_state_without_time_is_not_given_a_fake_history(self):
        payload = self.delivery()
        body = {
            "expected_revision": 0,
            "source_sha256": payload["source"]["source_sha256"],
            "author": "旧版确认人",
            "edits": [],
            "comments": [],
        }
        status, _, saved = self.json_request("POST", "/api/reports/detailed/delivery", body)
        self.assertEqual(status, 200)
        state_path = next((self.state_dir / "delivery").glob("detailed-*-state.json"))
        raw = json.loads(state_path.read_text(encoding="utf-8"))
        raw.pop("created_at")
        state_path.write_text(json.dumps(raw, ensure_ascii=False), encoding="utf-8")

        _, _, reread = self.json_request("GET", "/api/reports/detailed/delivery")
        self.assertTrue(reread["available"])
        self.assertIsNone(reread["state"]["created_at"])
        status, _, error = self.json_request(
            "POST",
            "/api/reports/detailed/delivery/export",
            {
                "expected_revision": 1,
                "source_sha256": payload["source"]["source_sha256"],
                "review_revision": 0,
                "variants": ["clean"],
            },
        )
        self.assertEqual(status, 500)
        self.assertIn("未伪造历史时间", error["error"])

        body["expected_revision"] = 1
        status, _, renewed = self.json_request("POST", "/api/reports/detailed/delivery", body)
        self.assertEqual(status, 200)
        self.assertEqual(renewed["state"]["revision"], 2)
        self.assertRegex(renewed["state"]["created_at"], r"Z$")


def source_docx_text(data: bytes) -> str:
    return "\n".join(paragraph.text for paragraph in Document(io.BytesIO(data)).paragraphs)


class ConfirmedDraftScopeTests(unittest.TestCase):
    setUp = ReaderDeliveryTests.setUp
    _stop = ReaderDeliveryTests._stop
    request = ReaderDeliveryTests.request
    json_request = ReaderDeliveryTests.json_request
    delivery = ReaderDeliveryTests.delivery

    def fixture(self):
        document = Document()
        for text in ("1.1 工资按时支付。", "1.2 交接另行处理。", "2.1 其他不相干约定。"):
            document.add_paragraph(text)
        stream = io.BytesIO()
        document.save(stream)
        self.detailed_source.write_bytes(stream.getvalue())
        run = self.root / "runtime" / REPORTS["detailed"]["directory"]
        # prepare refuses to overwrite a prior packet; the isolated fixture's
        # source/request history is rebuilt before any binding is read.
        (run / "request.json").unlink()
        packet = prepare(self.detailed_source, run / "request.json", detailed=True, sources=labor_sources())
        contract = ParsedContract.model_validate(packet["contract"])
        result = result_for(contract)
        result["findings"][0]["refs"].append({"id": "P001", "quote": contract.pieces[1].text})
        result["findings"][0]["clause_draft"] = "1.1 公司依法按时支付工资。"
        (run / "response.json").write_text(json.dumps({"request_digest": packet["request_digest"],
            "model": "test-sol", "result": result}, ensure_ascii=False), encoding="utf-8")
        payload = self.delivery()
        pieces = payload["source"]["pieces"]
        target = pieces[0]
        candidate = {"status": "ready", "candidate_id": "a" * 64,
            "source_sha256": payload["source"]["source_sha256"],
            "draft_sha256": hashlib.sha256(result["findings"][0]["clause_draft"].encode()).hexdigest(),
            "edits": [{"piece_id": target["id"], "paragraph_sha256": target["paragraph_sha256"],
                "start": 0, "end": len(target["text"]), "old_text": target["text"],
                "new_text": result["findings"][0]["clause_draft"], "location": "1.1"}],
            "insertions": [], "coverage_complete": True,
            "paragraph_actions": [
                {"piece_id": "P000", "action": "modify", "reason": "明确依法按时支付工资。", "location": "1.1"},
                {"piece_id": "P001", "action": "retain", "reason": "交接已有独立约定，不与支付工资挂钩。", "location": "1.2"}],
            "suggestion_actions": [{"unit_id": "S1", "text": "依法按时支付工资", "action": "implemented", "reason": "已落实。"}]}
        self.server.delivery.draft_candidate_resolver = lambda loaded, fid, cid, sha: copy.deepcopy(candidate)
        edit = {"id": "edit-F01-prepared-P000", "piece_id": "P000",
                "paragraph_sha256": target["paragraph_sha256"], "start": 4, "end": 6,
                "old_text": "工资", "new_text": "公司依法按时支付工资", "confirmed": True, "finding_ids": ["F01"]}
        # Match the exact candidate via unchanged-edge stripping, not edit IDs.
        from reader_delivery import _edit_signature
        edit.update({key: value for key, value in _edit_signature(candidate["edits"][0]).items() if key != "kind"})
        body = {"expected_revision": 0, "source_sha256": payload["source"]["source_sha256"],
                "author": "法务复核", "edits": [edit], "comments": [],
                "draft_confirmation": {"finding_id": "F01", "candidate_id": candidate["candidate_id"]}}
        return payload, body, candidate

    def save_body(self, body):
        status, _, result = self.json_request("POST", "/api/reports/detailed/delivery", body)
        self.assertEqual(status, 200, result)
        return result["state"]

    def test_confirmation_refresh_unrelated_edits_and_comments_then_remove_undo(self):
        payload, body, candidate = self.fixture()
        state = self.save_body(body)
        self.assertEqual(state["draft_confirmations"]["F01"]["paragraph_actions"], candidate["paragraph_actions"])
        self.assertTrue(state["draft_confirmations"]["F01"]["coverage_complete"])
        self.assertEqual(self.delivery()["state"]["draft_confirmations"], state["draft_confirmations"])
        # A fresh manager reads the persisted metadata without needing drafts.
        fresh = DeliveryManager(self.state_dir, {"detailed": self.detailed_source})
        self.assertEqual(fresh.describe(_load_report(self.root, "detailed"))["state"]["draft_confirmations"], state["draft_confirmations"])
        body.pop("draft_confirmation")
        body["expected_revision"] = state["revision"]
        unrelated = payload["source"]["pieces"][2]
        body["edits"].append({"id": "unrelated", "piece_id": "P002", "paragraph_sha256": unrelated["paragraph_sha256"],
            "start": 0, "end": len(unrelated["text"]), "old_text": unrelated["text"], "new_text": "2.1 无关文字调整。", "confirmed": True})
        body["comments"] = [{"id": "note", "edit_id": "unrelated", "text": "请确认。", "confirmed": True}]
        state = self.save_body(body)
        self.assertEqual(state["draft_confirmations"]["F01"]["status"], "confirmed")
        body["expected_revision"] = state["revision"]
        body["edits"] = [body["edits"][1]]
        state = self.save_body(body)
        self.assertEqual(state["draft_confirmations"], {})
        status, _, undone = self.json_request("POST", "/api/reports/detailed/delivery/undo",
            {"expected_revision": state["revision"], "source_sha256": body["source_sha256"]})
        self.assertEqual(status, 200)
        self.assertEqual(undone["state"]["draft_confirmations"]["F01"]["paragraph_actions"], candidate["paragraph_actions"])

    def test_manual_preview_is_not_falsely_approved_and_later_changes_drop_scope(self):
        _, body, _ = self.fixture()
        body["edits"][0]["new_text"] += "另行核实。"
        state = self.save_body(body)
        scope = state["draft_confirmations"]["F01"]
        self.assertEqual(scope["status"], "changed")
        self.assertFalse(scope["coverage_complete"])
        self.assertEqual(scope["paragraph_actions"][1]["action"], "unexplained")
        self.assertEqual(scope["suggestion_actions"][0]["action"], "unexplained")
        body.pop("draft_confirmation")
        body["expected_revision"] = state["revision"]
        body["edits"][0]["new_text"] += "再修改。"
        state = self.save_body(body)
        self.assertEqual(state["draft_confirmations"], {})

    def test_change_to_retained_clause_invalidates_reason_and_undo_restores(self):
        payload, body, candidate = self.fixture()
        state = self.save_body(body)
        body.pop("draft_confirmation")
        body["expected_revision"] = state["revision"]
        piece = payload["source"]["pieces"][1]
        body["edits"].append({"id": "edit-other-P001", "piece_id": "P001", "paragraph_sha256": piece["paragraph_sha256"],
            "start": 0, "end": len(piece["text"]), "old_text": piece["text"], "new_text": "1.2 交接规定改变。", "confirmed": True})
        state = self.save_body(body)
        self.assertEqual(state["draft_confirmations"], {})
        _, _, undone = self.json_request("POST", "/api/reports/detailed/delivery/undo",
            {"expected_revision": state["revision"], "source_sha256": body["source_sha256"]})
        self.assertEqual(undone["state"]["draft_confirmations"]["F01"]["paragraph_actions"], candidate["paragraph_actions"])

    def test_legacy_exact_recovery_is_read_only_and_does_not_fabricate_retain(self):
        _, body, candidate = self.fixture()
        body.pop("draft_confirmation")
        self.save_body(body)
        path = next((self.state_dir / "delivery").glob("detailed-*-state.json"))
        raw = json.loads(path.read_text(encoding="utf-8"))
        raw.pop("draft_confirmations")
        path.write_text(json.dumps(raw, ensure_ascii=False), encoding="utf-8")
        before = path.read_bytes()
        self.assertEqual(self.delivery()["state"]["draft_confirmations"]["F01"]["paragraph_actions"], candidate["paragraph_actions"])
        self.assertEqual(path.read_bytes(), before)
        # A historical candidate with no reasons remains unexplained; the
        # resolver's legacy public projection never invents a retain action.
        candidate["paragraph_actions"][1].update(action="unexplained", reason="尚未说明。")
        candidate["coverage_complete"] = False
        scope = self.delivery()["state"]["draft_confirmations"]["F01"]
        self.assertEqual(scope["paragraph_actions"][1]["action"], "unexplained")
        self.assertFalse(scope["coverage_complete"])
        self.assertEqual(path.read_bytes(), before)
        candidate["edits"][0]["new_text"] += "不同候选。"
        self.assertEqual(self.delivery()["state"]["draft_confirmations"], {})

    def test_invalid_confirmation_does_not_save_edits(self):
        _, body, _ = self.fixture()
        body["edits"] = []
        status, _, error = self.json_request("POST", "/api/reports/detailed/delivery", body)
        self.assertEqual(status, 400, error)
        self.assertEqual(self.delivery()["state"]["revision"], 0)

    def test_posted_retain_metadata_is_not_trusted(self):
        _, body, candidate = self.fixture()
        body.pop("draft_confirmation")
        body["draft_confirmations"] = {"F01": {"paragraph_actions": candidate["paragraph_actions"]}}
        state = self.save_body(body)
        self.assertEqual(state["draft_confirmations"], {})

    def test_insertion_only_candidate_preserves_scope_and_kind(self):
        payload, body, candidate = self.fixture()
        piece = payload["source"]["pieces"][0]
        insertion = {"kind": "insert_paragraph_after", "piece_id": "P000", "paragraph_sha256": piece["paragraph_sha256"],
            "start": len(piece["text"]), "end": len(piece["text"]), "old_text": "", "new_text": "1.3 新增送达条款。"}
        candidate["edits"] = []
        candidate["insertions"] = [insertion]
        body["edits"] = [{**insertion, "id": "edit-F01-insertion-P000", "confirmed": True, "finding_ids": ["F01"]}]
        state = self.save_body(body)
        self.assertEqual(state["edits"][0]["kind"], "insert_paragraph_after")
        self.assertEqual(state["draft_confirmations"]["F01"]["status"], "confirmed")
        self.assertEqual(self.delivery()["state"]["draft_confirmations"], state["draft_confirmations"])

    def test_real_cached_resolver_exact_old_body_binding_and_no_cache_rewrites(self):
        payload, body, candidate = self.fixture()
        bundle = _load_report(self.root, "detailed")
        finding = bundle["report"]["findings"][0]
        from reader_delivery import _digest
        operation = {"schema_version": "2026-10-04.1", "candidate_id": candidate["candidate_id"],
            "status": "ready", "report_id": bundle["id"], "request_digest": bundle["request_digest"],
            "finding_id": "F01", "finding_sha256": _digest(finding), "source_sha256": body["source_sha256"],
            "based_on_revision": 0, "draft_sha256": candidate["draft_sha256"],
            "clarification_sha256": hashlib.sha256(b"").hexdigest(), "edits": candidate["edits"],
            "questions": [], "comment": "请核对工资支付约定。", "error": None,
            "paragraph_actions": candidate["paragraph_actions"]}
        directory = self.server.drafts.root / candidate["candidate_id"]
        directory.mkdir(parents=True)
        targets = [{"piece_id": piece["id"], "location": piece["id"], "old_text": piece["text"]}
                   for piece in payload["source"]["pieces"][:2]]
        frozen = {"identity": {"draft_text": finding["clause_draft"]}, "messages": [
            {"role": "system", "content": "fixture"}, {"role": "user", "content":
            "UNTRUSTED DATA:\n" + json.dumps({"allowed_targets": targets})}]}
        (directory / "operation.json").write_text(json.dumps(operation, ensure_ascii=False), encoding="utf-8")
        (directory / "request.json").write_text(json.dumps(frozen, ensure_ascii=False), encoding="utf-8")
        pointer = self.server.drafts.root / "latest-detailed-F01.json"
        pointer.write_text(json.dumps({"candidate_id": candidate["candidate_id"]}), encoding="utf-8")
        self.server.delivery.draft_candidate_resolver = self.server.drafts.resolve_confirmation
        before = {path: path.read_bytes() for path in (directory / "operation.json", directory / "request.json", pointer)}
        status, _, cached = self.json_request("GET", "/api/reports/detailed/drafts/F01")
        self.assertEqual(status, 200, cached)
        self.assertFalse(cached["coverage_complete"])
        self.assertTrue(all(row["action"] == "unexplained" for row in cached["suggestion_actions"]))
        state = self.save_body(body)
        self.assertEqual(state["draft_confirmations"]["F01"]["paragraph_actions"], candidate["paragraph_actions"])
        self.assertFalse(state["draft_confirmations"]["F01"]["coverage_complete"])
        state_path = next((self.state_dir / "delivery").glob("detailed-*-state.json"))
        raw = json.loads(state_path.read_text(encoding="utf-8"))
        raw.pop("draft_confirmations")
        state_path.write_text(json.dumps(raw, ensure_ascii=False), encoding="utf-8")
        state_before = state_path.read_bytes()
        recovered = self.delivery()["state"]["draft_confirmations"]["F01"]
        self.assertEqual(recovered["paragraph_actions"], candidate["paragraph_actions"])
        self.assertFalse(recovered["coverage_complete"])
        self.assertEqual(state_path.read_bytes(), state_before)
        self.assertEqual({path: path.read_bytes() for path in before}, before)
        # The same candidate ID under a mismatching source is rejected before
        # saving even otherwise valid confirmed words.
        operation["source_sha256"] = "b" * 64
        (directory / "operation.json").write_text(json.dumps(operation), encoding="utf-8")
        body["expected_revision"] = state["revision"]
        status, _, error = self.json_request("POST", "/api/reports/detailed/delivery", body)
        self.assertEqual(status, 400, error)
        self.assertEqual(state_path.read_bytes(), state_before)


if __name__ == "__main__":
    unittest.main()


class ClientFileNameTests(unittest.TestCase):
    def test_contract_name_becomes_a_safe_delivery_name(self):
        from reader_delivery import client_stem, client_zip_name
        self.assertEqual(client_stem("软件委托开发合同.docx"), "软件委托开发合同")
        self.assertEqual(client_zip_name("软件委托开发合同.docx"), "软件委托开发合同-交付文件.zip")
        self.assertEqual(client_stem('a<b>:"c|d?e*.docx'), "abcde")
        self.assertEqual(client_stem(r'C:\上传\采购合同.docx'), "采购合同")
        self.assertEqual(client_stem("合同. .docx"), "合同")
        for empty in (None, "", ".docx", '<>.docx'):
            self.assertEqual(client_stem(empty), "合同")
        self.assertEqual(len(client_stem("长" * 200 + ".docx")), 60)
