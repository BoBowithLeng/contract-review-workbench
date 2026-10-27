from __future__ import annotations

from tests.source_fixtures import labor_sources, write_sources

import http.client
import io
import json
import tempfile
import threading
import unittest
import urllib.parse
from pathlib import Path
from unittest.mock import patch

from docx import Document

from contract_core import ParsedContract
from model_client import ModelClientError
from reader_server import REPORTS, create_server
from tool_review import prepare


PROJECT = Path(__file__).resolve().parents[1]


def model_result(contract: ParsedContract, *, bad_quote: bool = False) -> dict:
    quote = "不存在的原文" if bad_quote else contract.pieces[0].text
    return {
        "findings": [{
            "id": "F01",
            "rule_id": "LC-02",
            "title": "工资与交接",
            "status": "发现风险",
            "risk": "高",
            "evidence_status": "依据已核",
            "refs": [{"id": contract.pieces[0].id, "quote": quote}],
            "source_ids": ["LCL30_50"],
            "reason": "劳动合同法第三十条、第五十条需要分别核对工资与交接。",
            "suggestion": "工资按明确节点支付，交接责任另行处理。",
            "clause_draft": "公司依法、依约支付工资；交接责任另行处理。",
            "questions": [],
        }],
        "checks": [
            {"rule_id": "LC-01", "status": "未检出", "reason": "已审。"},
            {"rule_id": "LC-02", "status": "发现风险", "reason": "检出一项风险。"},
            {"rule_id": "LC-03", "status": "未检出", "reason": "已审。"},
        ],
        "options": [] if bad_quote else [{
            "id": "O1",
            "title": "工资与交接分开",
            "finding_ids": ["F01"],
            "preconditions": ["移除工资支付的交接前提"],
            "steps": ["确定工资节点", "另列交接清单"],
            "cost": "财务需单独核算。",
            "remaining_risk": "工资口径仍需核对。",
            "source_ids": ["LCL30_50"],
            "clause_draft": "公司依法、依约支付工资；交接责任另行处理。",
            "relationship": "可与提前交接流程并用。",
        }],
        "uncovered": [],
    }


class ReaderServerTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name) / "project"
        (self.root / "reader_ui").mkdir(parents=True)
        (self.root / "reader_ui" / "index.html").write_text("<h1>reader</h1>", encoding="utf-8")
        (self.root / "reader_ui" / "style.css").write_text("body{}", encoding="utf-8")
        (self.root / "reader_ui" / "app.js").write_text("window.reader=true;", encoding="utf-8")
        write_sources(self.root / "sources.json", labor_sources())

        for key, spec in REPORTS.items():
            run = self.root / "runtime" / spec["directory"]
            run.mkdir(parents=True)
            input_path = run / "input.txt"
            input_path.write_text("第一段合同文字。\n员工未完成交接，公司不支付工资。", encoding="utf-8")
            request_path = run / "request.json"
            packet = prepare(input_path, request_path, detailed=True, sources=labor_sources())
            contract = ParsedContract.model_validate(packet["contract"])
            result = model_result(contract, bad_quote=(key == "table"))
            if key == "detailed":
                # Model order is deliberately the reverse of source order.
                result["findings"][0]["title"] = "后段意见"
                result["findings"][0]["refs"] = [{"id": "P001", "quote": contract.pieces[1].text}]
                result["findings"].append({
                    "id": "F02",
                    "rule_id": "LC-02",
                    "title": "前段意见",
                    "status": "需要法务判断",
                    "risk": None,
                    "evidence_status": "适用待判断",
                    "refs": [{"id": "P000", "quote": contract.pieces[0].text}],
                    "source_ids": ["LCL30_50"],
                    "reason": "前段文字需要先行阅读。",
                    "suggestion": "结合后段一并核对。",
                    "clause_draft": "前段候选改稿。",
                    "questions": [],
                })
            response = {
                "request_digest": packet["request_digest"],
                "model": "test-sol",
                "result": result,
            }
            (run / "response.json").write_text(json.dumps(response, ensure_ascii=False), encoding="utf-8")

        self.state_dir = Path(self.temp.name) / "states"
        self.server = create_server(root=self.root, state_dir=self.state_dir)
        self.thread = threading.Thread(target=self.server.serve_forever, daemon=True)
        self.thread.start()
        self.addCleanup(self._stop_server)
        self.port = self.server.server_port

    def _stop_server(self):
        self.server.shutdown()
        self.server.server_close()
        self.thread.join(timeout=3)

    def request(self, method: str, path: str, body=None, *, origin=True, host=None, content_type="application/json"):
        connection = http.client.HTTPConnection("127.0.0.1", self.port, timeout=10)
        headers = {"Host": host or f"127.0.0.1:{self.port}"}
        if origin:
            headers["Origin"] = f"http://127.0.0.1:{self.port}"
        data = None
        if body is not None:
            data = json.dumps(body, ensure_ascii=False).encode("utf-8")
            headers["Content-Type"] = content_type
        connection.request(method, path, body=data, headers=headers)
        response = connection.getresponse()
        payload = response.read()
        result = (response.status, dict(response.getheaders()), payload)
        connection.close()
        return result

    def json_request(self, method: str, path: str, body=None, **kwargs):
        status, headers, data = self.request(method, path, body, **kwargs)
        return status, headers, json.loads(data.decode("utf-8"))

    def test_model_catalog_route_requires_same_origin_and_returns_only_catalog(self):
        config = {"base_url": "https://example.com/v1", "api_key": "sk-route-secret"}
        with patch("reader_server.list_models", return_value=[{"id": "m", "name": "Model"}]) as fetch:
            status, _, body = self.json_request("POST", "/api/models", {"config": config})
            self.assertEqual(status, 200)
            self.assertEqual(body, {"models": [{"id": "m", "name": "Model"}]})
            fetch.assert_called_once_with(config)
            self.assertNotIn(config["api_key"], json.dumps(body))
            status, _, _ = self.json_request("POST", "/api/models", {"config": config}, origin=False)
            self.assertEqual(status, 403)
            fetch.assert_called_once()
            status, _, _ = self.json_request("POST", "/api/models", {"config": config}, host="evil.example")
            self.assertEqual(status, 421)
            fetch.assert_called_once()
        status, _, body = self.json_request("POST", "/api/models", {"config": {"base_url": ""}})
        self.assertEqual(status, 400)
        self.assertNotIn(config["api_key"], json.dumps(body))
        with patch("reader_server.list_models", side_effect=ModelClientError("模型列表返回 HTTP 404；可手填模型 ID。")):
            status, _, body = self.json_request("POST", "/api/models", {"config": config})
            self.assertEqual(status, 502)
            self.assertIn("手填", body["error"])

    def test_model_settings_route_persists_encrypted_and_rejects_cross_origin(self):
        status, headers, empty = self.json_request("GET", "/api/model-settings", origin=False)
        self.assertEqual(status, 200)
        self.assertEqual(headers["Cache-Control"], "no-store")
        self.assertEqual(empty["revision"], 0)
        self.assertFalse(self.server.model_settings.path.exists())
        config = {"base_url": "https://api.deepseek.com", "api_key": "sk-route-private", "model": "deepseek-flash"}
        body = {"expected_revision": 0, "action": "save", "provider": "DeepSeek", "config": config}
        status, _, _ = self.json_request("POST", "/api/model-settings", body, origin=False)
        self.assertEqual(status, 403)
        status, _, saved = self.json_request("POST", "/api/model-settings", body)
        self.assertEqual(status, 200)
        self.assertEqual(saved["profiles"], [{"provider": "DeepSeek", **config}])
        self.assertNotIn(config["api_key"], self.server.model_settings.path.read_text(encoding="utf-8"))
        status, _, reloaded = self.json_request("GET", "/api/model-settings", origin=False)
        self.assertEqual(status, 200)
        self.assertEqual(reloaded, saved)
        status, _, _ = self.json_request("GET", "/api/model-settings", host="evil.example")
        self.assertEqual(status, 421)
        connection = http.client.HTTPConnection("127.0.0.1", self.port, timeout=10)
        connection.request("GET", "/api/model-settings", headers={
            "Host": f"127.0.0.1:{self.port}", "Origin": "http://evil.example",
        })
        response = connection.getresponse()
        self.assertEqual(response.status, 403)
        response.read()
        connection.close()
        status, _, conflict = self.json_request("POST", "/api/model-settings", body)
        self.assertEqual(status, 409)
        self.assertEqual(conflict["current_revision"], 1)
        status, _, removed = self.json_request("POST", "/api/model-settings", {
            "expected_revision": 1, "action": "remove", "provider": "DeepSeek",
        })
        self.assertEqual(status, 200)
        self.assertIsNone(removed["active_provider"])
        self.assertEqual(removed["profiles"], [])

    def test_reads_only_three_fixed_reports(self):
        status, headers, listing = self.json_request("GET", "/api/reports", origin=False)
        self.assertEqual(status, 200)
        self.assertEqual({item["id"] for item in listing["reports"]}, {"detailed", "minimal", "table"})
        self.assertEqual(headers["Cache-Control"], "no-store")
        for report_id in ("detailed", "minimal", "table"):
            status, _, payload = self.json_request("GET", f"/api/reports/{report_id}", origin=False)
            self.assertEqual(status, 200)
            self.assertEqual(payload["id"], report_id)
            self.assertTrue(payload["paragraphs"])
            self.assertEqual(len(payload["completeness"]), 8)
            self.assertEqual(payload["state"]["revision"], 0)
            self.assertFalse(payload["state"]["can_undo"])
            self.assertIn("LCL17", {item["id"] for item in payload["sources"]})
        _, _, detailed = self.json_request("GET", "/api/reports/detailed", origin=False)
        self.assertEqual([item["id"] for item in detailed["report"]["findings"]], ["F02", "F01"])
        _, _, table = self.json_request("GET", "/api/reports/table", origin=False)
        self.assertEqual(table["report"]["findings"][0]["status"], "处理失败")

    def test_retired_case_catalog_endpoint_is_not_found(self):
        for query in ("finding_id=F01", "finding_id=F01&q=synthetic"):
            status, _, _ = self.json_request(
                "GET", "/api/reports/detailed/cases?" + query, origin=False)
            self.assertEqual(status, 404)

    def test_ignore_restore_actions_work_without_docx_source(self):
        _, _, detail = self.json_request("GET", "/api/reports/minimal", origin=False)
        finding_id = detail["report"]["findings"][0]["id"]
        self.assertEqual(detail["actions"][finding_id]["status"], "待处理")
        status, _, ignored = self.json_request(
            "POST",
            "/api/reports/minimal/actions",
            {
                "action": "ignore",
                "finding_id": finding_id,
                "expected_review_revision": 0,
                "expected_delivery_revision": None,
            },
        )
        self.assertEqual(status, 200)
        self.assertEqual(ignored["review_state"]["revision"], 1)
        self.assertEqual(ignored["actions"][finding_id]["status"], "已忽略")
        self.assertEqual(
            ignored["review_state"]["reviews"][finding_id]["note"],
            "操作记录：不采用本条修改建议；不表示风险已解决，未填写复核理由。",
        )
        status, _, word = self.request("GET", "/api/reports/minimal/export.docx", origin=False)
        self.assertEqual(status, 200)
        word_text = "\n".join(p.text for p in Document(io.BytesIO(word)).paragraphs)
        self.assertTrue(any(cell.text == "不采用修改建议"
                            for table in Document(io.BytesIO(word)).tables
                            for row in table.rows for cell in row.cells))
        self.assertIn("不采用本条修改建议；不表示风险已解决", word_text)
        status, _, restored = self.json_request(
            "POST",
            "/api/reports/minimal/actions",
            {
                "action": "restore",
                "finding_id": finding_id,
                "expected_review_revision": 1,
                "expected_delivery_revision": None,
            },
        )
        self.assertEqual(status, 200)
        self.assertEqual(restored["review_state"]["revision"], 2)
        self.assertEqual(restored["review_state"]["reviews"][finding_id]["decision"], "未复核")
        self.assertEqual(restored["actions"][finding_id]["status"], "待处理")

    def test_save_reload_conflict_undo_and_export(self):
        _, _, detail = self.json_request("GET", "/api/reports/detailed", origin=False)
        reviews = detail["state"]["reviews"]
        reviews["F01"].update({
            "decision": "接受",
            "risk": "高",
            "note": "按公司流程执行并保留记录。",
            "clause_draft": "人工确认后的工资与交接条款。",
        })
        body = {"expected_revision": 0, "reviews": reviews, "selected_options": ["O1"]}
        status, _, saved = self.json_request("POST", "/api/reports/detailed/state", body)
        self.assertEqual(status, 200)
        self.assertEqual(saved["state"]["revision"], 1)
        self.assertTrue(saved["state"]["can_undo"])

        _, _, reread = self.json_request("GET", "/api/reports/detailed", origin=False)
        self.assertEqual(reread["state"], saved["state"])

        status, headers, data = self.request("GET", "/api/reports/detailed/export.docx", origin=False)
        self.assertEqual(status, 200)
        self.assertTrue(headers["Content-Disposition"].startswith('attachment; filename="review-report.docx"; filename*=UTF-8\'\''))
        self.assertTrue(urllib.parse.unquote(headers["Content-Disposition"].split("''", 1)[1]).endswith("-审查报告.docx"))
        text = "\n".join(paragraph.text for paragraph in Document(io.BytesIO(data)).paragraphs)
        self.assertIn("基本要素检查", text)
        self.assertIn("历史复核记录（不等于正文操作）：接受", text)
        self.assertIn("工资与交接分开（历史曾选择）", text)
        self.assertIn("人工确认后的工资与交接条款", text)
        self.assertLess(text.index("（一）前段意见"), text.index("（二）后段意见"))

        status, _, conflict = self.json_request("POST", "/api/reports/detailed/state", body)
        self.assertEqual(status, 409)
        self.assertEqual(conflict["current_revision"], 1)

        status, _, undone = self.json_request(
            "POST", "/api/reports/detailed/undo", {"expected_revision": 1}
        )
        self.assertEqual(status, 200)
        self.assertEqual(undone["state"]["revision"], 2)
        self.assertEqual(undone["state"]["reviews"]["F01"]["decision"], "未复核")
        self.assertEqual(undone["state"]["selected_options"], [])
        self.assertFalse(undone["state"]["can_undo"])

    def test_undo_action_diff_survives_reload(self):
        _, _, detail = self.json_request("GET", "/api/reports/detailed", origin=False)
        reviews = detail["state"]["reviews"]
        reviews["F01"]["decision"] = "接受"
        status, _, first = self.json_request(
            "POST", "/api/reports/detailed/state",
            {"expected_revision": 0, "reviews": reviews, "selected_options": []},
        )
        self.assertEqual(status, 200)
        self.assertEqual(first["state"]["undo_action"]["kind"], "review")
        self.assertEqual(first["state"]["undo_action"]["finding_id"], "F01")
        self.assertEqual(first["state"]["undo_action"]["decision"], "接受")
        _, _, reread = self.json_request("GET", "/api/reports/detailed", origin=False)
        self.assertEqual(reread["state"]["undo_action"], first["state"]["undo_action"])

        status, _, second = self.json_request(
            "POST", "/api/reports/detailed/state",
            {"expected_revision": 1, "reviews": reviews, "selected_options": ["O1"]},
        )
        self.assertEqual(status, 200)
        self.assertEqual(second["state"]["undo_action"]["kind"], "option")
        self.assertEqual(second["state"]["undo_action"]["option_id"], "O1")
        self.assertTrue(second["state"]["undo_action"]["selected"])

    def test_unreviewed_export_does_not_claim_human_risk_or_draft(self):
        status, _, data = self.request("GET", "/api/reports/detailed/export.docx", origin=False)
        self.assertEqual(status, 200)
        text = "\n".join(paragraph.text for paragraph in Document(io.BytesIO(data)).paragraphs)
        self.assertTrue(any(cell.text == "待处理"
                            for table in Document(io.BytesIO(data)).tables
                            for row in table.rows for cell in row.cells))
        self.assertNotIn("复核等级：", text)
        self.assertNotIn("复核后条款：", text)
        self.assertIn("4. 建议修改文字", text)
        self.assertIn("工资与交接分开（模型备选）", text)
        self.assertIn("合同以已确认的修改为准", text)

    def test_rejected_finding_exports_reason_not_default_human_draft(self):
        _, _, detail = self.json_request("GET", "/api/reports/detailed", origin=False)
        reviews = detail["state"]["reviews"]
        reviews["F01"].update({
            "decision": "驳回",
            "note": "公司决定不采用本条意见。",
            # Keep the default risk/draft exactly as the state currently does;
            # rejecting the opinion must not present either as human-adopted.
        })
        status, _, _ = self.json_request(
            "POST", "/api/reports/detailed/state",
            {"expected_revision": 0, "reviews": reviews, "selected_options": []},
        )
        self.assertEqual(status, 200)
        status, _, data = self.request("GET", "/api/reports/detailed/export.docx", origin=False)
        self.assertEqual(status, 200)
        text = "\n".join(paragraph.text for paragraph in Document(io.BytesIO(data)).paragraphs)
        self.assertIn("历史复核记录（不等于正文操作）：驳回", text)
        self.assertNotIn("复核等级：", text)
        self.assertNotIn("复核后条款：", text)
        self.assertIn("4. 建议修改文字", text)

    def test_tampering_cross_origin_and_static_escape_are_rejected(self):
        _, _, detail = self.json_request("GET", "/api/reports/detailed", origin=False)
        reviews = detail["state"]["reviews"]
        bad_reviews = {**reviews, "F99": dict(next(iter(reviews.values())))}
        status, _, _ = self.json_request(
            "POST", "/api/reports/detailed/state",
            {"expected_revision": 0, "reviews": bad_reviews, "selected_options": []},
        )
        self.assertEqual(status, 400)
        status, _, _ = self.json_request(
            "POST", "/api/reports/detailed/state",
            {"expected_revision": 0, "reviews": reviews, "selected_options": ["O99"]},
        )
        self.assertEqual(status, 400)

        _, _, failed = self.json_request("GET", "/api/reports/table", origin=False)
        failed_reviews = failed["state"]["reviews"]
        failed_reviews["F01"]["decision"] = "接受"
        status, _, message = self.json_request(
            "POST", "/api/reports/table/state",
            {"expected_revision": 0, "reviews": failed_reviews, "selected_options": []},
        )
        self.assertEqual(status, 400)
        self.assertIn("不能直接接受", message["error"])

        status, _, _ = self.json_request(
            "POST", "/api/reports/detailed/state",
            {"expected_revision": 0, "reviews": reviews, "selected_options": []},
            origin=False,
        )
        self.assertEqual(status, 403)
        status, _, _ = self.json_request(
            "POST", "/api/reports/detailed/state",
            {"expected_revision": 0, "reviews": reviews, "selected_options": []},
            host="evil.example:8518",
        )
        self.assertEqual(status, 421)

        for path in ("/sources.json", "/../sources.json", "/%2e%2e/sources.json", "/api/reports/unknown"):
            status, _, _ = self.json_request("GET", path, origin=False)
            self.assertEqual(status, 404)
        status, headers, body = self.request("GET", "/index.html", origin=False)
        self.assertEqual((status, body), (200, b"<h1>reader</h1>"))
        self.assertNotIn("Access-Control-Allow-Origin", headers)

    def test_corrupt_state_is_not_silently_reset(self):
        request = json.loads(
            (self.root / "runtime" / REPORTS["detailed"]["directory"] / "request.json").read_text(encoding="utf-8")
        )
        self.state_dir.mkdir(parents=True)
        state_path = self.state_dir / f"detailed-{request['request_digest']}.json"
        state_path.write_text("{broken", encoding="utf-8")
        status, _, error = self.json_request("GET", "/api/reports/detailed", origin=False)
        self.assertEqual(status, 500)
        self.assertIn("状态文件损坏", error["error"])
        self.assertEqual(state_path.read_text(encoding="utf-8"), "{broken")

    def test_word_outline_is_presentation_only_and_drives_navigation(self):
        source = self.root / "explicit-outline.docx"
        document = Document()
        document.add_paragraph("第一段合同文字。", style="Heading 3")
        document.add_paragraph("员工未完成交接，公司不支付工资。")
        document.save(source)
        request_path = self.root / "runtime" / REPORTS["detailed"]["directory"] / "request.json"
        original_request = request_path.read_bytes()
        self.server.delivery.source_paths["detailed"] = source
        status, _, data = self.json_request("GET", "/api/reports/detailed")
        self.assertEqual(status, 200, data)
        self.assertEqual(data["contract"]["headings"], {"P000": 3})
        self.assertEqual([(p["id"], p["level"]) for p in data["paragraphs"]], [("P000", 3), ("P001", 0)])
        self.assertEqual(request_path.read_bytes(), original_request)


if __name__ == "__main__":
    unittest.main()
