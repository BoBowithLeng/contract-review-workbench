"""Source-backed positions and typography for the agreed report sample."""
import re
from html import unescape
from urllib.parse import urlsplit

from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


def chinese_number(number):
    digits = '零一二三四五六七八九'
    if number < 10:
        return digits[number]
    if number < 100:
        tens, units = divmod(number, 10)
        return (digits[tens] if tens > 1 else '') + '十' + (digits[units] if units else '')
    hundreds, rest = divmod(number, 100)
    return digits[hundreds] + '百' + (('零' if rest < 10 else '') + chinese_number(rest) if rest else '')


def _short_title(text):
    return bool(text.strip()) and len(text.strip()) <= 36 and not re.search(r'[。；;！？?：:]|应当|不得|必须', text)


def _clause_number(part):
    """Interpret only source-backed digits, never a paragraph ordinal as a clause."""
    match = re.match(r'^(\d+(?:[.．]\d+)*)(?:[.．])?(?:\s|$)', part)
    return match.group(1).replace('．', '.') if match else ''


def source_locations(indexed):
    """Convert index_document positions into real clauses, titles and continuations.

    A numbered short title may be followed by its first unnumbered body paragraph.
    That is not a new clause. Further body paragraphs are explicitly continuations.
    Table coordinates remain coordinates rather than invented article numbers.
    """
    locations, titles, body_counts = {}, {}, {}
    heading_counts = {}
    for p in indexed:
        pid, text = p['id'], p.get('text', '').strip()
        parts = p.get('location', '').split(' / ')
        own_label = p.get('numbering_label', '')
        title = bool(p.get('level') and text) or (_short_title(text) and bool(own_label))
        if title:
            titles[p['location']] = text
        cell = re.fullmatch(r'T(\d+)R(\d+)C(\d+)P(\d+)', pid)
        if cell:
            t, r, c, para = map(int, cell.groups())
            parent = ' / '.join(parts[:-1])
            locations[pid] = (('「' + parent + '」' if parent else '')
                              + f'表格{t+1}第{r+1}行第{c+1}列第{para+1}段')
            continue
        if p.get('numbering_level') is not None and not own_label:
            # Unsupported Word numbering is not an unnumbered continuation of
            # a preceding clause, even if an ancestor remains in the outline.
            context = p.get('location', '')
            locations[pid] = (('「' + context + '」下段落') if context and not context.startswith('正文第')
                              else context or '正文段落') + '（条号未识别）'
            continue
        last = parts[-1] if parts else ''
        clause = _clause_number(last)
        article = re.match(r'^(第[一二三四五六七八九十百零〇\d]+条)', last)
        if clause or article:
            key = p['location']
            if own_label:
                body_counts[key] = 0 if title else 1
                continuation = False
            else:
                count = body_counts.get(key, 0)
                continuation = count >= 1
                if text and not title:
                    body_counts[key] = count + 1
            title_text = next((titles[key] for key in reversed(list(titles))
                               if p['location'] == key or p['location'].startswith(key + ' / ')), '')
            # Unnumbered outline roots may identify a separate attachment. Its
            # numbering can restart at 1 and must not collide with body clauses.
            scope = []
            for part in parts[:-1]:
                if _clause_number(part) or re.match(r'^第[一二三四五六七八九十百零〇\d]+条', part):
                    break
                if not part.endswith(('合同', '协议', '协议书', '合同书')):
                    scope.append(part)
            locations[pid] = (('「' + ' / '.join(scope) + '」') if scope else '') + ('第' + clause + '条' if clause else article.group(1))
            if continuation:
                locations[pid] += '续段'
            if title_text:
                locations[pid] += '「' + title_text + '」'
            continue
        if p.get('level') or title:
            locations[pid] = '「' + text + '」标题'
            heading_counts[p['location']] = 0
        elif p['location'] and not p['location'].startswith('正文第'):
            key = p['location']
            heading_counts[key] = heading_counts.get(key, 0) + 1
            heading_title = ' / '.join(parts)
            locations[pid] = '「' + heading_title + '」下第' + str(heading_counts[key]) + '段'
        else:
            # The source index has an ordinal, not a contractual article number.
            locations[pid] = p['location'] or '正文段落（条号未识别）'
    return locations


def _font(style, name, size, bold=False):
    style.font.name, style.font.size, style.font.bold = name, Pt(size), bold
    style.font.color.rgb = RGBColor(0, 0, 0)
    rpr = style.element.get_or_add_rPr()
    fonts = rpr.rFonts
    if fonts is None:
        fonts = OxmlElement('w:rFonts')
        rpr.insert(0, fonts)
    for attr in list(fonts.attrib):
        if attr.rsplit('}', 1)[-1].lower().endswith('theme'):
            del fonts.attrib[attr]
    for attr in ('ascii', 'hAnsi', 'eastAsia', 'cs'):
        fonts.set(qn('w:' + attr), name)
    color = rpr.find(qn('w:color'))
    if color is not None:
        for attr in list(color.attrib):
            if attr != qn('w:val'):
                del color.attrib[attr]


FIELD_LABELS = {'draft_plan': '建议改稿', 'research_memo': '检索记录'}


class ReportWriter:
    def __init__(self, doc, locations, findings, rules):
        self.doc, self.locations, self.findings, self.rules = doc, locations, findings, rules
        self.section_count = 0
        # Geometry from the agreed sample, not Word's installation defaults.
        for section in doc.sections:
            section.page_width, section.page_height = Cm(21.59), Cm(27.94)
            section.top_margin, section.bottom_margin = Cm(1.82915), Cm(1.778)
            section.left_margin = section.right_margin = Cm(2.10785)
            section.footer_distance = Cm(0.76)
        for style in doc.styles:
            # Remove latent theme selectors even on unused/default styles: Word
            # can inherit them through the style hierarchy and font substitution.
            for node in style.element.iter(qn('w:rFonts')):
                for attr in list(node.attrib):
                    if attr.rsplit('}', 1)[-1].lower().endswith('theme'):
                        del node.attrib[attr]
        for node in doc.styles.element.iter(qn('w:rFonts')):
            for attr in list(node.attrib):
                if attr.rsplit('}', 1)[-1].lower().endswith('theme'):
                    del node.attrib[attr]
            for attr in ('ascii', 'hAnsi', 'eastAsia', 'cs'):
                node.set(qn('w:' + attr), '宋体')
        _font(doc.styles['Normal'], '宋体', 12)
        normal = doc.styles['Normal'].paragraph_format
        normal.line_spacing, normal.first_line_indent = 1.5, Pt(24)
        normal.space_before, normal.space_after = Pt(0), Pt(3)
        normal.widow_control = True
        for name, font, size, bold in [('Title', '黑体', 16, False), ('Heading 1', '黑体', 14, False),
                                       ('Heading 2', '宋体', 12, True), ('Heading 3', '宋体', 12, True)]:
            _font(doc.styles[name], font, size, bold)
            for border in list(doc.styles[name].element.iter(qn('w:pBdr'))):
                border.getparent().remove(border)
            paragraph = doc.styles[name].paragraph_format
            paragraph.first_line_indent = Pt(0)
            paragraph.line_spacing = 1.5
            paragraph.space_before, paragraph.space_after = Pt(12 if name != 'Title' else 0), Pt(6)
            paragraph.keep_with_next = True
            paragraph.keep_together = True
            paragraph.page_break_before = False
            if name == 'Heading 3':
                paragraph.space_before, paragraph.space_after = Pt(5), Pt(3)
        doc.styles['Title'].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for section in doc.sections:
            paragraph = section.footer.paragraphs[0]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.first_line_indent = Pt(0)
            paragraph.paragraph_format.space_before = paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1
            paragraph.add_run('第 ')
            field = OxmlElement('w:fldSimple')
            field.set(qn('w:instr'), 'PAGE')
            run = OxmlElement('w:r')
            properties = OxmlElement('w:rPr')
            fonts = OxmlElement('w:rFonts')
            for attr in ('ascii', 'hAnsi', 'eastAsia', 'cs'):
                fonts.set(qn('w:' + attr), '宋体')
            properties.append(fonts)
            size = OxmlElement('w:sz')
            size.set(qn('w:val'), '18')
            properties.append(size)
            run.append(properties)
            text = OxmlElement('w:t')
            text.text = '1'
            run.append(text)
            field.append(run)
            paragraph._p.append(field)
            paragraph.add_run(' 页')
            for run in paragraph.runs:
                run.font.name, run.font.size = '宋体', Pt(9)
                run._r.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), '宋体')

    def clean(self, text):
        if text is None:
            return ''
        # Only narrative references change. Contract quotes, drafts and legal
        # source titles/URLs bypass this so original business wording is intact.
        # Prose names a position briefly; the clause title is already in the
        # quoted original, so repeating it at every mention only adds noise.
        def joined(match):
            first, last = self.short_location(match.group(1)), self.short_location(match.group(3))
            if not first or not last:
                return match.group(0)
            if first.removesuffix('续段') == last.removesuffix('续段'):
                # Paragraphs under one clause, e.g. an inserted block of (1)...(10) items.
                return first.removesuffix('续段')
            return first + '至' + last

        def replace(match):
            token = match.group(0)
            return (self.short_location(token) or self.findings.get(token) or self.rules.get(token)
                    or ('相关段落' if token.startswith(('P', 'T')) else '相关审查项'))
        text = re.sub(r'(?<![A-Za-z0-9_])(P\d+|T\d+R\d+C\d+P\d+)\s*([-—–~～至])\s*(P\d+|T\d+R\d+C\d+P\d+)(?![A-Za-z0-9_])',
                      joined, str(text))
        # The model sometimes names its own input fields ("按draft_plan修改…", "research_memo").
        for field, label in FIELD_LABELS.items():
            text = re.sub(r'(?<![A-Za-z0-9_])' + field + r'(?![A-Za-z0-9_])', label, text)
        return re.sub(r'(?<![A-Za-z0-9_])(?:T\d+R\d+C\d+P\d+|T\d+|[PCFG]\d{2,})(?![A-Za-z0-9_])', replace, text)

    def short_location(self, pid):
        label = self.locations.get(pid)
        if not label:
            return None
        return re.sub(r'「[^「」]*」$', '', label) or label

    def location(self, pid):
        return self.locations.get(pid, '原文位置待核对')

    def title(self, text):
        return self.doc.add_paragraph(self.clean(text), style='Title')

    def section(self, text):
        self.section_count += 1
        return self.doc.add_paragraph(chinese_number(self.section_count) + '、' + self.clean(text), style='Heading 1')

    def subheading(self, text):
        return self.doc.add_paragraph(self.clean(text), style='Heading 2')

    def detail_heading(self, text):
        return self.doc.add_paragraph(self.clean(text), style='Heading 3')

    def body(self, text, *, raw=False):
        return self.doc.add_paragraph(str(text or '') if raw else self.clean(text))

    def meta(self, text):
        paragraph = self.body(text)
        paragraph.paragraph_format.first_line_indent = Pt(0)
        paragraph.paragraph_format.line_spacing = 1.2
        return paragraph

    def optional(self, label, value):
        if value:
            self.body(label + str(value))

    def located(self, pid, text, *, quote=False, insertion=False, prefix=''):
        label = self.location(pid) + ('之后新增段落' if insertion else '')
        paragraph = self.doc.add_paragraph()
        paragraph.add_run(prefix + label + '：').bold = True
        # Preserve exact source/draft wording; the label describes placement.
        paragraph.add_run('“' + str(text) + '”' if quote else str(text))
        return paragraph

    def labeled(self, label, text):
        paragraph = self.doc.add_paragraph()
        paragraph.add_run(label).bold = True
        paragraph.add_run(str(text))  # Contract wording stays exact.
        return paragraph

    def links(self, ids):
        linked = [self.findings[fid] for fid in ids or [] if fid in self.findings]
        if linked:
            self.body('对应审查意见：' + '、'.join(linked))

    def sources(self, ids, source_map):
        for sid in ids:
            source = source_map.get(sid)
            if source is None:
                self.body('所引用资料记录缺失，需补齐原文后核对。')
                continue
            label = '待核资料：' if source.get('verified') is False else '依据：'
            # A source's name is evidence, not an internal identifier to redact.
            title = unescape(re.sub(r'<[^>]+>', '', source.get('title', '来源名称未提供'))).replace('**', '')
            text = label + title
            if source.get('article'):
                text += ' ' + source['article']
            if source.get('verified') and source.get('checked_at'):
                text += '（核验日 ' + source['checked_at'] + '）'
            paragraph = self.body(text, raw=True)
            url = source.get('url', '')
            if url and urlsplit(url).scheme in {'http', 'https'}:
                paragraph.add_run('  ')
                link = OxmlElement('w:hyperlink')
                link.set(qn('r:id'), self.doc.part.relate_to(url, RT.HYPERLINK, is_external=True))
                run = OxmlElement('w:r')
                properties = OxmlElement('w:rPr')
                color = OxmlElement('w:color'); color.set(qn('w:val'), '245A81')
                underline = OxmlElement('w:u'); underline.set(qn('w:val'), 'single')
                properties.extend([color, underline]); run.append(properties)
                label_node = OxmlElement('w:t'); label_node.text = '查看原文（' + (urlsplit(url).hostname or '来源页面') + '）'
                run.append(label_node); link.append(run); paragraph._p.append(link)

    def overview(self, findings, processing=None):
        if not findings:
            self.body('暂无具体审查意见。')
            return
        table = self.doc.add_table(rows=1, cols=4)
        table.autofit = False
        widths = (1.5, 11.1, 1.5, 3.274)
        for column, width in zip(table.columns, widths):
            column.width = Cm(width)
        properties = table._tbl.tblPr
        borders = OxmlElement('w:tblBorders')
        for side in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
            border = OxmlElement('w:' + side)
            for attr, value in [('val', 'single'), ('sz', '4'), ('color', 'D9D9D9')]:
                border.set(qn('w:' + attr), value)
            borders.append(border)
        properties.append(borders)
        margins = OxmlElement('w:tblCellMar')
        for side, value in [('top', '80'), ('bottom', '80'), ('left', '100'), ('right', '100')]:
            margin = OxmlElement('w:' + side)
            margin.set(qn('w:w'), value)
            margin.set(qn('w:type'), 'dxa')
            margins.append(margin)
        properties.append(margins)
        rows = [('序号', '需要关注的问题', '风险', '处理情况')]
        rows += [(f'{i:02d}', self.clean(f.get('title', '审查意见')), str(f.get('risk') or '待判断'),
                  self.clean((processing or {}).get(f.get('id', f.get('rule_id', '')), '待处理')))
                 for i, f in enumerate(findings, 1)]
        for i, values in enumerate(rows):
            row = table.rows[0] if i == 0 else table.add_row()
            if i == 0:
                row._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
            # Each short overview record stays intact; the table can span pages.
            row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
            for j, (cell, value, width) in enumerate(zip(row.cells, values, widths)):
                cell.width = Cm(width)
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                paragraph = cell.paragraphs[0]
                paragraph.paragraph_format.first_line_indent = Pt(0)
                paragraph.paragraph_format.line_spacing = 1.2
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT if j == 1 else WD_ALIGN_PARAGRAPH.CENTER
                run = paragraph.add_run(value)
                run.bold = i == 0
                run.font.size = Pt(10.5)
                if i == 0:
                    shading = OxmlElement('w:shd')
                    shading.set(qn('w:fill'), 'EEEEEE')
                    cell._tc.get_or_add_tcPr().append(shading)
